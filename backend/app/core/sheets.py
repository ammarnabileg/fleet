"""Rows as CSV (column keys as headers, for other systems) or as Excel for people (BRD FR-RPT-08, FR-AUD-03): headers
in the user's language, right to left in Arabic, numbers as numbers so the spreadsheet can total them, and text that a
spreadsheet would run as a formula kept as text. Shared by the reports and the audit log."""

import csv
import io
import re
from collections.abc import Iterable
from decimal import Decimal

NUMBER = re.compile(r"-?(0|[1-9]\d{0,14})(\.\d+)?")
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _safe(value) -> str:
    """CSV cells that a spreadsheet would run as a formula are prefixed (CSV injection)."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in FORMULA_START and not NUMBER.fullmatch(text) else text


def to_csv(header: list[str], rows: Iterable[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows([[_safe(v) for v in row] for row in rows])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # the BOM makes Excel read Arabic correctly


def _cell(value):
    """A number written as text by the report ("12.500", "3", "-4.000") becomes a number with as many decimals."""
    if value is None or value == "":
        return None, None
    if isinstance(value, int | float | Decimal):
        return value, None
    text = str(value)
    m = NUMBER.fullmatch(text)
    if m:
        decimals = len(m.group(2) or "") - 1 if m.group(2) else 0
        return (Decimal(text) if decimals else int(text)), ("#,##0." + "0" * decimals if decimals else "0")
    return text, "@"  # text stays text, even "=1+1"


def _hf(text: str) -> str:
    return text.replace("&", "&&")[:250]  # Excel's header codes start with &


def page_setup(ws, heading: dict | None) -> None:
    """Printed from Excel like the panel's printouts: A4, and with a heading the company and the title on top of every
    page, the filters, who printed it and when, and «page X of Y» at the bottom."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"  # the column headers on every page
    if not heading:
        return
    ws.oddHeader.center.text = "&B" + _hf(heading.get("company", "")) + "&B\n" + _hf(heading.get("title", ""))
    if heading.get("filters"):
        ws.oddHeader.right.text = _hf(heading["filters"])
    ws.oddFooter.center.text = _hf(heading.get("page", "&P / &N")).replace("&&P", "&P").replace("&&N", "&N")
    printed = " · ".join(x for x in (heading.get("printed_at"), heading.get("printed_by")) if x)
    if printed:
        ws.oddFooter.right.text = _hf(printed)


def to_xlsx(
    title: str,
    header: list[str],
    rows: Iterable[list],
    *,
    rtl: bool,
    text_columns: Iterable[int] = (),
    heading: dict | None = None,
) -> bytes:
    """`text_columns`: positions kept as text however they look, such as account codes ("0101" stays "0101").
    `heading`: the print header (company, title, filters, printed_at, printed_by, page); every sheet prints on A4."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:31] or "Report"
    ws.sheet_view.rightToLeft = rtl
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    widths = [len(h) for h in header]
    keep = set(text_columns)
    # the row number counted here: ws.max_row scans every cell written so far, which made a sheet quadratic
    for n, row in enumerate(rows, start=2):
        values = [("" if v is None else str(v), "@") if i in keep else _cell(v) for i, v in enumerate(row)]
        ws.append([v for v, _ in values])
        for i, (v, number_format) in enumerate(values):
            cell = ws.cell(row=n, column=i + 1)
            if number_format:
                cell.number_format = number_format
            if number_format == "@":
                cell.data_type = "s"  # openpyxl takes any text starting with "=" for a formula: a name must not run
            widths[i] = max(widths[i], len(str(v)) if v is not None else 0)
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = min(max(w + 2, 8), 50)
    ws.freeze_panes = "A2"
    page_setup(ws, heading)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
