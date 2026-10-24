"""A platform's month from Excel, checked before it is applied: the platform's partner report of orders by batch
(Rider ID, Batch No., Total Completed Deliveries; the rider found by the platform driver ID on his file), or a simple
template (the driver's civil ID or platform driver ID, the month, a column per monthly field, and an approved
exception with what it excuses and its note).

The check lists what would happen and what stops it: riders the system does not know, rows that cannot be read
(none is applied in part: the file is corrected first), rows given twice with different figures, batch orders that
differ from the approved daily reports, months already paid. A month that already has batch rows takes another
file only when «استبدال الشهر كله» is chosen: its batch rows are then replaced as a whole by the file's; the same file
again changes nothing, so a month is never counted twice."""

import hashlib
import io
import zipfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.core.text import cell_text, norm
from app.modules.audit import service as audit
from app.modules.payroll import forms, month
from app.modules.payroll.month_models import MonthException, MonthImport, MonthValue
from app.modules.payroll.rules.catalog import EXCUSES
from app.modules.payroll.service import month_start
from app.modules.people import service as people

# what a file's column may feed: who the row is, the month, a batch row (its batch, its completed orders, its period
# when the platform sets the batch week by week), then each monthly field of the platform and an approved exception
BASE_FIELDS = (
    "platform_driver_id",
    "civil_id",
    "month",
    "batch",
    "orders",
    "period_start",
    "period_end",
)
EXCEPTION_COLUMNS = ("exception_kind", "exception_days", "exception_excuses", "exception_note")
ALIASES = {
    "platform_driver_id": ("platform_driver_id", "platform driver id", "driver id", "rider id", "rider", "رقم السائق"),
    "civil_id": ("civil_id", "civil id", "الرقم المدني"),
    "month": ("month", "الشهر"),
    "batch": ("batch no.", "batch no", "batch", "الباتش", "رقم الباتش"),
    "orders": (
        "total completed deliveries", "completed deliveries", "deliveries", "batch_orders", "batch orders",
        "طلبات الباتش", "الطلبات المكتمله", "عدد الطلبات",
    ),
    "period_start": ("period_start", "period start", "week start", "start date", "from", "بدايه الفتره", "من"),
    "period_end": ("period_end", "period end", "week end", "end date", "to", "نهايه الفتره", "الي"),
    "exception_kind": ("exception_kind", "نوع الاستثناء"),
    "exception_days": ("exception_days", "ايام الاستثناء"),
    "exception_excuses": ("exception_excuses", "ما يعذر عنه"),
    "exception_note": ("exception_note", "ملاحظه الاستثناء"),
}  # fmt: skip
# the partner's report of orders by batch (Rider ID, Batch No., Total Completed Deliveries): a preset of the mapping
PARTNER = ("platform_driver_id", "batch", "orders")
PARTNER_ALIASES = {
    "platform_driver_id": ("rider id", "rider", "رقم السائق", "driver id"),
    "batch": ALIASES["batch"],
    "orders": ("total completed deliveries", "completed deliveries", "deliveries", "الطلبات المكتمله", "عدد الطلبات"),
}
GENERIC = ("platform_driver_id", "civil_id", "month", "batch", "orders", *EXCEPTION_COLUMNS)
GENERIC_ALIASES = {
    k: (("batch_orders", "batch orders", "طلبات الباتش") if k == "orders" else ALIASES[k]) for k in GENERIC
}
MAX_ROWS = 5000
MAX_BATCH = 20
HEADER_ROWS = 20


def _workbook(data: bytes):
    import openpyxl  # defusedxml makes openpyxl parse the XML safely

    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise AppError(422, "import_not_xlsx") from None


def _sheet_rows(data: bytes, sheet: str | None = None) -> list[list]:
    wb = _workbook(data)
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise AppError(422, "month_import_sheet", sheet=sheet[:60])
        ws = wb[sheet] if sheet is not None else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)][: MAX_ROWS + HEADER_ROWS]
    finally:
        wb.close()


def _header(rows: list[list], wanted: dict) -> tuple[int, dict]:
    """The first row naming at least two of the wanted columns: its index and where each column is."""
    names = {k: {norm(a) for a in v} for k, v in wanted.items()}
    for i, r in enumerate(rows[:HEADER_ROWS]):
        found = {}
        for j, c in enumerate(r):
            h = norm(c)
            for k, aliases in names.items():
                if h in aliases and k not in found:
                    found[k] = j
        if len(found) >= 2:
            return i, found
    raise AppError(422, "month_import_headers")


def _int(v) -> int:
    d = Decimal(cell_text(v))
    if d != d.to_integral_value() or d < 0:
        raise ValueError
    return int(d)


def _month_of(v) -> date | None:
    if isinstance(v, datetime):
        return month_start(v.date())
    if isinstance(v, date):
        return month_start(v)
    text = cell_text(v)
    if not text:
        return None
    try:
        y, m = text.replace("/", "-").split("-")[:2]
        return date(int(y), int(m), 1)
    except ValueError:
        raise ValueError from None


def _day(v) -> date | None:
    """A date cell: an Excel date, or the text 2026-09-07 or 07/09/2026."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = cell_text(v)
    if not text:
        return None
    parts = text.replace("/", "-").replace(".", "-").split(" ")[0].split("-")
    if len(parts) != 3:
        raise ValueError
    a, b, c = (int(x) for x in parts)
    return date(a, b, c) if a > 31 else date(c, b, a)


class _Parsed:
    def __init__(self):
        self.entries: dict[tuple, dict] = {}  # (driver id, key, sub) -> {"value", "row", "period"}
        self.exceptions: list[dict] = []
        self.unknown: list[dict] = []
        self.duplicates: list[dict] = []
        self.invalid: list[dict] = []
        self.other_month: list[dict] = []
        self.rows = 0

    def add(self, profile: dict, key: str, sub: str, value: Decimal, row: int, period: tuple | None = None) -> None:
        k = (profile["id"], key, sub)
        if k in self.entries:
            same = self.entries[k]["value"] == value and self.entries[k]["period"] == period
            self.duplicates.append(
                {
                    "row": row,
                    "first_row": self.entries[k]["row"],
                    "employee": _who(profile),
                    "key": key,
                    "sub": sub,
                    "same": same,
                }
            )
            if same:
                return
        self.entries[k] = {"value": value, "row": row, "profile": profile, "period": period}


def _who(p: dict) -> dict:
    return {"id": p["public_id"], "name": p["name"], "platform_driver_id": p["platform_driver_id"]}


def _drivers(db: Session, platform_id: int, scope: dict) -> tuple[dict, dict]:
    profiles = people.platform_profiles(db, platform_id, departed=True, **scope)
    profiles = [p for p in profiles if p["platform_id"] == platform_id]
    by_ref = {norm(p["platform_driver_id"]): p for p in profiles if p["platform_driver_id"]}
    by_civil = {norm(p["civil_id"]): p for p in profiles if p["civil_id"]}
    return by_ref, by_civil


def _batch(v) -> int:
    b = _int(v)
    if not 1 <= b <= MAX_BATCH:
        raise ValueError  # a batch 0 is no batch: the row is wrong
    return b


def _exception(cell, p: dict) -> dict | None:
    kind = cell_text(cell("exception_kind"))
    if not kind:
        return None
    if kind not in ("accepted_excuse", "exception_day"):
        raise ValueError  # a company data error is corrected in month review, with its figures
    excuses = [norm(x) for x in cell_text(cell("exception_excuses")).replace("،", ",").split(",") if norm(x)]
    note = cell_text(cell("exception_note"))
    days = _int(cell("exception_days")) if cell_text(cell("exception_days")) else 0
    if not excuses or set(excuses) - set(EXCUSES) or len(note) < 3 or days > 31:
        raise ValueError
    if "valid_days" in excuses and not days:
        raise ValueError
    return {"profile": p, "kind": kind, "days": days, "excuses": sorted(set(excuses)), "note": note}


# ------------------------------------------------------------------ the fields a column may feed


def fields(db: Session, platform_id: int) -> list[dict]:
    """What a file's columns may feed for this platform, in the order the panel lists them: who the row is, the
    month, a batch row, each monthly field of the platform (a task type each), an approved exception."""
    out = [{"key": k, "group": "base", "label": None, "aliases": list(ALIASES[k])} for k in BASE_FIELDS]
    kinds = month.monthly_kinds(db, platform_id)
    for it in forms.get(db, platform_id)["monthly"]:
        kind = kinds.get(it["key"])
        if kind == "task_counts":
            for o in it["options"]:
                out.append(
                    {"key": f"task:{o['value']}", "group": "task", "label": o["label"],
                     "aliases": [o["value"], *o["label"].values()]}
                )  # fmt: skip
        elif kind not in (None, "batch_orders", "choice"):
            out.append(
                {"key": it["key"], "group": "field", "label": it["label"], "kind": kind,
                 "aliases": [it["key"], *it["label"].values()]}
            )  # fmt: skip
    out += [{"key": k, "group": "exception", "label": None, "aliases": list(ALIASES[k])} for k in EXCEPTION_COLUMNS]
    return out


def _aliases(db: Session, platform_id: int, keys=None) -> dict:
    return {f["key"]: tuple(f["aliases"]) for f in fields(db, platform_id) if keys is None or f["key"] in keys}


def _resolve(rows: list[list], columns: dict, known: set) -> tuple[int, dict]:
    """A mapping {field: header text} on the file: the header row (the first naming every chosen header) and where
    each chosen column is."""
    unknown = sorted(set(columns) - known)
    if unknown:
        raise AppError(422, "month_import_field_unknown", field=unknown[0][:60])
    chosen = {k: norm(v) for k, v in columns.items() if cell_text(v)}
    if not chosen.keys() & {"platform_driver_id", "civil_id"}:
        raise AppError(422, "month_import_driver_column")
    if ("batch" in chosen) != ("orders" in chosen):
        raise AppError(422, "month_import_batch_columns")
    if ("period_end" in chosen) and ("period_start" not in chosen):
        raise AppError(422, "month_import_batch_columns")
    for i, r in enumerate(rows[:HEADER_ROWS]):
        where = {norm(c): j for j, c in reversed(list(enumerate(r))) if cell_text(c)}
        if all(h in where for h in chosen.values()):
            return i, {k: where[h] for k, h in chosen.items()}
    missing = next(v for v in columns.values() if cell_text(v))
    raise AppError(422, "month_import_column", column=cell_text(missing)[:80])


def _guess(headers: list, aliases: dict, saved: dict | None) -> dict:
    """The column for each field: the one chosen last time when the file has that header, else one whose header
    names it."""
    present = {norm(h): cell_text(h) for h in headers if cell_text(h)}
    out = {}
    for key, names in aliases.items():
        before = norm((saved or {}).get(key))
        if before and before in present:
            out[key] = present[before]
            continue
        hit = next((present[norm(a)] for a in names if norm(a) in present), None)
        if hit is not None and hit not in out.values():
            out[key] = hit
    return out


def columns(db: Session, *, platform_id: int, data: bytes) -> dict:
    """The file's sheets, each with its header row (the row naming most fields, or the first full one), its headers
    and its first rows; the fields a column may feed, and the mapping pre-selected (last time's, matched by header
    text, else by the headers' names)."""
    from app.modules.payroll.models import Platform

    p = db.get(Platform, platform_id)
    if p is None:
        raise AppError(404, "platform_not_found")
    saved = p.import_mapping or {}
    aliases = _aliases(db, platform_id)
    wb = _workbook(data)
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=HEADER_ROWS + 8)]
            best, score = None, 0
            for i, r in enumerate(rows[:HEADER_ROWS]):
                n = len(_guess(r, aliases, saved.get("columns")))
                if n > score:
                    best, score = i, n
            if best is None:
                best = next((i for i, r in enumerate(rows) if sum(1 for c in r if cell_text(c)) >= 2), None)
            if best is None:
                sheets.append({"name": ws.title, "header_row": None, "headers": [], "sample": []})
                continue
            headers = [cell_text(c) for c in rows[best]]
            while headers and not headers[-1]:
                headers.pop()
            sample = [[_shown(c) for c in r[: len(headers)]] for r in rows[best + 1 :] if any(cell_text(c) for c in r)]
            sheets.append({"name": ws.title, "header_row": best + 1, "headers": headers, "sample": sample[:5]})
    finally:
        wb.close()
    usable = [s for s in sheets if s["headers"]]
    if not usable:
        raise AppError(422, "month_import_headers")
    sheet = next((s for s in usable if s["name"] == saved.get("sheet")), None) or max(
        usable, key=lambda s: len(_guess(s["headers"], aliases, saved.get("columns")))
    )
    return {
        "sheets": sheets,
        "sheet": sheet["name"],
        "fields": [{k: v for k, v in f.items() if k != "aliases"} for f in fields(db, platform_id)],
        "mapping": _guess(sheet["headers"], aliases, saved.get("columns")),
        "saved": bool(saved),
    }


def _shown(c):
    if isinstance(c, datetime | date):
        return c.isoformat()[:10]
    return cell_text(c)


# ------------------------------------------------------------------ reading the rows


def _read(db: Session, rows: list[list], at: int, cols: dict, platform_id: int, the_month: date, refs: tuple,
          out: _Parsed) -> None:  # fmt: skip
    """Each row after the header: who (the platform driver ID, else the civil ID), the month when given, a batch row
    with its period, the platform's fields and task counts, an approved exception. A row is read whole or not at all."""
    by_ref, by_civil = refs
    kinds = month.monthly_kinds(db, platform_id)
    for n, r in enumerate(rows[at + 1 :], start=at + 2):
        if not any(cell_text(c) for c in r):
            continue
        out.rows += 1

        def cell(k, r=r):
            j = cols.get(k)
            return r[j] if j is not None and j < len(r) else None

        ref, civil = cell_text(cell("platform_driver_id")), cell_text(cell("civil_id"))
        staged: list[tuple] = []
        try:  # a batch row that cannot be read is a wrong row, whoever it names
            if cell_text(cell("batch")) or cell_text(cell("orders")):
                batch, orders = _batch(cell("batch")), _int(cell("orders"))
                start, end = _day(cell("period_start")), _day(cell("period_end"))
                if (start is None) != (end is None) and "period_end" in cols:
                    raise ValueError  # a week needs both its dates
                end = end or start
                if start and (end < start or start > statements_month_end(the_month) or end < the_month):
                    raise ValueError
                period = (start, end) if start else None
                staged.append(("batch_orders", month.batch_sub(batch, start), Decimal(orders), period))
        except (ValueError, InvalidOperation, TypeError):
            out.invalid.append({"row": n, "rider": ref or civil})
            continue
        p = (by_ref.get(norm(ref)) if ref else None) or (by_civil.get(norm(civil)) if civil else None)
        if p is None:
            out.unknown.append({"row": n, "rider": ref or civil})
            continue
        try:
            m = _month_of(cell("month"))
            if m is not None and m != the_month:
                out.other_month.append({"row": n, "employee": _who(p), "month": m.isoformat()})
                continue
            for key, kind in kinds.items():
                if key in cols and kind not in ("batch_orders", "task_counts", "choice") and cell_text(cell(key)):
                    raw = cell(key)
                    if kind == "bool":
                        raw = norm(raw) in ("1", "yes", "true", "نعم", "x", "✓")
                    staged.append((key, "", month._number(kind, raw), None))
            for key in cols:
                if key.startswith("task:") and cell_text(cell(key)):
                    staged.append(("task_counts", key[5:], Decimal(_int(cell(key))), None))
            exc = _exception(cell, p)
        except (ValueError, InvalidOperation, TypeError):
            out.invalid.append({"row": n, "rider": ref or civil})
            continue
        for key, sub, value, period in staged:
            out.add(p, key, sub, value, n, period)
        if exc:
            out.exceptions.append(exc | {"row": n})


def statements_month_end(d: date) -> date:
    from app.modules.payroll.statements import month_end

    return month_end(d)


def _layout(db: Session, rows: list[list], *, platform_id: int, fmt: str, mapping: dict | None) -> tuple[int, dict]:
    if fmt == "mapped":
        known = {f["key"] for f in fields(db, platform_id)}
        return _resolve(rows, dict((mapping or {}).get("columns") or {}), known)
    if fmt == "partner_batches":  # the preset of the partner's report: its three columns, all of them
        at, cols = _header(rows, PARTNER_ALIASES)
        if set(cols) != set(PARTNER):
            raise AppError(422, "month_import_headers")
        return at, cols
    wanted = GENERIC_ALIASES | {
        k: v for k, v in _aliases(db, platform_id).items() if k not in BASE_FIELDS and k not in EXCEPTION_COLUMNS
    }
    at, cols = _header(rows, wanted)
    if "civil_id" not in cols and "platform_driver_id" not in cols:
        raise AppError(422, "month_import_headers")
    return at, cols


def _parse(
    db: Session, *, platform_id: int, the_month: date, fmt: str, data: bytes, scope: dict, mapping: dict | None = None
) -> _Parsed:
    rows = _sheet_rows(data, (mapping or {}).get("sheet") if fmt == "mapped" else None)
    at, cols = _layout(db, rows, platform_id=platform_id, fmt=fmt, mapping=mapping)
    out = _Parsed()
    _read(db, rows, at, cols, platform_id, the_month, _drivers(db, platform_id, scope), out)
    return out


def _month_batches(db: Session, platform_id: int, the_month: date) -> dict:
    q = select(MonthValue.employee_id, MonthValue.sub, MonthValue.value, MonthValue.period_end).where(
        MonthValue.platform_id == platform_id, MonthValue.month == the_month, MonthValue.key == "batch_orders"
    )
    return {(e, sub): (v, end) for e, sub, v, end in db.execute(q)}


def _report(
    db: Session, parsed: _Parsed, *, platform_id: int, the_month: date, digest: str, replace_month: bool, approver: bool
) -> dict:
    from app.modules.payroll import statements

    by_driver: dict[int, dict] = defaultdict(lambda: {"batches": [], "values": {}, "tasks": {}})
    for (emp, key, sub), e in parsed.entries.items():
        d = by_driver[emp]
        d["profile"] = e["profile"]
        if key == "batch_orders":
            period = e["period"] or (None, None)
            d["batches"].append((month.batch_of(sub), int(e["value"]), *period))
        elif key == "task_counts":
            d["tasks"][sub] = int(e["value"])
        else:
            d["values"][key] = str(e["value"])
    for x in parsed.exceptions:
        by_driver[x["profile"]["id"]]["profile"] = x["profile"]
    ids = list(by_driver)
    system = statements.system_counts(db, ids, the_month)
    current = month.load(db, dict.fromkeys(ids, platform_id), the_month)
    locked = statements.locked_months(db, {(d["profile"]["company_id"], the_month) for d in by_driver.values()})
    drivers, conflicts, paid = [], [], []
    for emp, d in by_driver.items():
        p = d["profile"]
        in_file = sum(b[1] for b in d["batches"])
        daily = system[emp]["orders"]
        if d["batches"] and daily and in_file != daily:
            conflicts.append({"employee": _who(p), "file_orders": in_file, "daily_orders": daily})
        if (p["company_id"], the_month) in locked:
            paid.append(_who(p))
        had = current[emp]
        drivers.append(
            {
                "employee": _who(p),
                "batches": [month.batch_out(b) for b in sorted(d["batches"], key=lambda b: (b[2] or date.min, b[0]))],
                "values": d["values"],
                "tasks": d["tasks"],
                "daily_orders": daily,
                "replaces": bool(had.batches or had.values or had.tasks),
            }
        )
    again = db.scalar(
        select(MonthImport.id).where(
            MonthImport.platform_id == platform_id, MonthImport.month == the_month, MonthImport.file_sha256 == digest
        )
    )
    # the month's batch rows already there: another file replaces them only when asked to, as a whole
    existing = _month_batches(db, platform_id, the_month)
    in_file = {
        (e, sub): (v["value"], (v["period"] or (None, None))[1])
        for (e, key, sub), v in parsed.entries.items()
        if key == "batch_orders"
    }
    same = bool(in_file) and existing == in_file
    reasons = []
    if parsed.invalid:
        reasons.append("invalid_rows")
    if any(not x["same"] for x in parsed.duplicates):
        reasons.append("duplicate_conflict")
    if paid:
        reasons.append("month_locked")
    if not drivers:
        reasons.append("no_drivers")
    if in_file and existing and not same and not replace_month:
        reasons.append("month_has_batches")
    if parsed.exceptions and not approver:
        reasons.append("exceptions_need_approval")
    return {
        "platform_id": platform_id,
        "month": the_month,
        "rows": parsed.rows,
        "drivers": drivers,
        "unknown": parsed.unknown,
        "invalid": parsed.invalid,
        "duplicates": parsed.duplicates,
        "other_month": parsed.other_month,
        "conflicts": conflicts,
        "locked": paid,
        "exceptions": [
            {
                "row": x["row"],
                "employee": _who(x["profile"]),
                "kind": x["kind"],
                "days": x["days"],
                "excuses": x["excuses"],
                "note": x["note"],
            }  # fmt: skip
            for x in parsed.exceptions
        ],
        "imported_before": again is not None,
        "existing_batches": len(existing),
        "same_as_existing": same,
        "replace_month": replace_month,
        "reasons": reasons,
        "blocking": bool(reasons),
    }


def check(
    db: Session, *, platform_id: int, month_: date, fmt: str, data: bytes, replace_month: bool = False,
    approver: bool = False, mapping: dict | None = None, **scope,
) -> dict:  # fmt: skip
    the_month = month_start(month_)
    parsed = _parse(db, platform_id=platform_id, the_month=the_month, fmt=fmt, data=data, scope=scope, mapping=mapping)
    digest = hashlib.sha256(data).hexdigest()
    return _report(
        db, parsed, platform_id=platform_id, the_month=the_month, digest=digest, replace_month=replace_month,
        approver=approver,
    )  # fmt: skip


def apply(
    db: Session, *, platform_id: int, month_: date, fmt: str, data: bytes, file_name: str | None, actor_user_id: int,
    replace_month: bool = False, approver: bool = False, mapping: dict | None = None, **scope,
) -> dict:  # fmt: skip
    the_month = month_start(month_)
    digest = hashlib.sha256(data).hexdigest()
    parsed = _parse(db, platform_id=platform_id, the_month=the_month, fmt=fmt, data=data, scope=scope, mapping=mapping)
    report = _report(
        db, parsed, platform_id=platform_id, the_month=the_month, digest=digest, replace_month=replace_month,
        approver=approver,
    )  # fmt: skip
    if report["blocking"]:
        raise AppError(409, "month_import_blocked")
    imp = MonthImport(
        platform_id=platform_id,
        month=the_month,
        format=fmt,
        file_sha256=digest,
        file_name=(file_name or "")[:200] or None,
        summary={
            "rows": report["rows"],
            "drivers": len(report["drivers"]),
            "unknown": len(report["unknown"]),
            "replace_month": replace_month,
        },  # fmt: skip
        created_by=actor_user_id,
    )
    db.add(imp)
    db.flush()
    batches = {(e, sub): v for (e, key, sub), v in parsed.entries.items() if key == "batch_orders"}
    if batches and not report["same_as_existing"]:  # the month's batch rows, replaced as a whole by the file's
        db.execute(
            delete(MonthValue).where(
                MonthValue.platform_id == platform_id, MonthValue.month == the_month, MonthValue.key == "batch_orders"
            )
        )
        db.flush()
    applied = set()
    for (emp, key, sub), e in parsed.entries.items():
        if key == "batch_orders" and report["same_as_existing"]:
            continue  # the same rows: nothing changes
        ref = people.ref(db, emp)
        month._put(db, ref, the_month, key, sub, e["value"], source="import", actor_user_id=actor_user_id,
                   import_id=imp.id, period=e["period"])  # fmt: skip
        applied.add(emp)
    for x in parsed.exceptions:
        ref = people.ref(db, x["profile"]["id"])
        live = db.scalars(
            select(MonthException).where(
                MonthException.employee_id == ref.id,
                MonthException.month == the_month,
                MonthException.cancelled_at.is_(None),
            )
        ).all()
        if any((e.kind, e.days, sorted(e.excuses or []), e.note) == (x["kind"], x["days"], x["excuses"], x["note"])
               for e in live):  # fmt: skip
            continue  # imported before: never twice
        db.add(
            MonthException(
                employee_id=ref.id,
                company_id=ref.company_id,
                month=the_month,
                kind=x["kind"],
                days=x["days"],
                excuses=x["excuses"],
                corrections={},
                note=x["note"],
                approved_by=actor_user_id,
            )  # fmt: skip
        )
        applied.add(ref.id)
    if fmt == "mapped":  # the next file of this platform starts from these columns
        from app.modules.payroll.models import Platform

        db.get(Platform, platform_id).import_mapping = {
            "sheet": (mapping or {}).get("sheet"),
            "columns": {k: cell_text(v) for k, v in ((mapping or {}).get("columns") or {}).items() if cell_text(v)},
        }
    db.flush()
    audit.record(
        db,
        action="payroll.month_imported",
        entity_type="month_import",
        entity_id=imp.public_id,
        actor_user_id=actor_user_id,
        after={
            "platform_id": platform_id,
            "month": the_month.isoformat(),
            "format": fmt,
            "drivers": len(applied),
            "exceptions": len(parsed.exceptions),
            "replace_month": replace_month,
            "file_sha256": digest,
        },
    )
    db.commit()
    return report | {"import_id": str(imp.public_id), "applied": len({d["employee"]["id"] for d in report["drivers"]})}


def template(db: Session, platform_id: int) -> bytes:
    """The simple template for a platform: who, the month, a column per monthly field, an approved exception."""
    import openpyxl

    f = forms.get(db, platform_id)
    headers = ["platform_driver_id", "civil_id", "month"]
    for it in f["monthly"]:
        if it["type"] == "batch_orders":
            headers += ["batch", "batch_orders"]
        elif it["type"] == "task_counts":
            headers += [o["value"] for o in it["options"]]
        elif it["type"] != "choice":
            headers.append(it["key"])
    headers += list(EXCEPTION_COLUMNS)
    wb = openpyxl.Workbook()
    wb.active.title = "month"
    wb.active.append(headers)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
