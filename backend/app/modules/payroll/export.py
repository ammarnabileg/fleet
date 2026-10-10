"""The payroll as the client's Excel: one sheet per platform, with exactly that platform's columns, in its order and
with its headers (the template it was set up from), right to left. Employees without a platform get a sheet with
every column. This is the file for the bank and for the platforms (FR-PAY-06)."""

import io
import re
from decimal import Decimal

from sqlalchemy.orm import Session

from app.modules.i18n import service as i18n
from app.modules.payroll import platforms
from app.modules.payroll.columns import BY_CODE
from app.modules.payroll.models import Line, Run
from app.modules.payroll.runs import sheet_columns


def _title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "Sheet"
    title, n = base, 2
    while title in used:
        title, n = f"{base} {n}", n + 1
    used.add(title)
    return title


def _value(code: str, raw, labels: dict):
    kind = BY_CODE[code].kind if code in BY_CODE else "text"
    if raw is None or raw == "":
        return None
    if code == "payment_method":
        return labels.get("payment_method", {}).get(raw, raw)
    if kind == "money":
        return Decimal(str(raw))
    if kind == "number":
        return float(raw)
    if kind == "int":
        return int(raw)
    return str(raw)


def workbook(db: Session, run: Run, lines: list[Line]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    lang = i18n.default_language(db).code
    labels = i18n.effective_catalog(db, lang)
    plats = platforms.all_by_id(db)
    groups: dict[int | None, list[Line]] = {}
    for line in lines:
        groups.setdefault(line.platform_id, []).append(line)
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    order = sorted(groups, key=lambda pid: (pid is None, pid or 0))
    for pid in order or [None]:
        platform = plats.get(pid)
        name = (
            (platform.name.get(lang) or next(iter(platform.name.values())))
            if platform
            else labels.get("common", {}).get("no_platform", "—")
        )
        ws = wb.create_sheet(_title(name, used))
        ws.sheet_view.rightToLeft = True
        cols = sheet_columns(platform, labels.get("payroll_column", {}), [line.cells for line in groups.get(pid, [])])
        ws.append([c["header"] for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for line in groups.get(pid, []):
            ws.append([_value(c["code"], line.cells.get(c["code"]), labels) for c in cols])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # openpyxl takes any text starting with "=" for a formula: a name must not run
        for i, c in enumerate(cols, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = max(12, min(40, len(c["header"]) + 4))
            if BY_CODE.get(c["code"]) and BY_CODE[c["code"]].kind == "money":
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = "#,##0.000"
            elif c["code"] in ("civil_id", "iban", "platform_driver_id"):
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = "@"  # text: Excel would show a civil ID as 2.86E+11
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
