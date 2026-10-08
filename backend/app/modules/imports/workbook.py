"""Reading the onboarding workbook (the data template annexed to the contract) exactly as it is laid out:
same sheet names, same column headers in the same order, dates as day/month/year, phones without country code.
After the people columns come optional driver columns (app, initial password, its days), found by their header, so a
file made before them still imports."""

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.core.errors import AppError
from app.core.text import norm

VEHICLES = "السيارات"
PEOPLE = "المستخدمون والسائقون"
VEHICLE_COLUMNS = [
    "#",
    "رقم اللوحة *",
    "النوع *",
    "الموديل *",
    "السنة *",
    "اللون *",
    "رقم الهيكل (VIN) *",
    "الشركة *",
    "الفرع *",
    "تاريخ انتهاء التأمين *",
    "تاريخ انتهاء الاستمارة *",
]
PEOPLE_COLUMNS = [
    "#",
    "الاسم *",
    "الدور *",
    "الرقم المدني *",
    "رقم الهاتف (لرمز التحقق) *",
    "الشركة *",
    "الفرع *",
    "القسم *",
    "الوظيفة *",
    "الحالة الوظيفية *",
    "الراتب الأساسي (د.ك) *",
    "تاريخ انتهاء الإقامة *",
    "موديل الهاتف (للسائقين)",
]
# optional, after the people columns (N to P in the template), each found by its header wherever it is
DRIVER_COLUMNS = ["تفعيل التطبيق (نعم/لا)", "كلمة المرور المبدئية", "صلاحية كلمة المرور (يوم)"]
# what a client's own sheet may call them too (compared through norm())
PEOPLE_OPTIONAL = {
    "app_access": (
        DRIVER_COLUMNS[0],
        "تفعيل التطبيق",
        "التطبيق",
        "تطبيق السائق",
        "التطبيق مفعل",
        "app",
        "app access",
        "driver app",
        "app active",
        "activate app",
    ),
    "initial_password": (
        DRIVER_COLUMNS[1],
        "كلمة المرور",
        "كلمة السر",
        "الرقم السري",
        "الباسورد",
        "البسورد",
        "باسورد",
        "password",
        "initial password",
        "default password",
    ),
    "password_days": (
        DRIVER_COLUMNS[2],
        "صلاحية كلمة المرور",
        "مدة الصلاحية",
        "صالحة لمدة",
        "عدد الايام",
        "ايام الصلاحية",
        "days",
        "valid days",
        "validity days",
        "password days",
    ),
}
OPENING = "الأرصدة الافتتاحية"
OPENING_COLUMNS = [
    "#",
    "اسم السائق *",
    "الرقم المدني *",
    "الرصيد الافتتاحي (د.ك) *",
    "تاريخ الرصيد *",
    "اعتماد المحاسب (الاسم) *",
]
INSTRUCTIONS = "التعليمات"
EXAMPLE = "مثال"  # the grey example row is never imported
DRIVER_ROLE = "سائق"
COUNTRY_CODE = "965"


@dataclass
class Issue:
    sheet: str
    row: int
    code: str
    params: dict = field(default_factory=dict)


@dataclass
class Row:
    sheet: str
    number: int  # the row number in Excel, for the error list
    values: dict


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def parse_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", _text(v))
    if not m:
        return None
    try:
        return date(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        return None


def parse_phone(v) -> str | None:
    digits = re.sub(r"\D", "", _text(v))
    if len(digits) == 8:
        return f"+{COUNTRY_CODE}{digits}"
    if len(digits) == 11 and digits.startswith(COUNTRY_CODE):
        return f"+{digits}"
    return None


YES = {"نعم", "ن", "yes", "y", "1", "true", "مفعل"}
NO = {"لا", "no", "n", "0", "false", "غير مفعل"}
_DIACRITICS = re.compile("[\u064b-\u0652]")  # "مفعّل" is "مفعل"
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "0123456789" * 2)


def parse_yes_no(v) -> bool | None:
    """نعم or لا, yes or no, 1 or 0, an Excel TRUE/FALSE cell; None when it is none of them."""
    if isinstance(v, bool):
        return v
    t = _DIACRITICS.sub("", norm(v))
    return True if t in YES else False if t in NO else None


def parse_password(v) -> str | None:
    """The password as the driver will type it: text as written (spaces inside count), a number without ".0". None for
    a date or TRUE/FALSE cell, a number longer than Excel keeps exactly (15 digits), or a length outside 8 to 64."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int) and not isinstance(v, bool) and len(str(abs(v))) <= 15:
        v = str(v)
    if not isinstance(v, str):
        return None
    v = v.strip()
    return v if 8 <= len(v) <= 64 else None


def parse_days(v) -> int | None:
    """How many days a password stays valid, 1 to 60, also written in Arabic-Indic digits ("١٤")."""
    t = _text(v).translate(_DIGITS)
    return int(t) if re.fullmatch(r"[0-9]{1,2}", t) and 1 <= int(t) <= 60 else None


def parse_decimal(v, *, signed: bool = False) -> Decimal | None:
    try:
        d = Decimal(_text(v).replace(",", ""))
    except InvalidOperation:
        return None
    if not d.is_finite() or d.as_tuple().exponent < -3:
        return None
    return d if signed or d >= 0 else None


def read(data: bytes) -> dict[str, list[Row]]:
    """Rows of the sheets imported now (vehicles, people, opening balances); maintenance centres come with M3."""
    import openpyxl  # with defusedxml installed, openpyxl parses the XML safely

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise AppError(422, "import_not_xlsx") from None
    optional = {norm(a): key for key, aliases in PEOPLE_OPTIONAL.items() for a in aliases}
    out: dict[str, list[Row]] = {}
    for sheet, columns in ((VEHICLES, VEHICLE_COLUMNS), (PEOPLE, PEOPLE_COLUMNS), (OPENING, OPENING_COLUMNS)):
        if sheet not in wb.sheetnames:
            raise AppError(422, "import_bad_template", sheet=sheet)
        rows = wb[sheet].iter_rows(values_only=True)
        header = list(next(rows, ()))
        if [_text(c) for c in header[: len(columns)]] != columns:
            raise AppError(422, "import_bad_template", sheet=sheet)
        extra: dict[str, int] = {}  # optional column -> its index, the first one with that header
        if sheet == PEOPLE:
            for i in range(len(columns), len(header)):
                key = optional.get(norm(header[i]))
                if key and key not in extra:
                    extra[key] = i
        width = max([len(columns), *(i + 1 for i in extra.values())])
        out[sheet] = []
        for number, values in enumerate(rows, start=2):
            values = list(values) + [None] * (width - len(values))
            # a row with only the optional columns filled (an app column filled down) is still an empty row
            if _text(values[0]) == EXAMPLE or not any(_text(v) for v in values[1 : len(columns)]):
                continue
            cells = dict(zip(columns, values[: len(columns)], strict=True))
            out[sheet].append(Row(sheet, number, cells | {key: values[i] for key, i in extra.items()}))
    wb.close()
    return out


def text(row: Row, column: str) -> str:
    return _text(row.values.get(column))


GUIDE = (
    "قالب استيراد السيارات والمستخدمين والسائقين والأرصدة الافتتاحية",
    "لا تغيّر أسماء الأوراق ولا عناوين الأعمدة ولا ترتيبها. الصف الرمادي (مثال) لا يُستورد، ولا الصفوف الفارغة.",
    "التواريخ يوم/شهر/سنة (31/12/2027)، والهاتف 8 أرقام دون مفتاح الدولة، والشركة والفرع والحالة الوظيفية "
    "بأسمائها في النظام، والدور «سائق» للسائقين.",
    "أعمدة اختيارية للسائقين في ورقة «المستخدمون والسائقون» (N إلى P):",
    "• تفعيل التطبيق: نعم أو لا. فارغ أو نعم: يُفعَّل التطبيق للسائق. لا: لا يُفعَّل، ولا يوقف الاستيراد تطبيق "
    "سائق مفعّل من قبل (ذلك من ملف السائق).",
    "• كلمة المرور المبدئية: يدخل بها السائق أول مرة مع رقمه المدني. من 8 إلى 64 حرفاً، وليست رقمه المدني. "
    "تحتاج صلاحية «إدارة أجهزة السائقين». إعادة الاستيراد لا تغيّرها لسائق له كلمة مفتوحة أو دخل بها من قبل.",
    "• صلاحية كلمة المرور (يوم): من 1 إلى 60، وإن تُركت فارغة فـ 14 يوماً من وقت الاستيراد.",
    "الملف بعد كتابة كلمات المرور سري: احذفه بعد الاستيراد. النظام لا يحفظ الملف، ويحفظ كلمة المرور مشفّرة فقط.",
)
EXAMPLES = {
    VEHICLES: [
        "12/34567",
        "تويوتا",
        "يارس",
        2024,
        "أبيض",
        "MR0AA0000BB000001",
        "اسم الشركة",
        "الفرع الرئيسي",
        "31/12/2027",
        "31/12/2027",
    ],
    PEOPLE: [
        "اسم الموظف",
        DRIVER_ROLE,
        "290010112345",
        "98765432",
        "اسم الشركة",
        "الفرع الرئيسي",
        "العمليات",
        "سائق توصيل",
        "على رأس العمل",
        "150.000",
        "30/09/2027",
        "Samsung A15",
        "نعم",
        "(8 أحرف على الأقل)",
        14,
    ],
    OPENING: ["اسم السائق", "290010112345", "12.500", "01/10/2026", "اسم المحاسب"],
}
TEMPLATE_ROWS = 1000  # how far down the sheet the password column is text and N and P are checked


def template() -> bytes:
    """The empty workbook to fill: the sheets and headers read() expects with the grey example row. The password
    column is text, so a password such as 01234567 keeps its leading zero; Excel offers نعم/لا in N and 1 to 60 in P."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = INSTRUCTIONS
    ws.sheet_view.rightToLeft = True
    for line in GUIDE:
        ws.append([line])
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 140
    people = PEOPLE_COLUMNS + DRIVER_COLUMNS
    for title, header in ((VEHICLES, VEHICLE_COLUMNS), (PEOPLE, people), (OPENING, OPENING_COLUMNS)):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(header)
        ws.append([EXAMPLE, *EXAMPLES[title]])
        for i, h in enumerate(header, start=1):
            ws.cell(1, i).font = Font(bold=True)
            ws.cell(2, i).fill = PatternFill("solid", fgColor="D9D9D9")
            ws.column_dimensions[get_column_letter(i)].width = max(len(h) + 4, 12)
        ws.freeze_panes = "A2"
    ws = wb[PEOPLE]
    app, password, days = (get_column_letter(len(PEOPLE_COLUMNS) + i) for i in (1, 2, 3))
    ws.column_dimensions[password].number_format = "@"
    for r in range(2, TEMPLATE_ROWS + 1):
        ws[f"{password}{r}"].number_format = "@"
    yes_no = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True, showErrorMessage=True)
    yes_no.error = "نعم أو لا"
    whole = DataValidation(
        type="whole", operator="between", formula1="1", formula2="60", allow_blank=True, showErrorMessage=True
    )
    whole.error = "عدد أيام من 1 إلى 60"
    for validation, column in ((yes_no, app), (whole, days)):
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{TEMPLATE_ROWS}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
