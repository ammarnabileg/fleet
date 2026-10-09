"""The accountant's settings and his own entries: automatic or manual approval, the books' start, the fiscal year,
manual entries, opening balances to the opening balances account, the old
system's books from Excel, and the treasury's deposit rule."""

import io
from datetime import date, timedelta

import openpyxl
import pytest

from app.core.clock import today
from tests.conftest import fund_treasury, login, make_driver, make_employee, make_user, upload
from tests.test_finance import main_branch

F = "/api/v1/finance"


def settings(client, section: str, **value) -> dict:
    cur = client.get("/api/v1/settings").json()[section]
    r = client.put(f"/api/v1/settings/{section}", json={"version": cur["version"], "value": cur["value"] | value})
    assert r.status_code == 200, r.text
    return r.json()


def accounts(client) -> dict[str, int]:
    return {a["code"]: a["id"] for a in client.get(f"{F}/accounts").json()}


def period() -> dict:
    first = today().replace(day=1)
    return {"date_from": str(first), "date_to": str(today())}


def manual(client, lines, **kw) -> dict:
    acc = accounts(client)
    body = {
        "entry_date": str(today()),
        "description": "Rent accrual",
        "lines": [
            {"account_id": acc[code], "debit": d, "credit": c} | extra
            for code, d, c, *rest in lines
            for extra in [rest[0] if rest else {}]
        ],
    } | kw
    return client.post(f"{F}/entries/manual", json=body)


@pytest.fixture
def float_(db):
    """Cash in the main treasury from long ago, for the expenses paid from it."""
    fund_treasury(db)


def expense(client, company_id, amount="10"):
    types = {t["code"]: t["id"] for t in client.get(f"{F}/expense-types").json()}
    r = client.post(
        f"{F}/expenses",
        json={
            "company_id": company_id,
            "type_id": types["fuel"],
            "expense_date": str(today()),
            "amount": amount,
            "payment_method": "treasury",
            "branch_id": main_branch(client),
        },
    )
    assert r.status_code == 201, r.text
    assert client.post(f"{F}/expenses/{r.json()['id']}/approve", json={}).status_code == 200
    return r.json()


def balances(client, **params) -> dict[str, str]:
    r = client.get(f"{F}/trial-balance", params=period() | params)
    assert r.status_code == 200, r.text
    return {b["account"]["code"]: b["closing"] for b in r.json()}


# ------------------------------------------------------------------ approval: manual or automatic


def test_generated_and_manual_entries_follow_the_approval_setting(admin_client, company, float_):
    expense(admin_client, company["id"], "10")
    r = admin_client.post(f"{F}/entries/post", json=period())
    assert r.json()["created"] == 1
    [draft] = admin_client.get(f"{F}/entries", params=period()).json()
    assert draft["status"] == "draft"
    r = manual(admin_client, [("6130", "5", "0"), ("2110", "0", "5")])
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "draft" and r.json()["source_ref"] == f"MAN-{r.json()['number']}"

    # automatic: the drafts already there stay drafts; what is made from now on is approved by the system
    settings(admin_client, "finance", entry_approval="auto")
    assert admin_client.get(f"{F}/config").json()["entry_approval"] == "auto"
    assert admin_client.get(f"{F}/entries/{draft['id']}").json()["status"] == "draft"
    expense(admin_client, company["id"], "7")
    assert admin_client.post(f"{F}/entries/post", json=period()).json()["created"] == 1
    made = [e for e in admin_client.get(f"{F}/entries", params=period()).json() if e["amount"] == "7.000"]
    assert made[0]["status"] == "approved" and made[0]["approved_by"] is None and made[0]["approved_at"]
    r = manual(admin_client, [("6130", "3", "0"), ("2110", "0", "3")])
    assert r.json()["status"] == "approved" and r.json()["approved_by"] is None
    assert balances(admin_client)["6130"] == "3.000"

    # the nightly posting too
    from app.modules.finance import tasks

    expense(admin_client, company["id"], "2")
    assert tasks.post_entries()["created"] == 1
    assert [
        e["status"] for e in admin_client.get(f"{F}/entries", params=period()).json() if e["amount"] == "2.000"
    ] == ["approved"]
    # a reversal stays as it was: approved at once, by who reversed it
    rev = admin_client.post(f"{F}/entries/{made[0]['id']}/reverse", json={"reason": "wrong one"}).json()
    assert rev["status"] == "approved" and rev["approved_by"]


# ------------------------------------------------------------------ manual entries


def test_a_manual_entry_is_balanced_one_sided_on_open_accounts(admin_client, company):
    acc = accounts(admin_client)
    r = manual(admin_client, [("6130", "5", "0"), ("2110", "0", "4")])
    assert r.status_code == 422 and r.json()["code"] == "entry_unbalanced"
    assert r.json()["params"] == {"difference": "1.000"}
    r = manual(admin_client, [("6130", "5", "5"), ("2110", "0", "5")])
    assert r.status_code == 422 and r.json()["code"] == "line_one_side"
    r = manual(admin_client, [("6130", "0", "0"), ("2110", "0", "0")])
    assert r.status_code == 422 and r.json()["code"] == "line_one_side"
    assert manual(admin_client, [("6130", "5", "0")]).status_code == 422  # two lines at least
    many = [("6130", "1", "0")] * 100 + [("2110", "0", "1")] * 101
    assert manual(admin_client, many).status_code == 422  # 200 at most
    r = manual(admin_client, [("6130", "5", "0"), ("2110", "0", "5")], description="ab")
    assert r.status_code == 422  # a description of 3 characters at least
    # a closed account
    new = admin_client.post(f"{F}/accounts", json={"code": "6999", "name": {"ar": "س", "en": "x"}, "type": "expense"})
    a = new.json()
    admin_client.patch(f"{F}/accounts/{a['id']}", json={"version": a["version"], "active": False})
    r = manual(admin_client, [("6999", "5", "0"), ("2110", "0", "5")])
    assert r.status_code == 422 and r.json()["code"] == "account_inactive"
    r = admin_client.post(
        f"{F}/entries/manual",
        json={"entry_date": str(today()), "description": "xyz", "lines": [
            {"account_id": 999_999, "debit": "1"}, {"account_id": acc["2110"], "credit": "1"}]},
    )  # fmt: skip
    assert r.status_code == 404 and r.json()["code"] == "account_not_found"

    # with a company and a party: the statement shows the employee, the entry its lines
    emp = make_employee(admin_client, company["id"], name={"ar": "سالم", "en": "Salem"})
    r = manual(
        admin_client,
        [("1150", "20", "0", {"employee_id": emp["id"], "memo": "loan"}), ("1110", "0", "20")],
        company_id=company["id"],
    )
    assert r.status_code == 201, r.text
    entry = r.json()
    assert entry["company_id"] == company["id"] and entry["source_kind"] == "manual"
    assert entry["lines"][0]["employee"]["id"] == emp["id"] and entry["lines"][0]["memo"] == "loan"
    assert admin_client.post(f"{F}/entries/approve", json={"ids": [entry["id"]]}).json() == {"count": 1}
    book = admin_client.get(f"{F}/ledger", params=period() | {"account_id": acc["1150"]}).json()
    assert [(ln["party"]["name"]["en"], ln["debit"]) for ln in book["lines"]] == [("Salem", "20.000")]
    assert book["parties"][0]["balance"] == "20.000"
    # reversed: the reversal is with him too
    admin_client.post(f"{F}/entries/{entry['id']}/reverse", json={"reason": "wrong amount"})
    book = admin_client.get(f"{F}/ledger", params=period() | {"account_id": acc["1150"]}).json()
    assert [ln["party"]["name"]["en"] for ln in book["lines"]] == ["Salem", "Salem"] and book["parties"] == []


def test_a_manual_draft_is_deleted_alone_and_kept_by_the_periods_discard(admin_client, company, float_):
    expense(admin_client, company["id"])
    admin_client.post(f"{F}/entries/post", json=period())
    mine = manual(admin_client, [("6130", "5", "0"), ("2110", "0", "5")]).json()
    # discarding the period's drafts (to make them again) leaves the accountant's own
    assert admin_client.post(f"{F}/entries/discard", json=period()).json() == {"count": 1}
    left = admin_client.get(f"{F}/entries", params=period()).json()
    assert [e["id"] for e in left] == [mine["id"]]
    admin_client.post(f"{F}/entries/post", json=period())
    generated = next(
        e for e in admin_client.get(f"{F}/entries", params=period()).json() if e["source_kind"] == "expense"
    )
    r = admin_client.delete(f"{F}/entries/{generated['id']}")
    assert r.status_code == 409 and r.json()["code"] == "entry_not_deletable"
    assert admin_client.delete(f"{F}/entries/{mine['id']}").status_code == 204
    assert admin_client.get(f"{F}/entries/{mine['id']}").status_code == 404
    approved = manual(admin_client, [("6130", "5", "0"), ("2110", "0", "5")]).json()
    admin_client.post(f"{F}/entries/approve", json={"ids": [approved["id"]]})
    assert admin_client.delete(f"{F}/entries/{approved['id']}").json()["code"] == "entry_not_deletable"


def test_permissions_of_manual_and_opening_entries(admin_client, new_client, companies):
    make_user(admin_client, "writer", permissions=["finance.view", "finance.create"])
    c = new_client()
    login(c, "writer")
    assert manual(c, [("6130", "5", "0"), ("2110", "0", "5")]).status_code == 201
    acc = accounts(c)
    r = c.post(f"{F}/opening", json={"lines": [{"account_id": acc["1120"], "debit": "1"}]})
    assert r.status_code == 403  # finance.approve
    assert c.post(f"{F}/import", files={"file": ("b.xlsx", b"x")}).status_code == 403
    make_user(
        admin_client, "limited", permissions=["finance.view", "finance.create"], company_ids=[companies["a"]["id"]]
    )
    c = new_client()
    login(c, "limited")
    r = manual(c, [("6130", "5", "0"), ("2110", "0", "5")])
    assert r.status_code == 403 and r.json()["code"] == "all_companies_required"


# ------------------------------------------------------------------ the books' start and opening balances


def test_nothing_before_the_books_start(admin_client, company, float_):
    start = today() - timedelta(days=3)
    old = expense(admin_client, company["id"], "4")
    settings(admin_client, "finance", books_start_date=str(start))
    r = manual(admin_client, [("6130", "5", "0"), ("2110", "0", "5")], entry_date=str(start - timedelta(days=1)))
    assert r.status_code == 422 and r.json()["code"] == "before_books_start"
    assert r.json()["params"] == {"date": str(start)}
    assert manual(admin_client, [("6130", "5", "0"), ("2110", "0", "5")], entry_date=str(start)).status_code == 201
    # a document dated before the start is not entered; one after is
    types = {t["code"]: t["id"] for t in admin_client.get(f"{F}/expense-types").json()}
    early = admin_client.post(
        f"{F}/expenses",
        json={"company_id": company["id"], "type_id": types["fuel"], "expense_date": str(start - timedelta(days=2)),
              "amount": "9", "payment_method": "bank"},
    ).json()  # fmt: skip
    admin_client.post(f"{F}/expenses/{early['id']}/approve", json={})
    r = admin_client.post(
        f"{F}/entries/post", json={"date_from": str(start - timedelta(days=20)), "date_to": str(today())}
    )
    assert r.json()["by_kind"] == {"expense": 1} and old  # today's, not the early one
    # a reversal dated before the start is refused too
    e = next(x for x in admin_client.get(f"{F}/entries", params=period()).json() if x["source_kind"] == "expense")
    admin_client.post(f"{F}/entries/approve", json={"ids": [e["id"]]})
    r = admin_client.post(
        f"{F}/entries/{e['id']}/reverse", json={"reason": "wrong one", "entry_date": str(start - timedelta(days=1))}
    )
    assert r.status_code == 422 and r.json()["code"] == "before_books_start"


def test_opening_balances_with_the_difference_to_opening_equity(admin_client, companies):
    acc = accounts(admin_client)
    lines = [
        {"account_id": acc["1120"], "debit": "1500"},
        {"account_id": acc["1110"], "debit": "200"},
        {"account_id": acc["2110"], "credit": "300"},
    ]
    r = admin_client.post(f"{F}/opening", json={"lines": lines})
    assert r.status_code == 422 and r.json()["code"] == "books_start_required"
    start = today().replace(day=1)
    settings(admin_client, "finance", books_start_date=str(start))
    r = admin_client.post(f"{F}/opening", json={"lines": lines})
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["difference"] == "1400.000" and out["equity_account"]["code"] == "3900"
    entry = out["entry"]
    assert entry["entry_date"] == str(start - timedelta(days=1)) and entry["source_kind"] == "opening"
    assert [(ln["account"]["code"], ln["debit"], ln["credit"]) for ln in entry["lines"]] == [
        ("1120", "1500.000", "0.000"),
        ("1110", "200.000", "0.000"),
        ("2110", "0.000", "300.000"),
        ("3900", "0.000", "1400.000"),
    ]
    # once for the combined books, until reversed
    again = admin_client.post(f"{F}/opening", json={"lines": lines})
    assert again.status_code == 409 and again.json()["code"] == "opening_exists"
    # a company's own: once each
    r = admin_client.post(
        f"{F}/opening",
        json={"company_id": companies["a"]["id"], "lines": [{"account_id": acc["2110"], "credit": "50"}]},
    )
    assert r.status_code == 201 and r.json()["difference"] == "-50.000"
    assert [(ln["account"]["code"], ln["debit"]) for ln in r.json()["entry"]["lines"]][-1] == ("3900", "50.000")
    r = admin_client.post(
        f"{F}/opening", json={"company_id": companies["a"]["id"], "lines": [{"account_id": acc["2110"], "credit": "1"}]}
    )
    assert r.status_code == 409
    admin_client.post(f"{F}/entries/approve", json={"ids": [entry["id"]]})
    # in the trial balance as opening (before the period), summing to zero
    tb = balances(admin_client)
    assert tb["1120"] == "1500.000" and tb["3900"] == "-1400.000"
    rows = admin_client.get(f"{F}/trial-balance", params=period()).json()
    assert next(x for x in rows if x["account"]["code"] == "1120")["opening"] == "1500.000"
    admin_client.post(f"{F}/entries/{entry['id']}/reverse", json={"reason": "wrong figures"})
    assert admin_client.post(f"{F}/opening", json={"lines": lines}).status_code == 201


# ------------------------------------------------------------------ the fiscal year


def test_the_fiscal_year_start(admin_client):
    settings(admin_client, "finance", fiscal_year_start_month=4)
    start = date.fromisoformat(admin_client.get(f"{F}/config").json()["fiscal_year_start"])
    t = today()
    assert start.month == 4 and start.day == 1 and start <= t < start.replace(year=start.year + 1)


# ------------------------------------------------------------------ the old system's books


def workbook(opening=(), entries=()) -> bytes:
    from app.modules.finance import oldbooks

    wb = openpyxl.load_workbook(io.BytesIO(oldbooks.template()))
    for row in opening:
        wb[oldbooks.OPENING].append(list(row))
    for row in entries:
        wb[oldbooks.ENTRIES].append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send(client, data: bytes, *, apply=False, **params):
    return client.post(
        f"{F}/import",
        params={"apply": str(apply).lower()} | params,
        files={"file": ("old.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_the_template_has_both_sheets_and_an_example(admin_client):
    r = admin_client.get(f"{F}/import/template")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["أرصدة افتتاحية", "قيود"]
    rows = list(wb["قيود"].iter_rows(values_only=True))
    assert rows[0] == ("رقم القيد في النظام القديم", "التاريخ", "رمز الحساب", "مدين", "دائن", "البيان", "الرقم المدني")
    assert rows[1][0] == "مثال"
    # the empty template checks clean: the example rows are skipped
    r = send(admin_client, r.content)
    assert r.status_code == 200 and r.json()["errors"] == [] and r.json()["entries"] == 0


def test_importing_old_books_checks_then_imports_once(admin_client, company):
    start = today().replace(day=1)
    settings(admin_client, "finance", books_start_date=str(start))
    emp = make_employee(admin_client, company["id"], civil_id="290010100011")
    day = start.strftime("%d/%m/%Y")
    bad = workbook(
        opening=[("1120", "البنك", 1000, None, None), ("9999", "", 5, None, None), ("2110", "", 10, 10, None)],
        entries=[
            ("A-1", day, "6130", 100, None, "Rent", "290010100011"),
            ("A-1", day, "1120", None, 90, None, None),
            ("A-2", "31/02/2020", "6130", 5, None, None, "111111111111"),
            ("A-2", day, "1110", None, 5, None, None),
            ("A-3", str(start - timedelta(days=1)), "6130", 5, None, None, None),
            ("A-3", start - timedelta(days=1), "1110", None, 5, None, None),
        ],
    )
    r = send(admin_client, bad)
    assert r.status_code == 200, r.text
    out = r.json()
    codes = sorted((e["sheet"], e["row"], e["code"]) for e in out["errors"])
    assert codes == [
        ("أرصدة افتتاحية", 4, "unknown_account"),
        ("أرصدة افتتاحية", 5, "line_one_side"),
        ("قيود", 3, "entry_not_balanced"),
        ("قيود", 5, "date_invalid"),
        ("قيود", 7, "before_books_start"),
        ("قيود", 8, "before_books_start"),
    ]
    unbalanced = next(e for e in out["errors"] if e["code"] == "entry_not_balanced")
    assert unbalanced["params"] == {"number": "A-1", "difference": "10.000"}
    assert [(w["code"], w["params"]) for w in out["warnings"]] == [("civil_id_not_found", {"civil_id": "111111111111"})]
    # applying a file with errors imports nothing
    r = send(admin_client, bad, apply=True)
    assert r.json()["applied"] is False and admin_client.get(f"{F}/entries", params=period()).json() == []

    good = workbook(
        opening=[("1120", "البنك", 1000, None, None), ("2110", "", None, 300, None)],
        entries=[
            ("A-1", day, "6130", 100, None, "Rent", "290010100011"),
            ("A-1", day, "1120", None, 100, None, None),
            (7, start, "6140", 12.5, None, None, None),
            (7, start, "1110", None, 12.5, None, None),
        ],
    )
    r = send(admin_client, good)
    assert r.json()["errors"] == [] and r.json()["warnings"] == []
    assert r.json()["opening"] == {"debit": "1000.000", "credit": "300.000", "difference": "700.000", "lines": 2}
    assert (r.json()["entries"], r.json()["lines"]) == (2, 4)
    r = send(admin_client, good, apply=True)
    assert r.status_code == 200 and r.json()["applied"] is True and len(r.json()["numbers"]) == 3
    made = admin_client.get(
        f"{F}/entries", params={"date_from": str(start - timedelta(days=1)), "date_to": str(today())}
    )
    refs = sorted((e["source_kind"], e["source_ref"][:5], e["status"]) for e in made.json())
    assert refs == [("manual", "OLD-7", "draft"), ("manual", "OLD-A", "draft"), ("opening", "OPEN-", "draft")]
    rent = next(e for e in made.json() if e["source_ref"] == "OLD-A-1")
    detail = admin_client.get(f"{F}/entries/{rent['id']}").json()
    assert detail["description"] == "Rent" and detail["lines"][0]["employee"]["id"] == emp["id"]
    other = next(e for e in made.json() if e["source_ref"] == "OLD-7")
    assert other["description"] == "قيد مستورد #7" and other["amount"] == "12.500"
    opening = next(e for e in made.json() if e["source_kind"] == "opening")
    lines = admin_client.get(f"{F}/entries/{opening['id']}").json()["lines"]
    assert lines[-1]["account"]["code"] == "3900" and lines[-1]["credit"] == "700.000"

    # again: refused, nothing made twice
    again = send(admin_client, good, apply=True)
    assert again.status_code == 409 and again.json()["code"] == "old_entries_imported"
    assert again.json()["params"] == {"numbers": "7, A-1"}
    preview = send(admin_client, good).json()
    assert {e["code"] for e in preview["errors"]} == {"old_entry_imported", "opening_exists"}


def test_an_import_follows_the_approval_setting_and_refuses_a_bad_file(admin_client):
    settings(admin_client, "finance", entry_approval="auto")
    data = workbook(entries=[("9", today(), "6130", 3, None, "x", None), ("9", today(), "1110", None, 3, None, None)])
    r = send(admin_client, data, apply=True)
    assert r.json()["applied"] is True
    [e] = admin_client.get(f"{F}/entries", params=period()).json()
    assert e["status"] == "approved" and e["source_ref"] == "OLD-9"
    r = send(admin_client, b"not a workbook")
    assert r.status_code == 422 and r.json()["code"] == "import_not_xlsx"
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    r = send(admin_client, buf.getvalue())
    assert r.status_code == 422 and r.json()["code"] == "import_bad_template"


# ------------------------------------------------------------------ the treasury's deposit rule


def alerts(client, kind: str) -> list[dict]:
    return [a for a in client.get("/api/v1/alerts", headers={"Accept-Language": "en"}).json() if a["kind"] == kind]


def test_the_treasury_limit_alerts_until_a_deposit_brings_it_under(admin_client, company):
    settings(admin_client, "cash", treasury_deposit_limit="100.000")
    d = make_driver(admin_client, company["id"])
    r = admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "60"})
    assert r.status_code == 201, r.text
    assert alerts(admin_client, "treasury_deposit_due") == []
    branch = r.json()["branch_id"]
    admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "40"})
    [a] = alerts(admin_client, "treasury_deposit_due")
    assert a["entity_type"] == "branch" and "100.000" in a["message"] and a["params"]["limit"] == "100.000"
    admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "5"})
    [a] = alerts(admin_client, "treasury_deposit_due")  # one alert, its balance refreshed
    assert a["params"]["balance"] == "105.000"
    r = admin_client.post(
        "/api/v1/cash/bank-deposits",
        json={"branch_id": branch, "amount": "50", "reference": "NBK-1", "receipt_sha256": upload(admin_client)},
    )
    assert r.status_code == 201, r.text
    assert alerts(admin_client, "treasury_deposit_due") == []
    # nightly: raised again when the limit is lowered
    settings(admin_client, "cash", treasury_deposit_limit="50.000")
    from app.modules.cash import tasks

    tasks.scan_treasury("night")
    assert len(alerts(admin_client, "treasury_deposit_due")) == 1


def test_the_deposit_day_alert_in_the_morning_closes_at_the_end_of_the_day(admin_client, company):
    from app.modules.cash import service as cash
    from app.modules.cash import tasks

    assert [cash.kuwait_weekday(date(2026, 10, d)) for d in (3, 9)] == [0, 6]  # a Saturday, a Friday
    d = make_driver(admin_client, company["id"])
    admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "30"})
    others = [x for x in range(7) if x != cash.kuwait_weekday(today())]
    settings(admin_client, "cash", treasury_deposit_weekdays=others)
    assert tasks.scan_treasury("morning") == 0
    settings(admin_client, "cash", treasury_deposit_weekdays=[cash.kuwait_weekday(today())])
    assert tasks.scan_treasury("morning") == 1
    [a] = alerts(admin_client, "treasury_deposit_day")
    assert a["params"]["balance"] == "30.000" and a["entity_type"] == "branch"
    assert tasks.scan_treasury("morning") == 1 and len(alerts(admin_client, "treasury_deposit_day")) == 1
    tasks.scan_treasury("night")
    assert alerts(admin_client, "treasury_deposit_day") == []
    # an empty treasury raises none; one emptied by a deposit closes it
    tasks.scan_treasury("morning")
    branch = admin_client.get("/api/v1/cash/treasury").json()[0]["branch"]["id"]
    admin_client.post(
        "/api/v1/cash/bank-deposits",
        json={"branch_id": branch, "amount": "30", "reference": "NBK-2", "receipt_sha256": upload(admin_client)},
    )
    assert alerts(admin_client, "treasury_deposit_day") == []
    assert tasks.scan_treasury("morning") == 0


@pytest.mark.parametrize("days", [[7], [1, 1, 2]])
def test_deposit_weekdays_are_saturday_to_friday(admin_client, days):
    cur = admin_client.get("/api/v1/settings").json()["cash"]
    r = admin_client.put(
        "/api/v1/settings/cash",
        json={"version": cur["version"], "value": cur["value"] | {"treasury_deposit_weekdays": days}},
    )
    if days == [7]:
        assert r.status_code == 422
    else:
        assert r.json()["value"]["treasury_deposit_weekdays"] == [1, 2]


def test_review_fixes_manual_entries_are_never_stale_and_an_opening_reverses_on_its_date(admin_client):
    from tests.test_finance import period

    accounts = {a["code"]: a["id"] for a in admin_client.get("/api/v1/finance/accounts").json()}
    first = period()["date_from"]
    r = admin_client.post(
        "/api/v1/finance/entries/manual",
        json={
            "entry_date": first,
            "description": "rent accrual",
            "lines": [
                {"account_id": accounts["6130"], "debit": "5.000", "credit": "0"},
                {"account_id": accounts["2140"], "debit": "0", "credit": "5.000"},
            ],
        },
    )
    assert r.status_code == 201, r.text
    out = admin_client.post("/api/v1/finance/entries/post", json=period()).json()
    assert out["stale"] == [], out  # written by hand: no document to match


def test_review_fixes_the_books_start_holds_once_an_opening_stands_and_its_reversal_takes_its_date(admin_client):
    acc = accounts(admin_client)
    start = today().replace(day=1)
    settings(admin_client, "finance", books_start_date=str(start))
    out = admin_client.post(f"{F}/opening", json={"lines": [{"account_id": acc["1120"], "debit": "1000"}]}).json()
    entry = out["entry"]
    if entry["status"] == "draft":
        admin_client.post(f"{F}/entries/approve", json={"ids": [entry["id"]]})
    cur = admin_client.get("/api/v1/settings").json()["finance"]
    moved = str(start - timedelta(days=40))
    r = admin_client.put(
        "/api/v1/settings/finance",
        json={"version": cur["version"], "value": cur["value"] | {"books_start_date": moved}},
    )
    assert r.status_code == 409 and r.json()["code"] == "books_start_locked", r.text
    # a wrong opening is reversed where it stands, so the corrected one replaces it from the start
    r = admin_client.post(f"{F}/entries/{entry['id']}/reverse", json={"reason": "bank was 800"})
    assert r.status_code == 200, r.text
    assert r.json()["entry_date"] == str(start - timedelta(days=1))
    admin_client.post(f"{F}/opening", json={"lines": [{"account_id": acc["1120"], "debit": "800"}]})
    rows = admin_client.get(
        f"{F}/entries", params={"date_from": str(start - timedelta(days=1)), "date_to": str(start - timedelta(days=1))}
    ).json()
    assert sum(1 for e in rows if e["source_kind"] in ("opening", "reversal")) == 3
