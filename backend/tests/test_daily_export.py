"""The daily reports page's Excel export: the list as filtered (status, day, driver), in the user's language, with the
shared print header, numbers as numbers; only with daily_reports.view and reports.export, and only the companies in
the user's scope."""

import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.core.clock import today
from tests.conftest import bearer, bind_device, login, make_driver, make_user
from tests.test_cash import approve, report

X = "/api/v1/daily-reports/export"


@pytest.fixture
def two(admin_client, client, new_client, company):
    a = make_driver(admin_client, company["id"])
    a |= {"h": bearer(bind_device(client, a["phone"]))}
    other = new_client()
    b = make_driver(admin_client, company["id"])
    b |= {"h": bearer(bind_device(other, b["phone"])), "client": other}
    return a, b


def sheet(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(r.content)).active


def test_the_filtered_list_as_excel_with_the_print_header(admin_client, client, two):
    a, b = two
    ra = report(client, a, cash="12.500", orders=31).json()["id"]
    report(b["client"], b, cash="4", orders=7, day=today() - timedelta(days=1))
    assert approve(admin_client, ra).status_code == 200

    ws = sheet(admin_client.get(X, params={"status": "approved"}, headers={"Accept-Language": "en"}))
    header = [c.value for c in ws[1]]
    assert header[:7] == ["Driver", "Vehicle", "Business day", "Session", "Orders", "Day per platform", "Cash"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    row = dict(zip(header, rows[0], strict=True))
    assert row["Orders"] == 31 and float(row["Cash"]) == 12.5 and row["Status"] == "Approved"
    assert row["Business day"] == str(today())
    # the shared print header: company and title on top, the filter, who printed it
    assert "Daily reports" in ws.oddHeader.center.text
    assert "Approved" in ws.oddHeader.right.text
    assert "&P" in ws.oddFooter.center.text and ws.oddFooter.right.text

    # in Arabic: right to left, Arabic headers; the day filter narrows it
    ws = sheet(admin_client.get(X, params={"status": "submitted", "business_date": str(today() - timedelta(days=1))}))
    assert ws.sheet_view.rightToLeft is True
    assert [c.value for c in ws[1]][0] == "السائق"
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][4] == 7

    # no status: every report; one driver only
    assert ws.max_row == 2
    ws = sheet(admin_client.get(X, params={"status": "", "driver_id": a["id"]}))
    assert len(list(ws.iter_rows(min_row=2, values_only=True))) == 1
    ws = sheet(admin_client.get(X, params={"status": ""}))
    assert len(list(ws.iter_rows(min_row=2, values_only=True))) == 2


def test_needs_the_export_permission_and_keeps_the_company_scope(admin_client, client, new_client, two, companies):
    a, _ = two
    report(client, a, cash="3")
    make_user(admin_client, "viewer", permissions=["daily_reports.view"])
    viewer = new_client()
    login(viewer, "viewer")
    r = viewer.get(X)
    assert r.status_code == 403 and r.json()["params"] == {"permission": "reports.export"}

    make_user(
        admin_client,
        "elsewhere",
        permissions=["daily_reports.view", "reports.export"],
        company_ids=[companies["b"]["id"]],
    )
    elsewhere = new_client()
    login(elsewhere, "elsewhere")
    ws = sheet(elsewhere.get(X, params={"status": ""}))
    assert ws.max_row == 1  # the header only: the report is another company's
