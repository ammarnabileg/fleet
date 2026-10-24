"""The driver's languages: French, Swahili, Hindi and Filipino beside Arabic and English, translated from the dashboard
(drafts, then published versions the app downloads), and the server's own texts in the driver's language."""

import hashlib
import io
import json
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import text

from app.modules.i18n import icu, store
from app.modules.i18n import service as i18n
from tests.conftest import bearer, bind_device, login, make_driver, make_user

ROOT = Path(__file__).resolve().parents[2]
L10N = ROOT / "driver_app/lib/l10n"
NEW = ("fr", "sw", "hi", "fil")
ALL = ("ar", "en", *NEW)
T = "/api/v1/i18n/translations"


def arb(lang: str) -> dict:
    return {k: v for k, v in json.loads((L10N / f"app_{lang}.arb").read_text()).items() if not k.startswith("@")}


# ------------------------------------------------------------------ the shipped texts


def test_every_language_has_every_app_key_and_valid_messages():
    en = arb("en")
    for lang in ALL:
        strings = arb(lang)
        assert set(strings) == set(en), lang  # the same 531 keys in all six arb files
        for key, value in strings.items():
            assert icu.check(value, en[key]) is None, (lang, key, value)


def test_the_seed_is_the_apps_arb_files_and_translates_every_driver_text():
    for lang in ("ar", "en"):
        assert store.seed(lang)["app"] == arb(lang)  # tool/l10n_sync.py keeps them in step
    for lang in NEW:
        assert store.seed(lang)["app"] == arb(lang)
        for scope in store.SCOPES:
            base = store.base_texts(scope, lang)
            assert set(base) == set(store.keys(scope)), (lang, scope)
            for key, value in base.items():
                assert icu.check(value, store.reference(scope, key)) is None, (lang, scope, key)


def test_status_template_seed_matches_the_built_in_templates():
    from app.modules.driver_status.templates import DEFAULTS

    for lang in ("ar", "en"):
        for code, t in DEFAULTS.items():
            for field in ("name", "title", "subtitle"):
                assert store.seed(lang)["status_templates"][f"{code}.{field}"] == t[field][lang]


@pytest.mark.parametrize(
    ("value", "reference", "problem"),
    [
        ("Bonjour {name}", "Hello {name}", None),
        ("{n, plural, one{# commande} other{# commandes}}", "{n} orders", None),
        ("{n, plural, =0{aucune} one{une} other{#}}", "{n, plural, one{one} other{#}}", None),
        ("{g, select, male{il} female{elle} other{iel}}", "{g, select, male{he} other{they}}", None),
        ("Bonjour {nom}", "Hello {name}", "missing:{name}"),
        ("Bonjour {name} {x}", "Hello {name}", "unknown:{x}"),
        ("Bonjour {name", "Hello {name}", "invalid:expected_comma@13"),
        ("Bonjour }", "Hello", "invalid:unexpected_close@8"),
        ("{n, plural, one{x}}", "{n}", "invalid:missing_other@19"),
        ("{n, plural, single{x} other{y}}", "{n}", "invalid:bad_selector@18"),
        ("{n, select, a{x} other{y}}", "{n, plural, one{a} other{b}}", "kind:{n}"),
        ("   ", "Hello", "empty"),
    ],
)
def test_icu_validation(value, reference, problem):
    assert icu.check(value, reference) == problem


# ------------------------------------------------------------------ the app's side


def test_languages_and_translations_before_sign_in(client):
    langs = client.get("/api/v1/driver/languages").json()
    assert [(x["code"], x["dir"]) for x in langs] == [
        ("ar", "rtl"), ("en", "ltr"), ("fr", "ltr"), ("sw", "ltr"), ("hi", "ltr"), ("fil", "ltr")
    ]  # fmt: skip
    assert next(x for x in langs if x["code"] == "hi")["name_native"] == "हिन्दी"
    r = client.get("/api/v1/driver/translations", params={"lang": "fr"})
    body = r.json()
    assert r.status_code == 200 and body["version"] == 0 and body["dir"] == "ltr"
    assert body["strings"]["save"] == "Enregistrer" and len(body["strings"]) == len(arb("en"))
    assert client.get("/api/v1/driver/translations", params={"lang": "fr", "version": 0}).json() == {
        "lang": "fr", "version": 0, "dir": "ltr", "unchanged": True
    }  # fmt: skip
    assert client.get("/api/v1/driver/translations", params={"lang": "xx"}).status_code == 404
    assert "Accept-Language" in r.headers["vary"]


def test_publish_makes_a_version_drafts_stay_unseen(admin_client, client):
    m = client.get("/api/v1/driver/translations/manifest")
    etag = m.headers["etag"]
    fr = next(x for x in m.json()["languages"] if x["code"] == "fr")
    assert fr["version"] == 0 and fr["sha256"] is None and m.json()["schema"] == 1
    assert client.get("/api/v1/driver/translations/manifest", headers={"If-None-Match": etag}).status_code == 304

    r = admin_client.put(f"{T}/fr", json={"items": [{"scope": "app", "key": "save", "text": "Sauvegarder"}]})
    assert r.status_code == 200 and r.json()["version"] == 0  # a draft: nothing published
    assert client.get("/api/v1/driver/translations", params={"lang": "fr"}).json()["strings"]["save"] == "Enregistrer"
    assert client.get("/api/v1/driver/translations/manifest", headers={"If-None-Match": etag}).status_code == 304

    r = admin_client.post(f"{T}/fr/publish")
    assert r.status_code == 200 and r.json()["version"] == 1
    m2 = client.get("/api/v1/driver/translations/manifest", headers={"If-None-Match": etag})
    assert m2.status_code == 200
    fr = next(x for x in m2.json()["languages"] if x["code"] == "fr")
    assert fr["version"] == 1
    raw = client.get("/api/v1/driver/translations/bundle/fr/1")
    assert hashlib.sha256(raw.content).hexdigest() == fr["sha256"]
    doc = json.loads(raw.content)
    assert doc["lang"] == "fr" and doc["version"] == 1 and doc["schema"] == 1 and doc["min_app_build"] == 1
    assert doc["strings"]["save"] == "Sauvegarder"
    assert client.get("/api/v1/driver/translations", params={"lang": "fr", "version": 0}).json()["version"] == 1
    assert client.get("/api/v1/driver/translations", params={"lang": "fr", "version": 1}).json()["unchanged"] is True
    # an edit after publishing changes nothing published until the next publish; versions never change
    admin_client.put(f"{T}/fr", json={"items": [{"scope": "app", "key": "save", "text": "Garder"}]})
    assert json.loads(client.get("/api/v1/driver/translations/bundle/fr/1").content)["strings"]["save"] == "Sauvegarder"
    assert admin_client.post(f"{T}/fr/publish").json()["version"] == 2
    langs = {x["code"]: x for x in admin_client.get(f"{T}/languages").json()}
    assert langs["fr"]["published_version"] == 2 and langs["fr"]["unpublished"] == 0


def test_a_broken_message_is_refused_on_save_and_import(admin_client):
    bad = {"items": [{"scope": "app", "key": "greeting", "text": "Bonjour {nom"}]}
    r = admin_client.put(f"{T}/fr", json=bad)
    assert r.status_code == 422 and r.json()["code"] == "invalid_translation"
    r = admin_client.put(f"{T}/fr", json={"items": [{"scope": "app", "key": "greeting", "text": "Bonjour"}]})
    assert r.status_code == 422 and r.json()["params"]["reason"] == "missing:{name}"
    r = admin_client.put(f"{T}/fr", json={"items": [{"scope": "app", "key": "nope", "text": "x"}]})
    assert r.status_code == 422 and r.json()["code"] == "unknown_translation_key"


def test_review_status_and_progress(admin_client, db):
    rows = admin_client.get(f"{T}/sw", params={"scope": "app", "q": "Hifadhi"}).json()
    save = next(r for r in rows if r["key"] == "save")
    assert save["status"] == "needs_review" and save["en"] == "Save" and save["ar"] == "حفظ"
    before = next(x for x in admin_client.get(f"{T}/languages").json() if x["code"] == "sw")["progress"]
    assert before["reviewed"] == 0 and before["needs_review"] == before["total"]
    r = admin_client.post(
        f"{T}/sw/review", json={"items": [{"scope": "app", "key": "save"}, {"scope": "app", "key": "back"}]}
    )
    assert r.json()["reviewed"] == 2
    after = next(x for x in admin_client.get(f"{T}/languages").json() if x["code"] == "sw")["progress"]
    assert after["reviewed"] == 2 and after["percent"] == round(200 / after["total"])
    assert admin_client.get(f"{T}/sw", params={"status": "reviewed"}).json()[0]["key"] in ("save", "back")
    ar = next(x for x in admin_client.get(f"{T}/languages").json() if x["code"] == "ar")
    assert ar["progress"]["percent"] == 100 and ar["required"]
    actions = db.execute(text("SELECT action FROM audit.events WHERE entity_type = 'language'")).scalars().all()
    assert "translations.reviewed" in actions


def test_export_import_round_trip(admin_client):
    r = admin_client.get(f"{T}/hi/export")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["scope", "key", "en", "translation", "status"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "app" and row[1].value == "save":
            row[3].value = "सहेजें"
    buf = io.BytesIO()
    wb.save(buf)
    r = admin_client.post(f"{T}/hi/import", files={"file": ("hi.xlsx", buf.getvalue())}, params={"reviewed": True})
    assert r.status_code == 200 and r.json()["saved"] == 1
    row = next(x for x in admin_client.get(f"{T}/hi", params={"q": "सहेजें"}).json() if x["key"] == "save")
    assert row["text"] == "सहेजें" and row["status"] == "reviewed" and row["edited"]
    # a broken message anywhere: nothing saved, the rows named
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "app" and row[1].value == "greeting":
            row[3].value = "नमस्ते {name"
        if row[0].value == "app" and row[1].value == "back":
            row[3].value = "पीछे जाएँ"
    buf = io.BytesIO()
    wb.save(buf)
    r = admin_client.post(f"{T}/hi/import", files={"file": ("hi.xlsx", buf.getvalue())})
    assert r.status_code == 422 and r.json()["code"] == "translation_import_invalid"
    assert r.json()["params"]["rows"][0]["key"] == "greeting"
    assert (
        next(x for x in admin_client.get(f"{T}/hi", params={"scope": "app"}).json() if x["key"] == "back")["text"]
        == "वापस"
    )
    r = admin_client.post(f"{T}/hi/import", files={"file": ("x.xlsx", b"not a workbook")})
    assert r.json()["code"] == "import_not_xlsx"


def test_languages_on_and_off(admin_client, client):
    langs = {x["code"]: x for x in admin_client.get(f"{T}/languages").json()}
    r = admin_client.patch(f"{T}/languages/en", json={"version": langs["en"]["version"], "is_active": False})
    assert r.status_code == 422 and r.json()["code"] == "language_required"
    r = admin_client.patch(f"{T}/languages/ar", json={"version": langs["ar"]["version"], "is_active": False})
    assert r.status_code == 422
    r = admin_client.patch(f"{T}/languages/sw", json={"version": langs["sw"]["version"], "is_active": False})
    assert r.status_code == 200 and not r.json()["is_active"]
    assert "sw" not in [x["code"] for x in client.get("/api/v1/driver/languages").json()]
    assert client.get("/api/v1/driver/translations", params={"lang": "sw"}).status_code == 404
    assert "sw" not in [x["code"] for x in client.get("/api/v1/driver/translations/manifest").json()["languages"]]


def test_view_and_edit_permissions(admin_client, new_client):
    make_user(admin_client, "viewer", permissions=["settings.view"])
    c = new_client()
    login(c, "viewer")
    assert c.get(f"{T}/languages").status_code == 200
    assert c.get(f"{T}/fr").status_code == 200
    assert c.put(f"{T}/fr", json={"items": [{"scope": "app", "key": "save", "text": "x"}]}).status_code == 403
    assert c.post(f"{T}/fr/publish").status_code == 403
    make_user(admin_client, "nobody")
    c2 = new_client()
    login(c2, "nobody")
    assert c2.get(f"{T}/languages").status_code == 403


# ------------------------------------------------------------------ the server's texts in the driver's language


def test_server_texts_in_the_new_languages_with_fallback(admin_client, client, db):
    fr = client.get("/api/v1/i18n/catalog/fr", params={"ns": "errors"}).json()["messages"]["errors"]
    assert fr["otp_invalid"] == "Le code de vérification est faux ou expiré"
    common = client.get("/api/v1/i18n/catalog/fr", params={"ns": "common"}).json()["messages"]["common"]
    assert common["save"] == "Save"  # not a driver text: English, the language's fallback
    assert i18n.t(db, "hi", "errors.no_open_custody") == "आपके पास कोई गाड़ी नहीं है"
    assert i18n.t(db, "sw", "messages.otp", code="123456", minutes=5).startswith("Msimbo wako wa kuingia: 123456.")
    # a dashboard edit of a server text reaches the drivers once published
    admin_client.put(
        f"{T}/fr", json={"items": [{"scope": "server", "key": "errors.otp_locked", "text": "Trop d'essais"}]}
    )
    assert i18n.t(db, "fr", "errors.otp_locked") != "Trop d'essais"
    admin_client.post(f"{T}/fr/publish")
    db.expire_all()
    assert i18n.t(db, "fr", "errors.otp_locked") == "Trop d'essais"
    # names: the reader's language, else English (not Arabic) for the new languages; Arabic as before
    assert i18n.pick({"ar": "شركة", "en": "Company"}, "fr", "ar") == "Company"
    assert i18n.pick({"ar": "شركة", "en": "Company"}, "ar", "ar") == "شركة"
    assert (
        i18n.driver_lang(db, "fil") == "fil"
        and i18n.driver_lang(db, "de") == "en"
        and i18n.driver_lang(db, None) == "ar"
    )


def test_errors_carry_a_key_params_and_a_message_in_the_language(client):
    r = client.get("/api/v1/driver/translations", params={"lang": "xx"}, headers={"Accept-Language": "fr"})
    body = r.json()
    assert body["code"] == "unknown_language" and body["message_key"] == "errors.unknown_language"
    assert body["params"] == {"code": "xx"} and body["message"] == "Langue inconnue ou désactivée : xx"
    r = client.get("/api/v1/driver/translations", params={"lang": "xx"}, headers={"Accept-Language": "de-DE"})
    assert r.json()["message"] == "Unknown or inactive language: xx"  # unknown language: English


def test_the_drivers_language_for_notices_is_set_explicitly(admin_client, client, company, db):
    d = make_driver(admin_client, company["id"])
    h = bearer(bind_device(client, d["phone"]))
    r = client.put("/api/v1/driver/preferred-language", json={"lang": "hi"}, headers=h)
    assert r.status_code == 200 and r.json() == {"preferred_language": "hi"}
    assert db.execute(text("SELECT push_lang FROM identity.devices WHERE revoked_at IS NULL")).scalar() == "hi"
    r = client.put("/api/v1/driver/preferred-language", json={"lang": "xx"}, headers=h)
    assert r.status_code == 422
    # notices listed in the app's language
    h_fr = h | {"Accept-Language": "fr"}
    assert client.get("/api/v1/driver/notifications", headers=h_fr).status_code == 200
    # signed in, the translations are not counted against the address
    assert client.get("/api/v1/driver/translations", params={"lang": "fr"}, headers=h).status_code == 200
    assert (
        client.get(
            "/api/v1/driver/translations", params={"lang": "fr"}, headers={"Authorization": "Bearer x"}
        ).status_code
        == 401
    )


def test_live_status_cards_in_french(admin_client, client, company):
    d = make_driver(admin_client, company["id"])
    h = bearer(bind_device(client, d["phone"])) | {"Accept-Language": "fr"}
    body = client.get("/api/v1/driver/live-status", headers=h).json()
    templates = {t["code"]: t for t in body["config"]["templates"]}
    assert templates["shift_off"]["title"]["fr"] == "Hors service"
    assert templates["shift_off"]["title"]["ar"] == "خارج الدوام"
    assert templates["location_problem"]["subtitle"]["hi"] == "आख़िरी अपडेट {last_update}"
