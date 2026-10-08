"""Initial data import from the onboarding workbook (the contract's data template): check first, then import all at
once; importing the same file again updates and never duplicates."""

import io
from datetime import UTC, date, datetime, timedelta

import openpyxl
import pytest
from sqlalchemy import text

from app.modules.imports.api import NEEDED
from app.modules.imports.workbook import (
    DRIVER_COLUMNS,
    INSTRUCTIONS,
    OPENING,
    OPENING_COLUMNS,
    PEOPLE,
    PEOPLE_COLUMNS,
    VEHICLE_COLUMNS,
    VEHICLES,
    parse_days,
    parse_password,
    parse_yes_no,
)
from tests.conftest import login, make_user

COMPANY = "شركة أ"  # the "companies" fixture names


def workbook(vehicles=(), people=(), opening=(), *, break_header=False, driver_columns=False) -> bytes:
    """driver_columns: the people sheet has the optional app, password and days columns after the last one."""
    wb = openpyxl.Workbook()
    wb.active.title = "التعليمات"
    for title, columns, rows in (
        (VEHICLES, VEHICLE_COLUMNS, vehicles),
        (PEOPLE, PEOPLE_COLUMNS + (DRIVER_COLUMNS if driver_columns else []), people),
        (OPENING, OPENING_COLUMNS, opening),
    ):
        ws = wb.create_sheet(title)
        ws.append(columns[:-1] + ["something else"] if break_header and title == PEOPLE else columns)
        ws.append(["مثال"] + ["x"] * (len(columns) - 1))  # the grey example row
        for i, row in enumerate(rows, start=1):
            ws.append([i, *row])
        ws.append([len(rows) + 1])  # numbered but empty: ignored
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def vehicle(plate="12/34567", vin="MR0AA0000BB000001", **kw):
    row = {
        "plate": plate,
        "make": "تويوتا",
        "model": "يارس",
        "year": 2023,
        "color": "أبيض",
        "vin": vin,
        "company": COMPANY,
        "branch": "الفرع الرئيسي",
        "insurance": "15/03/2027",
        "registration": date(2027, 6, 20),
    }
    row.update(kw)
    return list(row.values())


def person(
    name="Ahmed Ali", role="سائق", civil="290010112345", phone="98765432", app=None, password=None, days=None, **kw
):
    row = {
        "name": name,
        "role": role,
        "civil": civil,
        "phone": phone,
        "company": COMPANY,
        "branch": "الفرع الرئيسي",
        "dept": "العمليات",
        "job": "سائق توصيل",
        "status": "على رأس العمل",
        "salary": "150.000",
        "residence": "30/09/2027",
        "model": "Samsung A15",
    }
    row.update(kw)
    driver_columns = [app, password, days]
    return list(row.values()) + (driver_columns if any(v is not None for v in driver_columns) else [])


def post(client, data, apply=False):
    return client.post(
        "/api/v1/imports/workbook",
        params={"apply": apply},
        files={"file": ("data.xlsx", data, "application/octet-stream")},
    )


@pytest.fixture
def main_branch(admin_client):
    branch = next(b for b in admin_client.get("/api/v1/branches").json() if b["is_default"])
    return branch["name"]["ar"]


def test_check_then_import_then_reimport(admin_client, companies, main_branch, db):
    data = workbook(
        [vehicle(branch=main_branch), vehicle("12/34568", "MR0AA0000BB000002", branch=main_branch)],
        [
            person(branch=main_branch),
            person("سالم", role="مشرف", civil="290010112346", phone="98765433", branch=main_branch),
        ],
    )
    check = post(admin_client, data).json()
    assert check["applied"] is False and check["errors"] == []
    assert check["vehicles"] == {"created": 2, "updated": 0} and check["people"] == {"created": 2, "updated": 0}
    assert check["documents"] == 6  # insurance + registration per vehicle, residence per person
    assert admin_client.get("/api/v1/vehicles").json() == []  # a check changes nothing

    done = post(admin_client, data, apply=True).json()
    assert done["applied"] is True and done["vehicles"]["created"] == 2
    vehicles = admin_client.get("/api/v1/vehicles").json()
    assert sorted(v["plate_number"] for v in vehicles) == ["12/34567", "12/34568"]
    driver = admin_client.get("/api/v1/employees", params={"is_driver": True}).json()[0]
    assert (driver["phone"], driver["civil_id"], driver["app_access"]) == ("+96598765432", "290010112345", "active")
    assert driver["name"] == {"ar": "Ahmed Ali", "en": "Ahmed Ali"} and driver["employee_number"] == "E00001"
    assert driver["basic_salary"] == "150.000"
    supervisor = admin_client.get("/api/v1/employees", params={"is_driver": False}).json()[0]
    assert supervisor["app_access"] == "none" and supervisor["employee_number"] == "E00002"
    docs = admin_client.get("/api/v1/documents", params={"owner_type": "vehicle", "owner_id": vehicles[0]["id"]})
    assert sorted((d["type_code"], d["expiry_date"]) for d in docs.json()) == [
        ("insurance", "2027-03-15"),
        ("registration", "2027-06-20"),
    ]

    again = post(
        admin_client,
        workbook(
            [vehicle(branch=main_branch, color="فضي"), vehicle("12/34568", "MR0AA0000BB000002", branch=main_branch)],
            [
                person(branch=main_branch, residence="30/09/2028"),
                person("سالم", role="مشرف", civil="290010112346", phone="98765433", branch=main_branch),
            ],
        ),
        apply=True,
    ).json()
    assert again["vehicles"] == {"created": 0, "updated": 2} and again["people"] == {"created": 0, "updated": 2}
    assert again["documents"] == 1  # only the renewed residence
    assert len(admin_client.get("/api/v1/vehicles").json()) == 2
    assert {v["color"] for v in admin_client.get("/api/v1/vehicles").json()} == {"فضي", "أبيض"}
    assert db.execute(text("SELECT count(*) FROM people.employees")).scalar() == 2
    assert db.execute(text("SELECT count(*) FROM audit.events WHERE action = 'import.applied'")).scalar() == 2


def test_every_row_error_is_reported_and_nothing_is_imported(admin_client, companies, main_branch, db):
    data = workbook(
        [
            vehicle(branch=main_branch),
            vehicle("12 /34567", "MR0AA0000BB000009", branch=main_branch),  # same plate, other spacing
            vehicle("55/55555", "MR0AA0000BB000003", branch="فرع لا يوجد", insurance="2027-15-03"),
        ],
        [
            person(branch=main_branch, phone="1234"),
            person("Ali", civil="12345", phone="98765439", branch=main_branch),
            person("Omar", civil="", phone="98765440", branch=main_branch),  # civil ID only for visa in process
            person("Visa", civil="", phone="98765441", status="تحت إجراء الفيزا", branch=main_branch),
            person(
                "Zed", civil="290010112350", phone="98765442", status="هارب", company="شركة لا توجد", branch=main_branch
            ),
        ],
    )
    r = post(admin_client, data, apply=True).json()
    assert r["applied"] is False
    found = {(e["sheet"], e["row"], e["code"]) for e in r["errors"]}
    assert found == {
        (VEHICLES, 3, "duplicate_in_file"),
        (VEHICLES, 4, "duplicate_in_file"),
        (VEHICLES, 5, "unknown_value"),
        (PEOPLE, 3, "invalid_phone"),
        (PEOPLE, 4, "invalid_civil_id"),
        (PEOPLE, 5, "field_required"),
        (PEOPLE, 7, "unknown_value"),
    }
    unknown_branch = next(e for e in r["errors"] if e["sheet"] == VEHICLES and e["row"] == 5)
    assert unknown_branch["params"] == {"field": "الفرع", "value": "فرع لا يوجد"}
    assert r["people"]["created"] == 1  # the visa row is fine on its own
    assert db.execute(text("SELECT count(*) FROM people.employees")).scalar() == 0
    assert db.execute(text("SELECT count(*) FROM fleet.vehicles")).scalar() == 0


def test_conflicts_with_existing_records_are_row_errors(admin_client, companies, main_branch):
    admin_client.post(
        "/api/v1/employees",
        json={
            "employee_number": "X1",
            "name": {"ar": "قديم"},
            "company_id": companies["a"]["id"],
            "civil_id": "290010112399",
        },
    )
    r = post(
        admin_client,
        workbook(
            [vehicle(branch=main_branch)],
            [
                person(branch=main_branch, phone="98765450", civil="290010112398"),
                person("Other", civil="290010112397", phone="98765450", branch=main_branch),
            ],  # same phone in file
        ),
    ).json()
    assert {(e["sheet"], e["code"]) for e in r["errors"]} == {(PEOPLE, "duplicate_in_file")}


def test_the_template_must_be_used_as_it_is(admin_client, companies):
    r = post(admin_client, workbook(break_header=True))
    assert r.status_code == 422 and r.json()["code"] == "import_bad_template" and r.json()["params"]["sheet"] == PEOPLE
    r = post(admin_client, b"PK\x03\x04 not really a zip")
    assert r.status_code == 422 and r.json()["code"] == "import_not_xlsx"


def test_permissions_scope_and_salary(admin_client, new_client, companies, main_branch):
    data = workbook([vehicle(branch=main_branch)], [person(branch=main_branch)])
    make_user(admin_client, "viewer", permissions=["employees.view", "vehicles.view"])
    c = new_client()
    login(c, "viewer")
    assert post(c, data).status_code == 403
    make_user(
        admin_client,
        "importer_b",
        company_ids=[companies["b"]["id"]],
        permissions=["employees.create", "employees.update", "vehicles.create", "vehicles.update", "documents.manage"],
    )
    c = new_client()
    login(c, "importer_b")
    r = post(c, data).json()
    assert {e["code"] for e in r["errors"]} == {"company_out_of_scope"}  # the rows belong to company A
    make_user(
        admin_client,
        "importer_a",
        company_ids=[companies["a"]["id"]],
        permissions=["employees.create", "employees.update", "vehicles.create", "vehicles.update", "documents.manage"],
    )
    c = new_client()
    login(c, "importer_a")
    r = post(c, data, apply=True).json()
    assert r["applied"] and [w["code"] for w in r["warnings"]] == ["salary_skipped"]
    assert admin_client.get("/api/v1/employees").json()[0]["basic_salary"] is None


def test_opening_balances_are_imported_once(admin_client, client, companies, main_branch):

    people = [person(branch=main_branch)]
    opening = [["Ahmed Ali", "290010112345", "35.750", "24/10/2026", "المحاسب"]]
    r = post(admin_client, workbook(people=people, opening=opening), apply=True).json()
    assert r["applied"] and r["opening_balances"] == 1
    driver = admin_client.get("/api/v1/employees", params={"is_driver": True}).json()[0]
    statement = admin_client.get(f"/api/v1/cash/drivers/{driver['id']}/statement").json()
    assert statement["posted"] == "35.750" and statement["lines"][0]["kind"] == "opening"
    assert statement["lines"][0]["reason"] == "approved by المحاسب"
    again = post(admin_client, workbook(people=people, opening=opening), apply=True).json()
    assert again["applied"] and again["opening_balances"] == 0  # same amount: unchanged
    changed = post(admin_client, workbook(people=people, opening=[opening[0][:2] + ["40", "24/10/2026", "x"]])).json()
    assert [e["code"] for e in changed["errors"]] == ["opening_balance_exists"]
    unknown = post(admin_client, workbook(opening=[["X", "290010199999", "1", "24/10/2026", "x"]])).json()
    assert [e["code"] for e in unknown["errors"]] == ["unknown_value"]


def test_opening_balances_need_the_cash_permission(admin_client, new_client, companies, main_branch):
    make_user(
        admin_client,
        "importer",
        permissions=["employees.create", "employees.update", "vehicles.create", "vehicles.update", "documents.manage"],
    )
    c = new_client()
    login(c, "importer")
    r = post(
        c,
        workbook(people=[person(branch=main_branch)], opening=[["Ahmed Ali", "290010112345", "1", "24/10/2026", "x"]]),
    ).json()
    assert [(e["sheet"], e["code"]) for e in r["errors"]] == [(OPENING, "permission_denied")]


# ------------------------------------------------------------------ the driver app, initial password and its days


def test_the_driver_columns_read_what_people_write():
    assert [parse_yes_no(v) for v in ("نعم", "ن", "YES", "y", 1, 1.0, True, "مفعّل", " لا ", "No", 0, False)] == [
        True,
        True,
        True,
        True,
        True,
        True,
        True,
        True,
        False,
        False,
        False,
        False,
    ]
    assert parse_yes_no("غير مفعّل") is False and parse_yes_no("ربما") is None and parse_yes_no(2) is None
    assert parse_password("  pass word 1 ") == "pass word 1"  # spaces inside count, around do not
    assert parse_password(12345678) == "12345678" and parse_password(12345678.0) == "12345678"
    assert parse_password(123456789012345) == "123456789012345"
    for bad in ("short", "x" * 65, 1234567890123456, 1234.5678, True, datetime(2026, 1, 1), date(2026, 1, 1)):
        assert parse_password(bad) is None, bad
    assert [parse_days(v) for v in (14, 14.0, "14", "١٤", "۷", 1, 60)] == [14, 14, 14, 14, 7, 1, 60]
    for bad in (0, 61, 14.5, "x", "-3", True, "١٤ يوم"):
        assert parse_days(bad) is None, bad


def test_template_app_password_and_days(admin_client, client, companies, main_branch, owner_db):
    a, b, c = "290010112361", "290010112362", "290010112363"
    data = workbook(
        people=[
            person("A", civil=a, phone="98765461", branch=main_branch, app="نعم", password="Init-pass-1", days=7),
            person("B", civil=b, phone="98765462", branch=main_branch, app="لا"),
            person("C", civil=c, phone="98765463", branch=main_branch),
            person("سالم", role="مشرف", civil="290010112364", phone="98765464", branch=main_branch),
        ],
        driver_columns=True,
    )
    check = post(admin_client, data)
    out = check.json()
    assert out["errors"] == [] and out["warnings"] == [] and out["applied"] is False, out
    assert (out["activated"], out["claims"]) == (2, 1)  # A and C get the app, A his password
    assert owner_db.execute(text("SELECT count(*) FROM identity.driver_claims")).scalar() == 0  # a check keeps nothing

    done = post(admin_client, data, apply=True)
    assert done.json()["applied"] and (done.json()["activated"], done.json()["claims"]) == (2, 1)
    staff = {e["civil_id"]: e for e in admin_client.get("/api/v1/employees", params={"limit": 50}).json()}
    assert [staff[x]["app_access"] for x in (a, b, c, "290010112364")] == ["active", "none", "active", "none"]
    claim = admin_client.get(f"/api/v1/employees/{staff[a]['id']}/claim").json()
    expires = datetime.fromisoformat(claim["expires_at"])
    assert claim["open"] and abs(expires - (datetime.now(UTC) + timedelta(days=7))) < timedelta(minutes=5)
    assert admin_client.get(f"/api/v1/employees/{staff[c]['id']}/claim").json() is None
    r = client.post(
        "/api/v1/driver/auth/claim", json={"civil_id": a, "password": "Init-pass-1", "device_uid": "phone-a-01"}
    )
    assert r.status_code == 200 and r.json()["next"] == "phone", r.text

    # the same file again: the open password stays as it was, nobody is activated twice
    again = post(admin_client, data, apply=True)
    out = again.json()
    assert out["applied"] and (out["activated"], out["claims"]) == (0, 0)
    assert [(w["row"], w["code"]) for w in out["warnings"]] == [(3, "claim_already_open")]
    audit = owner_db.execute(text("SELECT coalesce(before::text, '') || coalesce(after::text, '') FROM audit.events"))
    logged = " ".join(row[0] for row in audit)
    assert '"activated": 2' in logged and '"claims": 1' in logged
    for response in (check, done, again):
        assert "Init-pass-1" not in response.text
    assert "Init-pass-1" not in logged


def test_bad_app_password_and_days_values_are_row_errors(admin_client, companies, main_branch):
    rows = {
        3: dict(app="ربما"),
        4: dict(password="shortpw"),
        5: dict(days=0),
        6: dict(days=61),
        7: dict(days="x"),
        8: dict(role="مشرف", password="Super-pass-9"),
        9: dict(app="لا", password="Other-pass-1"),
        10: dict(password="290010112310"),  # the civil ID itself
        11: dict(app=True, password=12345678, days="١٤"),  # an Excel TRUE, a number, Arabic-Indic digits: fine
        12: dict(app="مفعّل", password="  with spaces  "),
    }
    data = workbook(
        people=[
            person(f"P{n}", civil=f"2900101123{n:02d}", phone=f"987654{n:02d}", branch=main_branch, **kw)
            for n, kw in rows.items()
        ],
        driver_columns=True,
    )
    r = post(admin_client, data)
    out = r.json()
    assert {(e["row"], e["code"]) for e in out["errors"]} == {
        (3, "invalid_yes_no"),
        (4, "invalid_initial_password"),
        (5, "claim_days_out_of_range"),
        (6, "claim_days_out_of_range"),
        (7, "claim_days_out_of_range"),
        (8, "not_a_driver"),
        (9, "password_without_app"),
        (10, "initial_password_is_civil_id"),
    }
    params = {e["row"]: e["params"] for e in out["errors"]}
    assert params[3] == {"field": "app_access", "value": "ربما"} and params[7] == {"value": "x"}
    assert params[4] == {"field": "initial_password"} and params[10] == {}  # never the password
    for secret in ("shortpw", "Super-pass-9", "Other-pass-1", "with spaces"):
        assert secret not in r.text
    assert out["claims"] == 0 and out["applied"] is False  # no password is set while the file has errors


def test_password_column_needs_devices_manage(admin_client, new_client, companies, main_branch):
    make_user(admin_client, "importer", permissions=list(NEEDED))
    c = new_client()
    login(c, "importer")
    data = workbook(
        people=[
            person("A", civil="290010112371", phone="98765471", branch=main_branch, password="Init-pass-1"),
            person("B", civil="290010112372", phone="98765472", branch=main_branch, app="نعم"),
            person("C", civil="290010112373", phone="98765473", branch=main_branch, days=7),
        ],
        driver_columns=True,
    )
    out = post(c, data).json()
    assert [(e["row"], e["code"], e["params"]) for e in out["errors"]] == [
        (3, "permission_denied", {"permission": "devices.manage"})
    ]
    warnings = [(w["row"], w["code"]) for w in out["warnings"] if w["code"] != "salary_skipped"]
    assert warnings == [(5, "password_days_ignored")]
    assert out["activated"] == 2  # the app column needs no more than importing does
    assert post(admin_client, data).json()["errors"] == []


def test_the_template_to_fill_has_the_driver_columns(admin_client, new_client, companies, main_branch):
    r = admin_client.get("/api/v1/imports/workbook/template")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats"), r.text
    book = openpyxl.load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == [INSTRUCTIONS, VEHICLES, PEOPLE, OPENING]
    for title, columns in (
        (VEHICLES, VEHICLE_COLUMNS),
        (PEOPLE, PEOPLE_COLUMNS + DRIVER_COLUMNS),
        (OPENING, OPENING_COLUMNS),
    ):
        ws = book[title]
        assert [c.value for c in ws[1]] == columns, title
        # every example filled, but the password: a hint there would be a public password copied with the row
        empty = 1 if title == PEOPLE else 0
        filled = len([c for c in ws[2] if c.value is not None])
        assert ws["A2"].value == "مثال" and filled == len(columns) - empty, title
    ws = book[PEOPLE]
    assert [ws[f"{col}2"].value for col in "NOP"] == ["نعم", None, 14]
    assert ws["O3"].number_format == "@" and ws.column_dimensions["O"].number_format == "@"  # 01234567 stays text
    rules = {str(v.sqref): (v.type, v.formula1, v.formula2) for v in ws.data_validations.dataValidation}
    assert rules == {"N2:N1000": ("list", '"نعم,لا"', None), "P2:P1000": ("whole", "1", "60")}
    assert post(admin_client, r.content).json()["errors"] == []  # as downloaded: nothing to import, nothing wrong

    # filled in by the office: imported as it is
    for i, value in enumerate([1, *person(branch=main_branch, app="نعم", password="01234567", days=3)], start=1):
        ws.cell(3, i, value)
    buf = io.BytesIO()
    book.save(buf)
    out = post(admin_client, buf.getvalue(), apply=True).json()
    assert out["applied"] and (out["people"]["created"], out["activated"], out["claims"]) == (1, 1, 1), out

    make_user(admin_client, "viewer", permissions=["employees.view"])
    c = new_client()
    login(c, "viewer")
    assert c.get("/api/v1/imports/workbook/template").status_code == 403


def test_driver_headers_written_differently_are_found_and_unknown_ones_said(admin_client, companies, main_branch):
    a, b = "290010112371", "290010112372"
    data = workbook(
        people=[
            person("A", civil=a, phone="98765471", branch=main_branch, app="نعم"),
            person("B", civil=b, phone="98765472", branch=main_branch, app="لا"),
        ],
        driver_columns=True,
    )
    book = openpyxl.load_workbook(io.BytesIO(data))
    ws = book[PEOPLE]
    ws.cell(1, len(PEOPLE_COLUMNS) + 1).value = "تفعيل التطبيق - App"  # a part of the header is enough
    ws.cell(1, len(PEOPLE_COLUMNS) + 4).value = "ملاحظات"
    buf = io.BytesIO()
    book.save(buf)
    out = post(admin_client, buf.getvalue(), apply=True).json()
    assert out["applied"] and out["activated"] == 1, out  # B said no, and the no was read
    assert [(w["code"], w["params"]) for w in out["warnings"]] == [("import_column_ignored", {"column": "ملاحظات"})]
    staff = {e["civil_id"]: e for e in admin_client.get("/api/v1/employees", params={"limit": 50}).json()}
    assert (staff[a]["app_access"], staff[b]["app_access"]) == ("active", "none")


def test_the_template_hint_is_never_a_password(admin_client, companies, main_branch):
    data = workbook(
        people=[person("A", civil="290010112381", branch=main_branch, password="(8 أحرف على الأقل)")],
        driver_columns=True,
    )
    out = post(admin_client, data).json()
    assert [(e["row"], e["code"]) for e in out["errors"]] == [(3, "invalid_initial_password")], out  # after the example
