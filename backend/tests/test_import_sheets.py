"""Importing a client's own workbook as it is (any sheet names and column order, a sheet without a header): the
preview suggests the kind and columns, the user confirms, then the same check-then-import as the template."""

import io
import json
from datetime import UTC, datetime, timedelta

import openpyxl
import pytest

from app.modules.imports.api import NEEDED
from tests.conftest import login, make_user
from tests.test_claims import choose, phone_codes

P = "/api/v1/imports/sheets"
W = (2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)


def civil(prefix: str) -> str:
    """A civil ID with a valid check digit."""
    check = 11 - sum(int(d) * w for d, w in zip(prefix, W, strict=True)) % 11
    assert check < 10, prefix
    return prefix + str(check)


ALI, OMAR, SARA, NOOR = civil("28609201527"), civil("29602200536"), civil("29704180474"), civil("28107030921")
HANI = civil("28501010001")
EXPIRED = datetime.now() - timedelta(days=90)


def client_workbook(*, employees=None, vehicles=None) -> bytes:
    """Laid out like the client's file: two empty salary templates, a headerless employees sheet
    (serial, civil ID, name, profession), and a vehicles sheet with bilingual headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "نموذج رواتب كيتا"
    ws.append(["driver id", "اسم السائق", "المهنه", "الرقم المدني", "رقم الايبان", "صافي الراتب"])
    wb.create_sheet("نموذج رواتب طلبات").append(["driver id", "اسم السائق", "الرقم المدني"])
    ws = wb.create_sheet("اجمالي السائقين")
    ws.append([None, None, None, None])
    rows = employees or [
        (1, int(ALI), "علي حسن", "سائق / سيارة خصوصي"),
        (2, int(OMAR), "عمر سالم", "سائق / دراجة نارية"),
        (3, int(SARA), "سارة أحمد", "مدير إداري"),
        (4, int(NOOR), "نور الدين", "سائق / سيارة خصوصي"),
    ]
    for r in rows:
        ws.append(list(r))
    ws = wb.create_sheet("اجمالي المركبات")
    ws.append(
        [
            "s",
            "رقم المركبة - car number",
            "نوع المركبة - vichel tybe",
            "الموديل - year",
            "اللون - color",
            "تاريخ انتهاء الدفتر - date of expier",
        ]
    )
    for r in vehicles or [
        (1, "60-76505", "شانجان - changan", 2024, "ابيض - white", datetime(2027, 6, 2)),
        (2, "60- 76303", "نيسان - nissan", 2022, "ابيض - white", EXPIRED),
        (3, "23-68409", "نيسان - nissan", 2024, "ابيض - white", datetime(2027, 9, 25)),
    ]:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def employees_book(header: list | None, rows: list) -> bytes:
    """One employees sheet, with a header row or without, as an office adds columns to its own file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "السائقون"
    for r in ([header] if header else []) + rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview(client, data: bytes) -> dict:
    r = client.post(f"{P}/preview", files={"file": ("x.xlsx", data, "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def plan_from(pv: dict, company_id: int, **extra) -> dict:
    sheets = [
        {"name": s["name"], "kind": s["kind"], "first_row": s["first_row"], "columns": s["mapping"]}
        for s in pv["sheets"]
        if s["kind"]
    ]
    return {"company_id": company_id, "sheets": sheets} | extra


def run(client, data: bytes, plan: dict, *, apply=False):
    return client.post(
        P,
        params={"apply": str(apply).lower()},
        files={"file": ("x.xlsx", data, "application/octet-stream")},
        data={"plan": json.dumps(plan)},
    )


def test_the_preview_reads_the_clients_layout(admin_client):
    pv = preview(admin_client, client_workbook())
    by = {s["name"]: s for s in pv["sheets"]}
    assert by["نموذج رواتب كيتا"]["kind"] is None and by["نموذج رواتب كيتا"]["rows"] == 0  # empty templates
    emp = by["اجمالي السائقين"]
    assert emp["kind"] == "employees" and emp["header_row"] is None and emp["first_row"] == 2 and emp["rows"] == 4
    assert emp["mapping"] == {"civil_id": 1, "name": 2, "job_title": 3}
    veh = by["اجمالي المركبات"]
    assert veh["kind"] == "vehicles" and veh["header_row"] == 1 and veh["first_row"] == 2
    assert veh["mapping"] == {"plate_number": 1, "make": 2, "year": 3, "color": 4, "registration_expiry": 5}
    assert veh["columns"][5]["samples"][0] == "02/06/2027"
    assert {f["key"] for f in pv["fields"]["employees"] if f["required"]} == {"civil_id", "name"}


def test_check_then_import_then_import_again(admin_client, company):
    data = client_workbook()
    plan = plan_from(preview(admin_client, data), company["id"])
    r = run(admin_client, data, plan)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["applied"] is False and out["errors"] == []
    assert out["vehicles"] == {"created": 3, "updated": 0} and out["people"] == {"created": 4, "updated": 0}
    assert out["documents"] == 3
    warnings = sorted((w["sheet"], w["row"], w["code"]) for w in out["warnings"])
    assert warnings == [
        ("اجمالي السائقين", 2, "driver_without_phone"),
        ("اجمالي السائقين", 3, "driver_without_phone"),
        ("اجمالي السائقين", 5, "driver_without_phone"),
        ("اجمالي المركبات", 3, "registration_expired"),
    ]
    assert admin_client.get("/api/v1/vehicles", params={"limit": 50}).json() == []  # the check changed nothing

    out = run(admin_client, data, plan, apply=True).json()
    assert out["applied"] is True
    vehicles = {v["plate_number"]: v for v in admin_client.get("/api/v1/vehicles", params={"limit": 50}).json()}
    assert set(vehicles) == {"60-76505", "60- 76303", "23-68409"}
    v = vehicles["60-76505"]
    assert (v["make"], v["year"], v["color"], v["company_id"]) == (
        "شانجان - changan",
        2024,
        "ابيض - white",
        company["id"],
    )
    docs = admin_client.get("/api/v1/documents", params={"owner_type": "vehicle", "owner_id": v["id"]}).json()
    assert [(d["type_code"], d["expiry_date"]) for d in docs] == [("registration", "2027-06-02")]
    staff = {e["civil_id"]: e for e in admin_client.get("/api/v1/employees", params={"limit": 50}).json()}
    assert staff[ALI]["is_driver"] and staff[ALI]["job_title"] == "سائق / سيارة خصوصي"
    assert staff[ALI]["app_access"] == "none" and staff[ALI]["phone"] is None
    assert staff[OMAR]["is_driver"] and not staff[SARA]["is_driver"]
    assert staff[SARA]["name"]["ar"] == "سارة أحمد" and staff[SARA]["employee_number"].startswith("E")

    # the same file again updates, never duplicates; a phone added later gives the driver the app
    out = run(admin_client, data, plan, apply=True).json()
    assert out["vehicles"] == {"created": 0, "updated": 3} and out["people"] == {"created": 0, "updated": 4}
    assert out["documents"] == 0
    with_phone = client_workbook(employees=[(1, int(ALI), "علي حسن", "سائق / سيارة خصوصي", "98765432")])
    p2 = plan_from(preview(admin_client, with_phone), company["id"])
    sheet = next(s for s in p2["sheets"] if s["kind"] == "employees")
    assert sheet["columns"].get("phone") == 4
    out = run(admin_client, with_phone, p2, apply=True).json()
    assert out["applied"] and not [w for w in out["warnings"] if w["code"] == "driver_without_phone"]
    ali = next(e for e in admin_client.get("/api/v1/employees", params={"q": ALI}).json() if e["civil_id"] == ALI)
    assert ali["phone"] == "+96598765432" and ali["app_access"] == "active"


def test_row_errors_block_the_import_and_name_the_row(admin_client, company):
    good = civil("29001011234")
    wrong_check = good[:11] + str((int(good[11]) + 1) % 10)
    data = client_workbook(
        employees=[
            (1, int(ALI), "علي", "سائق"),
            (2, int(ALI), "علي مكرر", "سائق"),
            (3, 12345, "رقم قصير", "سائق"),
            (4, int(wrong_check), "تحقق", "مدير"),
        ],
        vehicles=[(1, "11-22222", "x", 1900, "w", datetime(2027, 1, 1))],
    )
    plan = {
        "company_id": company["id"],
        "sheets": [
            {"name": "اجمالي السائقين", "kind": "employees", "first_row": 2, "columns": {"civil_id": 1, "name": 2}},
            {"name": "اجمالي المركبات", "kind": "vehicles", "first_row": 2, "columns": {"plate_number": 1, "year": 3}},
        ],
    }
    out = run(admin_client, data, plan, apply=True).json()
    assert out["applied"] is False
    errors = sorted((e["sheet"], e["row"], e["code"]) for e in out["errors"])
    assert errors == [
        ("اجمالي السائقين", 2, "duplicate_in_file"),
        ("اجمالي السائقين", 3, "duplicate_in_file"),
        ("اجمالي السائقين", 4, "invalid_civil_id"),
        ("اجمالي المركبات", 2, "invalid_number"),
    ]
    assert [w["row"] for w in out["warnings"] if w["code"] == "civil_id_check_digit"] == [5]
    assert admin_client.get("/api/v1/employees", params={"limit": 50}).json() == []


@pytest.mark.parametrize(
    "change, code",
    [
        (lambda p: p["sheets"][0]["columns"].pop("civil_id"), "import_bad_mapping"),
        (lambda p: p["sheets"][0]["columns"].update(name=99), "import_bad_mapping"),
        (lambda p: p["sheets"][0]["columns"].update(salary=4), "import_bad_mapping"),
        (lambda p: p["sheets"][0].update(name="غير موجودة"), "import_sheet_missing"),
        (lambda p: p.update(company_id=999), "company_not_found"),
    ],
)
def test_a_wrong_plan_is_refused(admin_client, company, change, code):
    data = client_workbook()
    plan = plan_from(preview(admin_client, data), company["id"])
    plan["sheets"].sort(key=lambda s: s["kind"] != "employees")
    change(plan)
    r = run(admin_client, data, plan)
    assert r.status_code == 422 and r.json()["code"] == code, r.text


def test_an_iban_is_checked_and_needs_the_salary_permission(admin_client, new_client, company):
    data = client_workbook(employees=[(1, int(ALI), "علي", "سائق", "KW81CBKU0000000000001234560101")])
    plan = plan_from(preview(admin_client, data), company["id"])
    plan["sheets"] = [s for s in plan["sheets"] if s["kind"] == "employees"]
    plan["sheets"][0]["columns"]["iban"] = 4
    out = run(admin_client, data, plan, apply=True).json()
    assert out["applied"], out
    bad = client_workbook(employees=[(1, int(OMAR), "عمر", "سائق", "KW81CBKU0000000000001234560102")])
    assert [e["code"] for e in run(admin_client, bad, plan).json()["errors"]] == ["invalid_iban"]

    needed = ["employees.create", "employees.update", "vehicles.create", "vehicles.update", "documents.manage"]
    make_user(admin_client, "importer", permissions=needed)
    c = new_client()
    login(c, "importer")
    good = client_workbook(employees=[(1, int(SARA), "سارة", "مدير", "KW81CBKU0000000000001234560101")])
    out = run(c, good, plan, apply=True).json()
    assert out["applied"] and [w["code"] for w in out["warnings"]] == ["salary_skipped"]
    make_user(admin_client, "viewer", permissions=["employees.view"])
    c2 = new_client()
    login(c2, "viewer")
    assert c2.post(f"{P}/preview", files={"file": ("x.xlsx", good, "x")}).status_code == 403


def test_a_nationality_as_the_sheet_writes_it_becomes_the_lists(admin_client, company):
    data = client_workbook(employees=[(1, int(ALI), "علي", "سائق", "هندي"), (2, int(OMAR), "عمر", "سائق", "Egyptian")])
    plan = plan_from(preview(admin_client, data), company["id"])
    plan["sheets"] = [s for s in plan["sheets"] if s["kind"] == "employees"]
    plan["sheets"][0]["columns"]["nationality"] = 4
    assert run(admin_client, data, plan, apply=True).json()["applied"]
    staff = {
        e["civil_id"]: e["nationality"] for e in admin_client.get("/api/v1/employees", params={"limit": 50}).json()
    }
    assert (staff[ALI], staff[OMAR]) == ("الهند", "مصر")
    bad = client_workbook(employees=[(1, int(SARA), "سارة", "مدير", "مريخية")])
    errors = run(admin_client, bad, plan).json()["errors"]
    assert [(e["code"], e["params"]["value"]) for e in errors] == [("nationality_not_matched", "مريخية")]


def test_drivers_without_a_phone_get_the_initial_password(admin_client, client, new_client, company, owner_db):
    from sqlalchemy import text

    data = client_workbook()
    plan = plan_from(preview(admin_client, data), company["id"]) | {"claim_password": "12345678", "claim_days": 7}
    out = run(admin_client, data, plan).json()
    assert out["claims"] == 3 and out["applied"] is False  # the three drivers, not the office manager
    assert owner_db.execute(text("SELECT count(*) FROM identity.driver_claims")).scalar() == 0  # a check keeps nothing
    out = run(admin_client, data, plan, apply=True).json()
    assert out["claims"] == 3 and out["applied"]
    r = client.post(
        "/api/v1/driver/auth/claim", json={"civil_id": ALI, "password": "12345678", "device_uid": "phone-ali-01"}
    )
    assert r.status_code == 200 and r.json()["name"]["ar"] == "علي حسن", r.text
    r = client.post(
        "/api/v1/driver/auth/claim", json={"civil_id": SARA, "password": "12345678", "device_uid": "phone-sara-1"}
    )
    assert r.status_code == 401  # not a driver: no initial password

    # passwords for drivers are given by whoever manages their phones, not by anyone allowed to import
    make_user(admin_client, "importer", permissions=list(NEEDED))
    c = new_client()
    login(c, "importer")
    r = run(c, data, plan)
    assert r.status_code == 403 and r.json()["params"]["permission"] == "devices.manage", r.text
    assert run(c, data, {k: v for k, v in plan.items() if k != "claim_password"}).status_code == 200


# ------------------------------------------------------------------ the driver app, initial password and its days

CLAIM = "/api/v1/driver/auth/claim"


def sign_in(client, civil_id: str, password: str, device: str = "phone-0001"):
    return client.post(CLAIM, json={"civil_id": civil_id, "password": password, "device_uid": device})


def staff(admin_client) -> dict:
    return {e["civil_id"]: e for e in admin_client.get("/api/v1/employees", params={"limit": 50}).json()}


def test_the_password_columns_are_found_and_never_shown(admin_client):
    data = employees_book(
        ["الرقم المدني", "الاسم", "المهنة", "التطبيق", "Password", "Days"],
        [(ALI, "علي", "سائق", "نعم", "Secret-pass-1", 7), (OMAR, "عمر", "سائق", "لا", None, None)],
    )
    r = admin_client.post(f"{P}/preview", files={"file": ("x.xlsx", data, "application/octet-stream")})
    sheet = r.json()["sheets"][0]
    assert sheet["kind"] == "employees" and sheet["mapping"] == {
        "civil_id": 0,
        "name": 1,
        "job_title": 2,
        "app_access": 3,
        "initial_password": 4,
        "password_days": 5,
    }
    assert sheet["columns"][4]["samples"] == ["••••••••", ""] and "Secret-pass-1" not in r.text
    assert sheet["columns"][3]["samples"] == ["نعم", "لا"]
    optional = {"app_access", "initial_password", "password_days"}
    assert {f["key"] for f in r.json()["fields"]["employees"] if not f["required"]} >= optional


def test_per_row_password_and_days(admin_client, client, company):
    header = ["الرقم المدني", "الاسم", "المهنة", "تفعيل التطبيق", "كلمة المرور", "صالحة لمدة"]
    data = employees_book(
        header, [(ALI, "علي", "سائق", None, "Row-pass-11", 5), (OMAR, "عمر", "سائق", "نعم", None, None)]
    )
    plan = plan_from(preview(admin_client, data), company["id"])
    out = run(admin_client, data, plan).json()
    assert [(e["row"], e["code"]) for e in out["errors"]] == [(3, "app_needs_phone_or_password")]  # phone codes on

    data = employees_book(header, [(ALI, "علي", "سائق", None, "Row-pass-11", 5)])
    out = run(admin_client, data, plan, apply=True).json()
    assert out["applied"] and (out["activated"], out["claims"]) == (0, 1) and out["warnings"] == [], out
    ali = staff(admin_client)[ALI]
    assert ali["app_access"] == "none"  # his first sign-in gives him the app, with the phone he verifies
    claim = admin_client.get(f"/api/v1/employees/{ali['id']}/claim").json()
    expires = datetime.fromisoformat(claim["expires_at"])
    assert claim["open"] and abs(expires - (datetime.now(UTC) + timedelta(days=5))) < timedelta(minutes=5)
    r = sign_in(client, ALI, "Row-pass-11")
    assert r.status_code == 200 and r.json()["next"] == "phone", r.text

    # without phone codes a driver needs no phone for the app: activated now, and told how he will sign in
    phone_codes(admin_client, False)
    data = employees_book(
        header, [(NOOR, "نور", "سائق", "نعم", None, None), (OMAR, "عمر", "سائق", None, "Row-pass-22", None)]
    )
    out = run(admin_client, data, plan, apply=True).json()
    assert out["applied"] and (out["activated"], out["claims"]) == (2, 1), out
    assert [(w["row"], w["code"]) for w in out["warnings"]] == [(2, "driver_no_sign_in_yet")]
    people = staff(admin_client)
    assert (people[NOOR]["app_access"], people[OMAR]["app_access"]) == ("active", "active")
    claim = admin_client.get(f"/api/v1/employees/{people[OMAR]['id']}/claim").json()
    expires = datetime.fromisoformat(claim["expires_at"])
    assert abs(expires - (datetime.now(UTC) + timedelta(days=14))) < timedelta(minutes=5)  # the plan's days
    assert sign_in(client, OMAR, "Row-pass-22", "phone-0002").json()["next"] == "password"


def test_a_row_password_wins_over_the_plan_one(admin_client, client, new_client, company):
    data = employees_book(
        ["الرقم المدني", "الاسم", "المهنة", "كلمة المرور"],
        [(ALI, "علي", "سائق", "Row-pass-33"), (OMAR, "عمر", "سائق", None), (SARA, "سارة", "مدير", None)],
    )
    plan = plan_from(preview(admin_client, data), company["id"])
    assert plan["sheets"][0]["columns"]["initial_password"] == 3
    out = run(admin_client, data, plan | {"claim_password": "Plan-pass-1", "claim_days": 7}, apply=True).json()
    assert out["applied"] and out["claims"] == 2 and out["warnings"] == [], out
    assert sign_in(client, ALI, "Plan-pass-1").status_code == 401
    assert sign_in(client, ALI, "Row-pass-33").status_code == 200
    assert sign_in(client, OMAR, "Plan-pass-1", "phone-0002").status_code == 200
    assert sign_in(client, SARA, "Plan-pass-1", "phone-0003").status_code == 401  # not a driver

    # a password column is a password, plan or not: it needs whoever manages the drivers' phones
    make_user(admin_client, "importer", permissions=list(NEEDED))
    c = new_client()
    login(c, "importer")
    r = run(c, data, plan)
    assert r.status_code == 403 and r.json()["params"]["permission"] == "devices.manage", r.text
    plan["sheets"][0]["columns"].pop("initial_password")
    assert run(c, data, plan).status_code == 200


def test_reimport_never_resets_open_or_used_passwords(admin_client, client, company):
    """Without phone codes a driver who signed in keeps no phone: the plan's password used to reach him again on every
    import, wiping his own password and reopening the shared one."""
    phone_codes(admin_client, False)
    header = ["الرقم المدني", "الاسم", "المهنة", "كلمة المرور"]
    data = employees_book(header, [(ALI, "علي", "سائق", "Row-pass-44")])
    plan = plan_from(preview(admin_client, data), company["id"])
    out = run(admin_client, data, plan, apply=True).json()
    assert (out["activated"], out["claims"]) == (1, 1)
    out = run(admin_client, data, plan, apply=True).json()
    assert out["claims"] == 0 and [(w["row"], w["code"]) for w in out["warnings"]] == [(2, "claim_already_open")]
    r = sign_in(client, ALI, "Row-pass-44")  # open as it was
    assert r.status_code == 200, r.text
    assert choose(client, r.json()["claim_token"], "Ali-own-pass-1").status_code == 200
    out = run(admin_client, data, plan, apply=True).json()
    assert out["claims"] == 0 and [(w["row"], w["code"]) for w in out["warnings"]] == [(2, "claim_already_used")]

    # the plan's password, for drivers with no phone and no password in the file: only OMAR gets it
    data = employees_book(header[:3], [(ALI, "علي", "سائق"), (OMAR, "عمر", "سائق")])
    plan = plan_from(preview(admin_client, data), company["id"]) | {"claim_password": "Plan-pass-55"}
    for expected in ([(2, "claim_already_used")], [(2, "claim_already_used"), (3, "claim_already_open")]):
        out = run(admin_client, data, plan, apply=True).json()
        assert out["applied"] and [(w["row"], w["code"]) for w in out["warnings"]] == expected, out
    assert out["claims"] == 0
    r = sign_in(client, ALI, "Ali-own-pass-1", "phone-0009")
    assert r.status_code == 200 and r.json()["next"] == "done", r.text  # his own password still signs him in
    assert sign_in(client, ALI, "Plan-pass-55", "phone-0010").status_code == 401
    assert sign_in(client, OMAR, "Plan-pass-55", "phone-0011").json()["next"] == "password"


def test_a_shared_password_or_a_yes_no_column_is_not_guessed(admin_client):
    """Without a header row: a shared password with 8 digits in it passes for a phone, نعم/لا for a profession."""
    rows = [
        (1, int(ALI), "علي حسن", "نعم", "Pass12345678", "سائق"),
        (2, int(OMAR), "عمر سالم", "لا", "Pass12345678", "سائق"),
        (3, int(HANI), "هاني", "نعم", "Pass12345678", "سائق"),
    ]
    sheet = preview(admin_client, employees_book(None, rows))["sheets"][0]
    assert sheet["header_row"] is None and sheet["kind"] == "employees"
    assert sheet["mapping"] == {"civil_id": 1, "name": 2, "job_title": 5}  # one profession for all is still one


def test_a_driver_on_a_phone_never_gets_the_shared_password_again(admin_client, client, company, owner_db):
    """He bound his phone by a code (or a link) and never used the initial password: once it expired, a re-import of
    the same file must not reopen it, or anyone with the batch password and his civil ID takes his account."""
    from sqlalchemy import text

    from tests.conftest import bind_device

    header = ["الرقم المدني", "الاسم", "المهنة", "الهاتف", "كلمة المرور"]
    data = employees_book(header, [(ALI, "علي", "سائق", "65001122", "Row-pass-66")])
    plan = plan_from(preview(admin_client, data), company["id"])
    out = run(admin_client, data, plan, apply=True).json()
    assert out["applied"] and out["claims"] == 1, out
    bind_device(client, "+96565001122")
    owner_db.execute(text("UPDATE identity.driver_claims SET expires_at = now() - interval '1 day'"))
    owner_db.commit()
    phone_codes(admin_client, False)
    out = run(admin_client, data, plan, apply=True).json()
    assert out["claims"] == 0 and [(w["row"], w["code"]) for w in out["warnings"]] == [(2, "device_already_bound")]
    assert sign_in(client, ALI, "Row-pass-66", "phone-0042").status_code == 401


def test_no_for_the_app_says_when_an_open_password_would_still_let_him_in(admin_client, company):
    header = ["الرقم المدني", "الاسم", "المهنة", "تفعيل التطبيق", "كلمة المرور"]
    data = employees_book(header, [(OMAR, "عمر", "سائق", None, "Row-pass-77")])
    plan = plan_from(preview(admin_client, data), company["id"])
    assert run(admin_client, data, plan, apply=True).json()["claims"] == 1
    data = employees_book(header, [(OMAR, "عمر", "سائق", "لا", None)])
    out = run(admin_client, data, plan, apply=True).json()
    assert [(w["row"], w["code"]) for w in out["warnings"]] == [(2, "claim_open_app_off")], out


def test_a_password_header_with_the_word_app_is_the_password_and_never_echoed(admin_client, company):
    header = ["الرقم المدني", "الاسم", "المهنة", "كلمة المرور (التطبيق)"]
    data = employees_book(header, [(ALI, "علي", "سائق", "Ali@Pass2026")])
    pv = preview(admin_client, data)
    mapping = pv["sheets"][0]["mapping"]
    assert mapping.get("initial_password") == 3 and "app_access" not in mapping, mapping
    # mapped by hand to the app column anyway: the error does not repeat the value
    plan = plan_from(pv, company["id"])
    plan["sheets"][0]["columns"] = {"civil_id": 0, "name": 1, "job_title": 2, "app_access": 3}
    r = run(admin_client, data, plan)
    assert "Ali@Pass2026" not in r.text and r.json()["errors"][0]["code"] == "invalid_yes_no", r.text
