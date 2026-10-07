import io
import json
import uuid
from datetime import UTC, datetime, time, timedelta

import openpyxl
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError

from app.core.clock import today
from app.modules.audit import service as audit_service
from tests.conftest import APP_ROLE_PASSWORD, login, make_user


def test_audit_trail_is_append_only(admin_client, owner_db):
    assert owner_db.execute(text("SELECT count(*) FROM audit.events")).scalar() > 0
    for sql in ("UPDATE audit.events SET action = 'x'", "DELETE FROM audit.events"):  # even for the owner
        with pytest.raises(DBAPIError, match="append-only"):
            owner_db.execute(text(sql))
        owner_db.rollback()


def test_audit_view_is_company_scoped(admin_client, new_client, companies):
    a, b = companies["a"], companies["b"]
    for company in (a, b):
        admin_client.patch(f"/api/v1/companies/{company['public_id']}", json={"version": 1, "phone": "22000000"})
    make_user(admin_client, "auditor_a", permissions=["audit.view"], company_ids=[a["id"]])
    c = new_client()
    login(c, "auditor_a")
    events = c.get("/api/v1/audit", params={"entity_type": "company"}).json()
    assert events and {e["company_id"] for e in events} == {a["id"]}
    everything = admin_client.get("/api/v1/audit", params={"action": "company"}).json()
    assert {e["company_id"] for e in everything} == {a["id"], b["id"]}
    assert {e["actor_name"] for e in everything} == {"System Administrator (admin)"}  # who, readable


def test_application_role_has_least_privilege(database_url, admin_client):
    eng = create_engine(make_url(database_url["url"]).set(username="fleet_app", password=APP_ROLE_PASSWORD))
    with eng.connect() as c:
        assert c.execute(text("SHOW statement_timeout")).scalar() == "30s"
        c.execute(text("INSERT INTO audit.events (actor_type, action, entity_type) VALUES ('system', 't', 't')"))
        c.commit()
        assert c.execute(text("SELECT count(*) FROM i18n.languages")).scalar() == 2  # later schemas are granted too
        with pytest.raises(DBAPIError, match="permission denied"):
            c.execute(text("UPDATE audit.events SET action = 'x'"))
        c.rollback()
        with pytest.raises(DBAPIError, match="permission denied"):
            c.execute(text("TRUNCATE audit.events"))
    eng.dispose()


def _event(db, *, at, action, actor=None, company_id=None, entity_type="employee", after=None):
    db.execute(
        text(
            "INSERT INTO audit.events (occurred_at, actor_type, actor_user_id, action, entity_type, entity_id, "
            "company_id, after) VALUES (:t, :ty, :u, :a, :e, 'rec-1', :c, CAST(:af AS jsonb))"
        ),
        {
            "t": at,
            "ty": "user" if actor else "system",
            "u": actor,
            "a": action,
            "e": entity_type,
            "c": company_id,
            "af": json.dumps(after) if after is not None else None,
        },
    )
    db.commit()


def _uid(db, username: str) -> int:
    return db.execute(text("SELECT id FROM identity.users WHERE username = :u"), {"u": username}).scalar()


def test_search_by_who_by_kuwait_days_and_by_action_then_export(admin_client, new_client, db, companies, monkeypatch):
    a, b = companies["a"]["id"], companies["b"]["id"]
    make_user(admin_client, "clerk", permissions=["employees.view"])
    clerk_public = next(
        u["id"] for u in admin_client.get("/api/v1/audit/users").json() if u["name"].endswith("(clerk)")
    )
    admin, clerk = _uid(db, "admin"), _uid(db, "clerk")
    d = today() - timedelta(days=10)
    late = datetime.combine(d, time(22, 30), tzinfo=UTC)  # 01:30 the next day in Kuwait
    _event(db, at=late - timedelta(hours=3), action="employeeXstatus.changed", actor=admin, company_id=a)
    _event(db, at=late - timedelta(hours=2), action="employee.status_changed", actor=admin, company_id=b)
    _event(db, at=late, action="employee.updated", actor=clerk, company_id=a, after={"name": "=HYPERLINK(1)"})

    def search(**params):
        r = admin_client.get("/api/v1/audit", params=params)
        assert r.status_code == 200, r.text
        return [e["action"] for e in r.json()]

    # the late event belongs to the next Kuwait day
    assert search(date_from=str(d + timedelta(days=1)), date_to=str(d + timedelta(days=1))) == ["employee.updated"]
    assert search(date_from=str(d), date_to=str(d)) == ["employee.status_changed", "employeeXstatus.changed"]
    assert search(actor_id=clerk_public, date_from=str(d)) == ["employee.updated"]
    # "_" is a letter, not a wildcard: "employee_status" finds nothing under "employeeXstatus"
    assert search(action="employee_status") == []
    assert search(action="employee", date_from=str(d), date_to=str(d)) == ["employee.status_changed"]
    assert admin_client.get("/api/v1/audit", params={"actor_id": str(uuid.uuid4())}).status_code == 404

    # Excel in Arabic: time in Kuwait, who by name, the record type's name; CSV with the codes
    span = {"date_from": str(d), "date_to": str(d + timedelta(days=1))}
    r = admin_client.get("/api/v1/audit/export", params=span, headers={"Accept-Language": "ar"})
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('audit.xlsx"')
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert ws.title == "سجل التدقيق" and ws.sheet_view.rightToLeft
    assert rows[0][:4] == ("الوقت (الكويت)", "بواسطة", "الإجراء", "نوع السجل")
    assert [r[2] for r in rows[1:]] == ["employee.updated", "employee.status_changed", "employeeXstatus.changed"]
    first = rows[1]
    assert first[0] == f"{d + timedelta(days=1)} 01:30:00" and first[1] == f"{name_of(admin_client, 'clerk')} (clerk)"
    assert (
        first[3] == "موظف"
        and first[5] == companies["a"]["name"]["ar"]
        and json.loads(first[8]) == {"name": "=HYPERLINK(1)"}
    )
    assert ws.cell(row=2, column=9).data_type == "s"  # stays text
    csv_rows = admin_client.get("/api/v1/audit/export", params=span | {"format": "csv"}).content.decode("utf-8-sig")
    assert csv_rows.splitlines()[0] == "occurred_at,actor,action,entity_type,entity_id,company,ip,before,after"
    assert ",employee," in csv_rows.splitlines()[1]

    # a company-limited auditor exports his companies only; and an export too large is refused, not cut short
    make_user(admin_client, "auditor_b", permissions=["audit.view"], company_ids=[b])
    c = new_client()
    login(c, "auditor_b")
    r = c.get("/api/v1/audit/export", params=span | {"format": "csv"})
    assert [line.split(",")[2] for line in r.content.decode("utf-8-sig").splitlines()[1:]] == [
        "employee.status_changed"
    ]
    monkeypatch.setattr(audit_service, "EXPORT_MAX", 2)
    r = admin_client.get("/api/v1/audit/export", params=span)
    assert r.status_code == 422 and r.json()["code"] == "export_too_large"
    # who may: audit.view
    make_user(admin_client, "no_audit", permissions=["employees.view"])
    c = new_client()
    login(c, "no_audit")
    assert c.get("/api/v1/audit/export").status_code == 403 and c.get("/api/v1/audit/users").status_code == 403


def name_of(admin_client, username: str) -> str:
    return next(u["full_name"] for u in admin_client.get("/api/v1/users").json() if u["username"] == username)
