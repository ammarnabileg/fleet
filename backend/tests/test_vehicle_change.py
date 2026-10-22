"""A change of vehicle names the other car (FR-ASG-04): the plate the driver types must be one of his company's
vehicles, typed with Arabic or Latin digits and any dash; the office sees the car asked for (free, or with whom) and
keeps every request, which it can filter and export."""

import io

import pytest
from openpyxl import load_workbook

from app.modules.fleet.service import normalize_plate
from tests.conftest import bearer, bind_device, hand_over, login, make_driver, make_user, make_vehicle, upload

ASK = "/api/v1/driver/vehicle-change-requests"
V = "/api/v1/vehicle-change-requests"


def holding(admin_client, client, company, plate=None):
    """A driver on his phone, holding a vehicle."""
    d = make_driver(admin_client, company["id"])
    v = make_vehicle(admin_client, company["id"], **({"plate_number": plate} if plate else {}))
    hand_over(admin_client, v, d)
    return d, v, bearer(bind_device(client, d["phone"]))


@pytest.mark.parametrize(
    ("typed", "plate"),
    [
        ("١٢-٣٤٥٦٧", "12-34567"),
        ("۱۲-۳۴۵۶۷", "12-34567"),
        ("12–34567", "12-34567"),
        ("12− 34567", "12- 34567"),
        ("  12   ab  ", "12 AB"),
    ],
)
def test_the_plate_is_normalized(typed, plate):
    assert normalize_plate(typed) == plate


def test_a_vehicle_added_with_arabic_digits_is_kept_in_latin(admin_client, company):
    v = make_vehicle(admin_client, company["id"], plate_number="٤٥–٦٧٨")
    assert v["plate_number"] == "45-678"


@pytest.mark.parametrize("typed", ["٣٣-٤٤٥٥٦", "33–44556", "33 44556", "3344556", "33/44556", " 33-44556 "])
def test_the_other_car_found_however_he_types_its_plate(admin_client, client, company, typed):
    other = make_vehicle(admin_client, company["id"], plate_number="33-44556")
    _, _, h = holding(admin_client, client, company)
    r = client.post(ASK, headers=h, json={"requested_plate": typed, "reason": "الفرامل ضعيفة"})
    assert r.status_code == 201, r.text
    assert r.json()["change_request"]["requested_plate"] == other["plate_number"]


def test_a_plate_not_found_his_own_car_and_another_companys_car_are_refused(admin_client, client, companies):
    a, b = companies["a"], companies["b"]
    make_vehicle(admin_client, b["id"], plate_number="55-12345")
    d, v, h = holding(admin_client, client, a, plate="77-1")

    def ask(plate):
        return client.post(ASK, headers=h, json={"requested_plate": plate, "reason": "الفرامل ضعيفة"})

    for plate in ("99-99999", "55-12345"):  # unknown, or another company's: the same answer, he types it again
        r = ask(plate)
        assert r.status_code == 422 and r.json()["code"] == "vehicle_plate_not_found", r.text
    r = ask("٧٧-١")
    assert r.status_code == 422 and r.json()["code"] == "vehicle_change_same_vehicle"
    assert client.get("/api/v1/driver/vehicle", headers=h).json()["change_request"] is None
    assert admin_client.get(V, params={"status": "all"}).json() == []


def test_the_office_list_shows_the_car_asked_for_and_stays_in_its_companies(
    admin_client, client, new_client, companies
):
    a, b = companies["a"], companies["b"]
    d1, v1, h1 = holding(admin_client, client, a)
    d2, v2, h2 = holding(admin_client, client, a)
    d3, v3, h3 = holding(admin_client, client, b)
    free = make_vehicle(admin_client, b["id"], plate_number="88-100")
    # d1 asks for d2's car; d2 asks for d1's; d3 (company B) for a free one
    assert (
        client.post(ASK, headers=h1, json={"requested_plate": v2["plate_number"], "reason": "أقرب لبيتي"}).status_code
        == 201
    )
    assert (
        client.post(ASK, headers=h2, json={"requested_plate": v1["plate_number"], "reason": "أكبر"}).status_code == 201
    )
    assert client.post(ASK, headers=h3, json={"requested_plate": "88-100", "reason": "أحدث"}).status_code == 201

    everything = admin_client.get(V, params={"status": "all"}).json()
    assert [x["driver"]["id"] for x in everything] == [d3["id"], d2["id"], d1["id"]]  # newest first
    assert [x["driver"]["id"] for x in admin_client.get(V).json()] == [d1["id"], d2["id"], d3["id"]]  # the queue
    first = everything[2]
    assert first["requested_vehicle"]["id"] == v2["id"]
    assert first["requested_vehicle"]["holder"]["id"] == d2["id"]
    assert first["requested_vehicle"]["status"] == "assigned"
    assert everything[0]["requested_vehicle"] == {
        "found": True,
        "id": free["id"],
        "plate": "88-100",
        "status": "available",
        "holder": None,
    }

    # filters: a driver's requests, a vehicle's (held or asked for)
    assert [x["driver"]["id"] for x in admin_client.get(V, params={"status": "all", "driver_id": d2["id"]}).json()] == [
        d2["id"]
    ]
    by_car = admin_client.get(V, params={"status": "all", "vehicle_id": v1["id"]}).json()
    assert sorted(x["driver"]["id"] for x in by_car) == sorted([d1["id"], d2["id"]])
    assert admin_client.get(V, params={"status": "nope"}).status_code == 422
    page = admin_client.get(V, params={"status": "all", "limit": 1, "offset": 1}).json()
    assert [x["driver"]["id"] for x in page] == [d2["id"]]

    # a supervisor of company A sees only its requests, and cannot filter on company B's
    make_user(admin_client, "sup_a", permissions=["custody.view", "custody.assign"], company_ids=[a["id"]])
    sup = new_client()
    login(sup, "sup_a")
    assert sorted(x["driver"]["id"] for x in sup.get(V, params={"status": "all"}).json()) == sorted(
        [d1["id"], d2["id"]]
    )
    assert sup.get(V, params={"status": "all", "driver_id": d3["id"]}).status_code == 404
    assert sup.get(V, params={"status": "all", "vehicle_id": v3["id"]}).status_code == 404

    # the export: the same rows, in Excel
    r = sup.get(f"{V}/export", headers={"Accept-Language": "ar"})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.values)
    assert rows[0][:4] == ("التاريخ", "السائق", "السيارة التي معه", "السيارة المطلوبة")
    assert len(rows) == 3
    assert {row[3] for row in rows[1:]} == {v1["plate_number"], v2["plate_number"]}
    assert {row[6] for row in rows[1:]} == {"قيد المراجعة"}
    assert any(str(row[4]).startswith("مع ") for row in rows[1:])
    assert sup.get(f"{V}/export", params={"driver_id": d3["id"]}).status_code == 404


# ------------------------------------------------------------------ "done" only when the requested car is handed over


def told(client, h) -> list[str]:
    return [n["kind"] for n in client.get("/api/v1/driver/notifications", headers=h).json()["items"]]


def asked(admin_client, client, h, plate, reason="الفرامل ضعيفة"):
    r = client.post(ASK, headers=h, json={"requested_plate": plate, "reason": reason})
    assert r.status_code == 201, r.text
    (row,) = admin_client.get(V, params={"status": "open"}).json()
    return row


def vehicle(admin_client, v):
    return admin_client.get(f"/api/v1/vehicles/{v['id']}").json()


def open_custody(admin_client, v):
    return [
        x
        for x in admin_client.get("/api/v1/custodies", params={"open": "true"}).json()
        if x["vehicle"]["id"] == v["id"]
    ]


def test_approved_waits_for_the_handover_and_the_transfer_makes_it_done_once(admin_client, client, company):
    """Approving never moves a car: the request waits; the transfer of the requested car to him (his own given back in
    the same step) is what makes it done, linked to that custody, the two cars and custodies consistent."""
    wanted = make_vehicle(admin_client, company["id"], km=20_000)
    d, v, h = holding(admin_client, client, company)
    row = asked(admin_client, client, h, wanted["plate_number"])
    assert row["status"] == "pending" and row["unverified"] is False

    r = admin_client.post(f"{V}/{row['id']}/approve", json={})
    assert r.status_code == 200 and r.json()["status"] == "approved" and r.json()["approved_at"], r.text
    assert admin_client.post(f"{V}/{row['id']}/approve", json={}).json()["code"] == "vehicle_change_decided"
    assert told(client, h)[0] == "vehicle_change_approved"
    # nothing moved: his car still his, the requested one still free
    assert vehicle(admin_client, v)["status"] == "assigned" and vehicle(admin_client, wanted)["status"] == "available"
    assert admin_client.get(V, params={"status": "approved"}).json()[0]["id"] == row["id"]
    assert [a for a in admin_client.get("/api/v1/alerts").json() if a["kind"] == "vehicle_change_requested"] == []

    body = {
        "vehicle_id": wanted["id"],
        "driver_id": d["id"],
        "mode": "release",
        "odometer_km": 20_010,
        "photo_sha256": upload(admin_client),
        "other_odometer_km": v["last_odometer_km"] + 30,
        "other_photo_sha256": upload(admin_client),
    }
    r = admin_client.post("/api/v1/custodies/transfer", json=body)
    assert r.status_code == 201, r.text
    started = r.json()["started"][0]
    (done,) = admin_client.get(V, params={"status": "done"}).json()
    assert done["id"] == row["id"] and done["handed_over_at"] == started["started_at"] and not done["unverified"]
    assert vehicle(admin_client, v)["status"] == "available" and open_custody(admin_client, v) == []
    assert vehicle(admin_client, wanted)["status"] == "assigned"
    assert [c["driver"]["id"] for c in open_custody(admin_client, wanted)] == [d["id"]]
    assert told(client, h)[:2] == ["vehicle_handed_over", "vehicle_change_done"]
    events = admin_client.get("/api/v1/audit", params={"action": "vehicle_change.done"}).json()
    items = events["items"] if isinstance(events, dict) else events
    assert len([e for e in items if e["action"] == "vehicle_change.done"]) == 1
    # closed once: nothing left to decide
    for decision in ("approve", "reject", "close"):
        assert admin_client.post(f"{V}/{row['id']}/{decision}", json={"note": "متأخر"}).status_code == 409


def test_returned_then_another_car_closes_it_without_the_change(admin_client, client, company):
    wanted = make_vehicle(admin_client, company["id"])
    other = make_vehicle(admin_client, company["id"])
    d, v, h = holding(admin_client, client, company)
    row = asked(admin_client, client, h, wanted["plate_number"])
    (c,) = open_custody(admin_client, v)
    r = admin_client.post(
        f"/api/v1/custodies/{c['id']}/return",
        json={"odometer_km": v["last_odometer_km"] + 5, "photo_sha256": upload(admin_client)},
    )
    assert r.status_code == 200, r.text
    # returning his car is not the change he asked for
    assert admin_client.get(V, params={"status": "pending"}).json()[0]["id"] == row["id"]
    hand_over(admin_client, other, d)
    (closed,) = admin_client.get(V, params={"status": "closed"}).json()
    assert closed["id"] == row["id"] and closed["note"] and closed["handed_over_at"] is None
    assert admin_client.get(V, params={"status": "done"}).json() == []
    assert told(client, h)[:2] == ["vehicle_handed_over", "vehicle_change_closed"]


def test_the_office_closes_without_the_change_only_with_a_note(admin_client, client, company):
    wanted = make_vehicle(admin_client, company["id"])
    _, v, h = holding(admin_client, client, company)
    row = asked(admin_client, client, h, wanted["plate_number"])
    assert admin_client.post(f"{V}/{row['id']}/close", json={}).json()["code"] == "reason_required"
    r = admin_client.post(f"{V}/{row['id']}/close", json={"note": "اتصلح المكيف"})
    assert r.status_code == 200 and r.json()["status"] == "closed"
    assert told(client, h)[0] == "vehicle_change_closed"
    assert vehicle(admin_client, v)["status"] == "assigned"
    # he may ask again: nothing open on his custody
    assert (
        client.post(ASK, headers=h, json={"requested_plate": wanted["plate_number"], "reason": "تاني"}).status_code
        == 201
    )
    (again,) = admin_client.get(V).json()
    r = admin_client.post(f"{V}/{again['id']}/approve", json={})
    assert r.status_code == 200
    r = admin_client.post(f"{V}/{again['id']}/reject", json={"note": "العربية رايحة الصيانة"})
    assert r.status_code == 200 and r.json()["status"] == "rejected"  # an approval can still be withdrawn


def test_a_claim_of_the_requested_car_makes_it_done(admin_client, client, company):
    """However the custody starts (here his own registration of the car, approved), the requested car makes it done."""
    from tests.test_vehicle_claims import ask as claim

    wanted = make_vehicle(admin_client, company["id"], plate_number="61-77001", km=10_000)
    d, v, h = holding(admin_client, client, company)
    row = asked(admin_client, client, h, "61-77001")
    admin_client.post(f"{V}/{row['id']}/approve", json={})
    (c,) = open_custody(admin_client, v)
    admin_client.post(
        f"/api/v1/custodies/{c['id']}/return",
        json={"odometer_km": v["last_odometer_km"] + 5, "photo_sha256": upload(admin_client)},
    )
    from app.core.clock import utcnow

    r, _ = claim(client, h, "61-77001", recorded_at=utcnow().isoformat())
    assert r.status_code == 201, r.text
    (cl,) = admin_client.get("/api/v1/vehicle-claims").json()
    assert admin_client.post(f"/api/v1/vehicle-claims/{cl['id']}/approve", json={}).status_code == 200
    (done,) = admin_client.get(V, params={"status": "done"}).json()
    assert done["id"] == row["id"] and done["handed_over_at"] is not None
    assert vehicle(admin_client, wanted)["status"] == "assigned"


def test_a_request_closed_as_done_without_the_handover_is_shown_for_review(admin_client, client, company, owner_db):
    """Requests marked done before (by hand, or by the return of the car) are kept as they are: the list flags those
    with no handover of the requested car after them, filters them, and exports them so."""
    from sqlalchemy import text

    wanted = make_vehicle(admin_client, company["id"])
    d1, _, h1 = holding(admin_client, client, company)
    d2, _, h2 = holding(admin_client, client, company)
    stale = asked(admin_client, client, h1, wanted["plate_number"])
    fine = make_vehicle(admin_client, company["id"])
    client.post(ASK, headers=h2, json={"requested_plate": fine["plate_number"], "reason": "أحدث"})
    for r in admin_client.get(V, params={"status": "open"}).json():
        owner_db.execute(
            text("UPDATE fleet.vehicle_change_requests SET status = 'done', decided_at = now() WHERE public_id = :p"),
            {"p": r["id"]},
        )
    owner_db.commit()
    # d2 did get the car he asked for later (handed over by the office before the link existed)
    owner_db.execute(
        text(
            "UPDATE fleet.vehicle_change_requests SET status = 'pending', decided_at = NULL "
            "WHERE employee_id = (SELECT id FROM people.employees WHERE public_id = :d)"
        ),
        {"d": d2["id"]},
    )
    owner_db.commit()
    c2 = [
        x
        for x in admin_client.get("/api/v1/custodies", params={"open": "true"}).json()
        if x["driver"]["id"] == d2["id"]
    ]
    admin_client.post(
        f"/api/v1/custodies/{c2[0]['id']}/return",
        json={"odometer_km": 10_100, "photo_sha256": upload(admin_client)},
    )
    hand_over(admin_client, fine, d2)
    rows = {r["driver"]["id"]: r for r in admin_client.get(V, params={"status": "done"}).json()}
    assert rows[d1["id"]]["unverified"] is True and rows[d1["id"]]["handed_over_at"] is None
    assert rows[d2["id"]]["unverified"] is False and rows[d2["id"]]["handed_over_at"] is not None
    assert [r["id"] for r in admin_client.get(V, params={"status": "unverified"}).json()] == [stale["id"]]
    ws = load_workbook(io.BytesIO(admin_client.get(f"{V}/export", headers={"Accept-Language": "ar"}).content)).active
    statuses = {row[1]: row[6] for row in list(ws.values)[1:]}
    assert "دون تسليم" in statuses[d1["name"]["ar"]] and statuses[d2["name"]["ar"]] == "تم التغيير"
    # nothing corrected automatically: still done, for a person to review
    assert admin_client.get(V, params={"status": "all", "driver_id": d1["id"]}).json()[0]["status"] == "done"
