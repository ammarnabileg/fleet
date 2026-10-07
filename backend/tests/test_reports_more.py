"""Kilometers (FR-RPT-03), cash and its age (FR-RPT-04), payroll by month (FR-RPT-07), and every report's filters and
Excel export (FR-RPT-08). The numbers are worked out by hand from a known history."""

import io
from datetime import date, datetime, time, timedelta
from decimal import Decimal

import openpyxl
from sqlalchemy import text

from app.core.clock import KUWAIT, today
from app.modules.cash import service as cash
from tests.conftest import bearer, bind_device, hand_over, login, make_driver, make_user, make_vehicle, name, upload
from tests.test_payroll_runs import IBAN, MONTH, P, platforms, set_cap, shot, staff  # noqa: F401 (fixtures)

R = "/api/v1/reports"


def at(day: date, hour: int, minute: int = 0) -> datetime:
    return datetime.combine(day, time(hour, minute), tzinfo=KUWAIT)


def _id(db, table: str, public_id: str) -> int:
    return db.execute(text(f"SELECT id FROM {table} WHERE public_id = :p"), {"p": public_id}).scalar()  # noqa: S608


def reading(
    db,
    vehicle: int,
    custody: int | None,
    driver: int | None,
    kind: str,
    km: int,
    when: datetime,
    photo,
    *,
    pending=False,
):
    db.execute(
        text(
            "INSERT INTO fleet.odometer_readings (vehicle_id, custody_id, driver_id, kind, value_km, photo_sha256, "
            "recorded_at, business_date, review_status, flags) VALUES (:v, :c, :d, :k, :km, :p, :t, :b, :s, :f)"
        ),
        {
            "v": vehicle,
            "c": custody,
            "d": driver,
            "k": kind,
            "km": km,
            "p": photo,
            "t": when,
            "b": when.astimezone(KUWAIT).date(),
            "s": "pending" if pending else "ok",
            "f": ["daily_limit"] if pending else [],
        },
    )


def positions(db, vehicle: int, custody: int, driver: int, points: list[tuple[datetime, float, float]]):
    for month in sorted({p[0].astimezone(KUWAIT).date().replace(day=1) for p in points}):
        db.execute(text("SELECT tracking.ensure_month_partition(:m)"), {"m": month})
    for seq, (when, lat, lng) in enumerate(points):
        db.execute(
            text(
                "INSERT INTO tracking.positions (vehicle_id, recorded_at, custody_id, driver_id, device_id, "
                "device_seq, lat, lng) VALUES (:v, :t, :c, :d, 1, :s, :lat, :lng)"
            ),
            {"v": vehicle, "t": when, "c": custody, "d": driver, "s": seq, "lat": lat, "lng": lng},
        )


def history(admin_client, db, company):
    """One vehicle over four days: Ali holds it, returns it, nobody has it, Omar takes it, it goes to the center."""
    day0 = today() - timedelta(days=4)
    d1, d2, d3 = (day0 + timedelta(days=i) for i in (1, 2, 3))
    v = make_vehicle(admin_client, company["id"], km=10_000)
    ali, omar = make_driver(admin_client, company["id"]), make_driver(admin_client, company["id"])
    c1 = hand_over(admin_client, v, ali, km=10_000, started_at=at(day0, 8).isoformat())
    vid, a, o = (
        _id(db, "fleet.vehicles", v["id"]),
        _id(db, "people.employees", ali["id"]),
        _id(db, "people.employees", omar["id"]),
    )
    k1 = _id(db, "fleet.custodies", c1["id"])
    photo = upload(admin_client)
    for kind, km, when in (
        ("start_day", 10_010, at(day0, 9)),
        ("end_day", 10_210, at(day0, 20)),
        ("start_day", 10_240, at(d1, 8)),  # 30 km between the end of a day and the next start: off duty
        ("end_day", 10_440, at(d1, 20)),
    ):
        reading(db, vid, k1, a, kind, km, when, photo)
    db.commit()
    r = admin_client.post(
        f"/api/v1/custodies/{c1['id']}/return",
        json={"odometer_km": 10_450, "photo_sha256": upload(admin_client), "ended_at": at(d2, 7).isoformat()},
    )
    assert r.status_code == 200, r.text
    c2 = hand_over(admin_client, v, omar, km=10_470, started_at=at(d2, 8).isoformat())  # 20 km with nobody
    k2 = _id(db, "fleet.custodies", c2["id"])
    flagged = admin_client.get("/api/v1/odometer/readings", params={"review_status": "pending"}).json()
    assert [x["flags"] for x in flagged] == [["off_duty_km"]]  # the reviewer accepts it as it is
    r = admin_client.post(f"/api/v1/odometer/readings/{flagged[0]['id']}/review", json={"reason": "seen on the photo"})
    assert r.status_code == 200, r.text
    reading(db, vid, k2, o, "start_day", 19_000, at(d2, 9), photo, pending=True)  # a typo, waiting for review
    reading(db, vid, k2, o, "end_day", 10_600, at(d2, 20), photo)
    reading(db, vid, k2, o, "maintenance_in", 10_610, at(d2, 21), photo)
    reading(db, vid, None, None, "maintenance_out", 10_640, at(d3, 10), photo)  # 30 km of test drives
    # the phone's points: 0.1 degree of latitude is 11.12 km
    positions(db, vid, k1, a, [(at(day0, 10), 29.0, 48.0), (at(day0, 11), 29.1, 48.0), (at(day0, 12), 29.2, 48.0)])
    positions(db, vid, k2, o, [(at(d2, 10), 29.0, 48.0), (at(d2, 11), 29.05, 48.0)])
    db.commit()
    return {"vehicle": v, "ali": ali, "omar": omar, "day0": day0, "d1": d1}


def test_kilometers_by_vehicle_and_driver_off_duty_and_against_the_gps(admin_client, db, company):
    h = history(admin_client, db, company)
    params = {"date_from": str(h["day0"]), "date_to": str(today())}
    rep = admin_client.get(f"{R}/kilometers", params=params).json()
    [v] = rep["by_vehicle"]
    # on duty: 10 + 200 + 200 + 10 (to the return) + 130 (the typo left out) + 10 (to the center) = 560
    assert (v["km"], v["on_duty"], v["off_duty"], v["unattended"], v["center"]) == (640, 560, 30, 20, 30)
    assert (v["with_driver"], v["gps"], v["difference"]) == (590, "27.8", "562.2")
    drivers = {d["driver"]["id"]: d for d in rep["by_driver"]}
    ali, omar = drivers[h["ali"]["id"]], drivers[h["omar"]["id"]]
    assert (ali["on_duty"], ali["off_duty"], ali["days"], ali["gps"], ali["km_per_day"]) == (
        420,
        30,
        2,
        "22.2",
        "225.0",
    )
    assert (omar["on_duty"], omar["off_duty"], omar["days"], omar["gps"]) == (140, 0, 1, "5.6")
    assert rep["pending_count"] == 1 and rep["pending"][0]["value_km"] == 19_000
    assert rep["totals"]["km"] == 640 and rep["totals"]["gps"] == "27.8"

    # the reviewer corrects the typo: its corrected value counts (10 then 120 km instead of nothing)
    typo = admin_client.get("/api/v1/odometer/readings", params={"review_status": "pending"}).json()[0]
    r = admin_client.post(
        f"/api/v1/odometer/readings/{typo['id']}/review",
        json={"corrected_km": 10_480, "reason": "the photo reads 10480"},
    )
    assert r.status_code == 200, r.text
    fixed = admin_client.get(f"{R}/kilometers", params=params).json()
    assert (fixed["by_vehicle"][0]["km"], fixed["by_vehicle"][0]["on_duty"], fixed["pending_count"]) == (640, 560, 0)

    # from the second day: counted from the last reading before it (the first day's end, 10 210)
    later = admin_client.get(f"{R}/kilometers", params=params | {"date_from": str(h["d1"])}).json()["by_vehicle"][0]
    assert (later["km"], later["on_duty"], later["off_duty"]) == (430, 350, 30)

    # one driver: his stretches and his points only
    mine = admin_client.get(f"{R}/kilometers", params=params | {"driver_id": h["ali"]["id"]}).json()
    assert (mine["by_vehicle"][0]["km"], mine["by_vehicle"][0]["gps"]) == (450, "22.2")
    assert [d["driver"]["id"] for d in mine["by_driver"]] == [h["ali"]["id"]]


def test_filters_never_reach_past_the_users_companies(admin_client, new_client, db, companies):
    history(admin_client, db, companies["a"])
    params = {"date_from": str(today() - timedelta(days=4)), "date_to": str(today())}
    b = admin_client.get(f"{R}/kilometers", params=params | {"company_id": companies["b"]["id"]}).json()
    assert b["by_vehicle"] == [] and b["pending_count"] == 0
    make_user(admin_client, "km_b", permissions=["reports.view", "odometer.view"], company_ids=[companies["b"]["id"]])
    c = new_client()
    login(c, "km_b")
    assert c.get(f"{R}/kilometers", params=params).json()["by_vehicle"] == []
    r = c.get(f"{R}/kilometers", params=params | {"company_id": companies["a"]["id"]})
    assert r.status_code == 403 and r.json()["code"] == "company_out_of_scope"
    v = admin_client.get("/api/v1/vehicles").json()[0]
    assert c.get(f"{R}/kilometers", params=params | {"vehicle_id": v["id"]}).status_code == 404
    make_user(admin_client, "no_odo", permissions=["reports.view"])
    c = new_client()
    login(c, "no_odo")
    assert c.get(f"{R}/kilometers", params=params).status_code == 403


def test_excel_in_the_users_language_with_numbers_as_numbers(admin_client, db, company):
    history(admin_client, db, company)
    params = {"date_from": str(today() - timedelta(days=4)), "date_to": str(today()), "format": "xlsx"}
    r = admin_client.get(f"{R}/kilometers", params=params, headers={"Accept-Language": "ar"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="kilometers-vehicles-' in r.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "الكيلومترات" and ws.sheet_view.rightToLeft
    header = [c.value for c in ws[1]]
    assert header[:5] == ["اللوحة", "النوع", "الموديل", "الكيلومترات", "في العمل"]
    row = [c.value for c in ws[2]]
    assert row[3] == 640 and row[4] == 560 and row[9] == 27.8  # Excel keeps every number as a double
    en = admin_client.get(f"{R}/kilometers", params=params | {"section": "drivers"}, headers={"Accept-Language": "en"})
    ws = openpyxl.load_workbook(io.BytesIO(en.content)).active
    assert [c.value for c in ws[1]][:3] == ["Driver", "Days", "On duty"] and not ws.sheet_view.rightToLeft
    # CSV keeps the column keys, for other systems
    csv = admin_client.get(f"{R}/kilometers", params=params | {"format": "csv"}).content.decode("utf-8-sig")
    assert csv.splitlines()[0].startswith("plate,make,model,km,on_duty")


def test_text_that_looks_like_a_formula_stays_text_in_excel(admin_client, client, company):
    from tests.test_reports import driver_report

    d = make_driver(admin_client, company["id"], name=name("=HYPERLINK(1)", "=HYPERLINK(1)"))
    driver_report(client, bearer(bind_device(client, d["phone"])), "12.250", orders=9)
    params = {"date_from": str(today()), "date_to": str(today()), "format": "xlsx"}
    ws = openpyxl.load_workbook(io.BytesIO(admin_client.get(f"{R}/daily-summary", params=params).content)).active
    cell = ws.cell(row=2, column=2)
    assert cell.value == "=HYPERLINK(1)" and cell.data_type == "s"  # a string, not a formula
    assert ws.cell(row=2, column=4).value == 9 and ws.cell(row=2, column=5).value == 12.25


def _journal(db, admin_id, driver_id, amount: str, day: date, *, post=True):
    drv = cash.account(db, "driver", driver_id=driver_id)
    other = cash.account(db, "adjustments")
    cash._journal(
        db,
        "adjustment",
        source_type="adjustment",
        source_id=driver_id,
        lines=[(drv, Decimal(amount)), (other, -Decimal(amount))],
        actor_user_id=admin_id,
        reason="test",
        post=post,
        business_date=day,
    )


def test_cash_balances_by_age_receipts_by_collector_and_the_treasury(admin_client, db, company):
    t = today()
    d, e = make_driver(admin_client, company["id"]), make_driver(admin_client, company["id"])
    di, ei = _id(db, "people.employees", d["id"]), _id(db, "people.employees", e["id"])
    admin_id = db.execute(text("SELECT id FROM identity.users WHERE username = 'admin'")).scalar()
    for amount, days_ago in (("50", 10), ("30", 5), ("8", 4), ("20", 2), ("10", 0), ("-60", 1)):
        _journal(db, admin_id, di, amount, t - timedelta(days=days_ago))
    _journal(db, admin_id, di, "15", t, post=False)  # a collection waiting for approval: not aged
    _journal(db, admin_id, ei, "-5", t)  # the company owes him
    db.commit()
    rep = admin_client.get(f"{R}/cash", params={"date_from": str(t - timedelta(days=30)), "date_to": str(t)}).json()
    rows = {r["driver"]["id"]: r for r in rep["aging"]["rows"]}
    # handed in 60: the oldest 50 and 10 of the 30; held 20 (5 days) + 8 (4 days), 20 (2 days), 10 (today)
    mine = rows[d["id"]]
    assert (mine["posted"], mine["pending"]) == ("58.000", "15.000")
    assert (mine["d0_1"], mine["d2_3"], mine["d4_7"], mine["d8_plus"]) == ("10.000", "20.000", "28.000", "0.000")
    assert (mine["oldest"], mine["oldest_days"]) == (str(t - timedelta(days=5)), 5)
    assert rows[e["id"]]["posted"] == "-5.000" and rows[e["id"]]["d0_1"] == "0.000"

    # receipts: one kept, one reversed; the treasury takes them in, and a deposit takes 4 to the bank
    r1 = admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "12"}).json()
    r2 = admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "7"}).json()
    journal = db.execute(
        text(
            "SELECT j.public_id FROM cash.journals j JOIN cash.receipts r ON r.id = j.source_id "
            "WHERE j.source_type = 'receipt' AND j.kind <> 'reversal' AND r.public_id = :p"
        ),
        {"p": r2["id"]},
    ).scalar()
    assert (
        admin_client.post(f"/api/v1/cash/journals/{journal}/reverse", json={"reason": "wrong driver"}).status_code
        == 201
    )
    branch = r1["branch_id"]
    assert (
        admin_client.post(
            "/api/v1/cash/bank-deposits", json={"branch_id": branch, "amount": "4", "reference": "DEP-1"}
        ).status_code
        == 201
    )
    rep = admin_client.get(f"{R}/cash", params={"date_from": str(t), "date_to": str(t)}).json()
    [who] = rep["collectors"]
    assert (who["receipts"], who["amount"], who["reversed"]) == (1, "12.000", 1)
    [tr] = [x for x in rep["treasury"] if x["branch"]["id"] == branch]
    # in: 12 + 7; out: the reversed 7 and the 4 deposited
    assert (tr["opening"], tr["received"], tr["paid_out"], tr["closing"]) == ("0.000", "19.000", "11.000", "8.000")
    assert [(x["amount"], x["reference"]) for x in rep["deposits"]] == [("4.000", "DEP-1")]


def test_treasury_needs_its_permission(admin_client, new_client, company):
    make_user(admin_client, "cash_only", permissions=["reports.view", "cash.view", "reports.export"])
    c = new_client()
    login(c, "cash_only")
    params = {"date_from": str(today()), "date_to": str(today())}
    rep = c.get(f"{R}/cash", params=params).json()
    assert rep["treasury"] is None and rep["deposits"] is None
    assert c.get(f"{R}/cash", params=params | {"format": "xlsx", "section": "treasury"}).status_code == 403
    assert c.get(f"{R}/cash", params=params | {"format": "xlsx", "section": "aging"}).status_code == 200


def test_payroll_by_month_with_installments_by_source_and_what_was_carried(
    admin_client,
    client,
    staff,  # noqa: F811 (the fixture imported above)
    company,
    companies,
    new_client,
):
    a, b = staff["a"], staff["b"]
    set_cap(admin_client, "50.00", "gross")
    h = bearer(bind_device(client, a["phone"]))
    client.post(
        "/api/v1/driver/statements",
        headers=h,
        json={"month": MONTH.isoformat(), "screenshots": [shot(client, h)], "valid_days": 22},
    )
    sid = admin_client.get(f"{P}/statements").json()[0]["id"]
    admin_client.post(
        f"{P}/statements/{sid}/approve",
        json={"working_days": 26, "valid_days": 22, "bonus": "15.000", "tips": "5.000", "cancelled_orders": "10.000"},
    )
    for source, total, inst in (("other", "300.000", 2), ("advance", "20.000", 1)):
        r = admin_client.post(
            "/api/v1/deductions",
            json={
                "employee_id": a["id"],
                "source_type": source,
                "reason": f"{source} test",
                "total": total,
                "installments": inst,
                "start_month": MONTH.isoformat(),
            },
        )
        assert r.status_code == 201, r.text
    run = admin_client.post(f"{P}/runs", json={"company_id": company["id"], "month": MONTH.isoformat()}).json()
    month = {"month_from": MONTH.isoformat(), "month_to": MONTH.isoformat()}
    rep = admin_client.get(f"{R}/payroll", params=month).json()
    [line] = rep["runs"]
    assert (line["id"], line["status"], line["lines"]) == (run["id"], "draft", len(run["lines"]))
    assert line["gross"] == f"{sum(Decimal(x['gross']) for x in run['lines']):.3f}"
    assert line["net"] == f"{sum(Decimal(x['net']) for x in run['lines']):.3f}"
    # the cap left 140 of the 150 installment and none of the 20 advance: 30 carried
    assert (line["installments_due"], line["installments_deducted"], line["carried"]) == (
        "170.000",
        "140.000",
        "30.000",
    )
    assert line["by_source"] == {
        "advance": {"due": "20.000", "deducted": "0.000", "carried": "20.000"},
        "other": {"due": "150.000", "deducted": "140.000", "carried": "10.000"},
    }
    assert [(c["source_type"], c["carried"]) for c in rep["carried"]] == [("advance", "20.000"), ("other", "10.000")]
    # one employee: his line and his installments only
    only_b = admin_client.get(f"{R}/payroll", params=month | {"driver_id": b["id"]}).json()["runs"][0]
    assert (only_b["lines"], only_b["installments_due"], only_b["carried"]) == (1, "0.000", "0.000")
    xlsx = admin_client.get(f"{R}/payroll", params=month | {"format": "xlsx", "section": "carried"})
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    assert [c.value for c in ws[2]][-1] == 20
    r = admin_client.get(f"{R}/payroll", params={"month_from": "2024-01-01", "month_to": "2026-02-01"})
    assert r.status_code == 422 and r.json()["code"] == "range_too_long_months"
    # a user of another company sees none of it
    make_user(admin_client, "pay_b", permissions=["reports.view", "payroll.view"], company_ids=[companies["b"]["id"]])
    c = new_client()
    login(c, "pay_b")
    assert c.get(f"{R}/payroll", params=month).json()["runs"] == []


def test_daily_summary_by_driver_and_branch(admin_client, client, company):
    from tests.test_reports import driver_report

    d1, d2 = make_driver(admin_client, company["id"]), make_driver(admin_client, company["id"])
    for d, cash_ in ((d1, "5"), (d2, "7")):
        driver_report(client, bearer(bind_device(client, d["phone"])), cash_)
    params = {"date_from": str(today()), "date_to": str(today())}
    rows = admin_client.get(f"{R}/daily-summary", params=params | {"driver_id": d2["id"]}).json()["rows"]
    assert [r["driver"]["id"] for r in rows] == [d2["id"]]
    other = admin_client.post("/api/v1/branches", json={"name": {"ar": "فرع آخر", "en": "Other"}}).json()
    rows = admin_client.get(f"{R}/daily-summary", params=params | {"branch_id": other["id"]}).json()["rows"]
    assert rows == []
    r = admin_client.get(f"{R}/daily-summary", params=params | {"branch_id": 999_999})
    assert r.status_code == 422 and r.json()["code"] == "branch_not_found"
