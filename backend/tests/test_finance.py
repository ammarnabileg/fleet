"""Finance (BRD FR-FIN-01..05): expenses with their payment, entries made from every kind of document on the chart's
roles, approved and then never changed, reversed to correct, exported, and a trial balance that adds up to zero."""

import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal as D

import openpyxl
import pytest
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from app.core.clock import today
from tests.conftest import bearer, bind_device, jpeg, login, make_driver, make_employee, make_user, make_vehicle, upload

F = "/api/v1/finance"


def month() -> tuple[date, date]:
    first = today().replace(day=1)
    nxt = (first + timedelta(days=32)).replace(day=1)
    return first, nxt - timedelta(days=1)


def period() -> dict:
    first, last = month()
    return {"date_from": str(first), "date_to": str(last)}


def types(admin_client) -> dict:
    return {t["code"]: t["id"] for t in admin_client.get(f"{F}/expense-types").json()}


def expense(admin_client, company_id, **kw) -> dict:
    body = {
        "company_id": company_id,
        "type_id": types(admin_client)[kw.pop("kind", "fuel")],
        "expense_date": str(today()),
        "amount": "12.500",
        "payment_method": "treasury",
    } | kw
    r = admin_client.post(f"{F}/expenses", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def approve(admin_client, e) -> dict:
    r = admin_client.post(f"{F}/expenses/{e['id']}/approve", json={})
    assert r.status_code == 200, r.text
    return r.json()


def post(admin_client) -> dict:
    r = admin_client.post(f"{F}/entries/post", json=period())
    assert r.status_code == 200, r.text
    return r.json()


def entries(admin_client, **params) -> list[dict]:
    return admin_client.get(f"{F}/entries", params=period() | params).json()


def balances(admin_client) -> dict[str, str]:
    r = admin_client.get(f"{F}/trial-balance", params=period())
    assert r.status_code == 200, r.text
    return {b["account"]["code"]: b["closing"] for b in r.json()}


def payroll_settings(admin_client):
    version = admin_client.get("/api/v1/settings").json()["payroll"]["version"]
    r = admin_client.put(
        "/api/v1/settings/payroll",
        json={"version": version, "value": {"max_deduction_percent": "50.00", "deduction_cap_base": "gross"}},
    )
    assert r.status_code == 200, r.text


@pytest.fixture
def documents(admin_client, client, companies):
    """One document of each kind, with amounts chosen so every balance is known."""
    a, b = companies["a"]["id"], companies["b"]["id"]
    # the cash ledger: a driver's 20.000 approved, 15.000 handed in by receipt, 10.000 taken to the bank
    d = make_driver(admin_client, a)
    h = bearer(bind_device(client, d["phone"]))
    shot = client.post(
        "/api/v1/driver/files", params={"source": "upload"}, headers=h, files={"file": ("s.jpg", jpeg(), "image/jpeg")}
    ).json()["sha256"]
    report = client.post(
        "/api/v1/driver/reports",
        headers=h,
        json={"business_date": str(today()), "orders_count": 9, "cash_amount": "20", "screenshot_sha256": shot},
    ).json()
    assert admin_client.post(f"/api/v1/daily-reports/{report['id']}/approve", json={}).status_code == 200
    # another driver's report still waiting for review: its cash is in no entry until approved
    h2 = bearer(bind_device(client, make_driver(admin_client, a)["phone"]))
    shot2 = client.post(
        "/api/v1/driver/files", params={"source": "upload"}, headers=h2, files={"file": ("s.jpg", jpeg(), "image/jpeg")}
    ).json()["sha256"]
    pending = client.post(
        "/api/v1/driver/reports",
        headers=h2,
        json={"business_date": str(today()), "orders_count": 3, "cash_amount": "7", "screenshot_sha256": shot2},
    )
    assert pending.status_code == 201, pending.text
    receipt = admin_client.post("/api/v1/cash/receipts", json={"driver_id": d["id"], "amount": "15"}).json()
    r = admin_client.post(
        "/api/v1/cash/bank-deposits", json={"branch_id": receipt["branch_id"], "amount": "10", "reference": "NBK-7"}
    )
    assert r.status_code == 201, r.text
    # fuel paid from the treasury; a supplier's invoice of 100.000 paid later from the bank
    approve(admin_client, expense(admin_client, a, quantity="40"))
    supplier = approve(admin_client, expense(admin_client, a, kind="other", amount="100", payment_method="payable"))
    r = admin_client.post(f"{F}/expenses/{supplier['id']}/pay", json={"paid_from": "bank", "payment_ref": "TR-9"})
    assert r.status_code == 200, r.text
    # a maintenance center's invoice of 45.000, approved and paid
    c = admin_client.post("/api/v1/maintenance/centers", json={"name": "Garage", "specialty": "mechanical"}).json()
    inv = admin_client.post(
        "/api/v1/maintenance/invoices",
        json={
            "center_id": c["id"],
            "company_id": a,
            "number": "G-1",
            "invoice_date": str(today()),
            "total": "45",
            "file_sha256": upload(admin_client),
            "items": [{"kind": "other", "description": "Brakes", "unit_price": "45"}],
        },
    ).json()
    assert admin_client.post(f"/api/v1/maintenance/invoices/{inv['id']}/approve", json={}).status_code == 200
    assert admin_client.post(f"/api/v1/maintenance/invoices/{inv['id']}/paid", json={}).status_code == 200
    # an advance of 50.000 to an office employee of company b, taken back in this month's payroll of 300.000
    office = make_employee(admin_client, b, basic_salary="300.000", payment_method="cash")
    first, _ = month()
    r = admin_client.post(
        "/api/v1/deductions",
        json={
            "employee_id": office["id"],
            "source_type": "advance",
            "reason": "advance",
            "total": "50",
            "start_month": str(first),
        },
    )
    assert r.status_code == 201, r.text
    payroll_settings(admin_client)
    run = admin_client.post("/api/v1/payroll/runs", json={"company_id": b, "month": str(first)}).json()
    assert admin_client.post(f"/api/v1/payroll/runs/{run['id']}/approve").status_code == 200
    assert admin_client.post(f"/api/v1/payroll/runs/{run['id']}/paid", json={"payment_ref": "WPS-1"}).status_code == 200
    # a fine of 30.000 paid to the traffic department
    v = make_vehicle(admin_client, a)
    fine = admin_client.post(
        "/api/v1/fines",
        json={
            "vehicle_id": v["id"],
            "occurred_at": datetime.now(UTC).isoformat(),
            "violation": "speeding",
            "amount": "30",
        },
    ).json()
    assert admin_client.post(f"/api/v1/fines/{fine['id']}/paid", json={"payment_ref": "MOI-1"}).status_code == 200
    return {"driver": d, "supplier": supplier, "run": run, "office": office}


def test_every_document_makes_its_entry_and_the_books_balance(admin_client, documents):
    result = post(admin_client)
    assert result["errors"] == [] and result["stale"] == []
    assert result["by_kind"] == {
        "cash_journal": 3,
        "expense": 2,
        "expense_payment": 1,
        "maintenance_invoice": 1,
        "invoice_payment": 1,
        "payroll_run": 1,
        "payroll_payment": 1,
        "deduction": 1,
        "fine_payment": 1,
    }
    assert post(admin_client)["created"] == 0  # run again: nothing twice

    drafts = entries(admin_client, status="draft")
    run = next(e for e in drafts if e["source_kind"] == "payroll_run")
    detail = admin_client.get(f"{F}/entries/{run['id']}").json()
    assert [(ln["account"]["code"], ln["debit"], ln["credit"]) for ln in detail["lines"]] == [
        ("6110", "300.000", "0.000"),
        ("2120", "0.000", "250.000"),
        ("1150", "0.000", "50.000"),
    ]
    assert detail["entry_date"] == str(month()[1])  # the month's salaries, at its end
    fuel = next(e for e in drafts if e["source_kind"] == "expense" and e["amount"] == "12.500")
    assert fuel["description"].startswith("مصروف #") and "وقود" in fuel["description"]

    # nothing in the trial balance until approved; then every balance as worked out by hand, summing to zero
    assert balances(admin_client) == {}
    r = admin_client.post(f"{F}/entries/approve", json=period())
    assert r.json() == {"count": 12}
    tb = balances(admin_client)
    assert tb == {
        "1110": "-57.500",  # +15 receipt -10 to the bank -12.5 fuel -50 advance
        "1120": "-415.000",  # +10 deposit -100 supplier -45 invoice -250 salaries -30 fine
        "1130": "5.000",  # +20 collected -15 handed in
        "1150": "0.000",  # +50 advance -50 taken back in payroll
        "2130": "-20.000",
        "2110": "0.000",
        "2120": "0.000",
        "6110": "300.000",
        "5130": "45.000",
        "5170": "30.000",
        "5120": "12.500",
        "6190": "100.000",
    }
    assert sum(D(v) for v in tb.values()) == 0

    # the export: one row per line, Arabic headers, numbers as numbers
    r = admin_client.get(f"{F}/entries/export", params=period(), headers={"Accept-Language": "ar"})
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:6] == ("رقم القيد", "التاريخ", "رمز الحساب", "اسم الحساب", "مدين", "دائن")
    assert len(rows) == 1 + 25 and ws.sheet_view.rightToLeft
    assert all(isinstance(r[2], str) for r in rows[1:])  # account codes stay text
    assert sum(r[4] or 0 for r in rows[1:]) == sum(r[5] or 0 for r in rows[1:]) == D("977.500")


def test_approved_entries_never_change_and_a_reversal_frees_the_document(admin_client, documents, owner_db, companies):
    post(admin_client)
    admin_client.post(f"{F}/entries/approve", json=period())
    supplier = next(
        e
        for e in entries(admin_client)
        if e["source_ref"] == f"EXP-{documents['supplier']['number']}" and e["source_kind"] == "expense"
    )
    # the database itself refuses any change
    for sql in (
        "UPDATE finance.entry_lines SET debit = debit + 1 WHERE debit > 0",
        "UPDATE finance.entries SET description = 'x'",
        "DELETE FROM finance.entries",
    ):
        with pytest.raises(DBAPIError):
            owner_db.execute(text(sql))
        owner_db.rollback()
    assert admin_client.post(f"{F}/entries/discard", json=period()).json() == {"count": 0}  # drafts only

    # reversed with a reason: the mirror image, approved; the document can be entered again
    r = admin_client.post(f"{F}/entries/{supplier['id']}/reverse", json={"reason": "wrong account"})
    assert r.status_code == 200, r.text
    rev = r.json()
    assert rev["source_kind"] == "reversal" and rev["status"] == "approved" and rev["reverses"] == supplier["number"]
    assert [(ln["account"]["code"], ln["debit"], ln["credit"]) for ln in rev["lines"]] == [
        ("6190", "0.000", "100.000"),
        ("2110", "100.000", "0.000"),
    ]
    again = admin_client.post(f"{F}/entries/{supplier['id']}/reverse", json={"reason": "again please"})
    assert again.status_code == 409 and again.json()["code"] == "entry_already_reversed"
    assert post(admin_client)["by_kind"] == {"expense": 1}  # entered anew
    assert balances(admin_client)["6190"] == "0.000"  # the new draft is not approved yet


def test_a_cancelled_expense_is_listed_to_reverse_and_not_entered_again(admin_client, company):
    e = approve(admin_client, expense(admin_client, company["id"], amount="8"))
    post(admin_client)
    admin_client.post(f"{F}/entries/approve", json=period())
    r = admin_client.post(f"{F}/expenses/{e['id']}/cancel", json={"reason": "entered twice"})
    assert r.status_code == 200 and r.json()["status"] == "cancelled" and r.json()["decided_by"]  # who approved stays
    stale = post(admin_client)["stale"]
    assert [(s["source_ref"], s["entered"], s["document"]) for s in stale] == [(f"EXP-{e['number']}", "8.000", "0.000")]
    entry = next(x for x in entries(admin_client) if x["id"] == stale[0]["id"])
    admin_client.post(f"{F}/entries/{entry['id']}/reverse", json={"reason": "expense cancelled"})
    result = post(admin_client)
    assert result["created"] == 0 and result["stale"] == []
    assert balances(admin_client)["5120"] == "0.000"


def test_the_chart_is_the_accountants(admin_client, company, owner_db):
    accounts = {a["code"]: a for a in admin_client.get(f"{F}/accounts").json()}
    assert accounts["1110"]["roles"] == ["deduction_advance", "treasury"] and not accounts["1110"]["used"]
    # his own fuel account, the fuel type pointed at it; drafts made before keep the old one until made again
    approve(admin_client, expense(admin_client, company["id"]))
    post(admin_client)
    r = admin_client.post(
        f"{F}/accounts", json={"code": "6201", "name": {"ar": "وقود السيارات", "en": "Vehicle fuel"}, "type": "expense"}
    )
    assert r.status_code == 201, r.text
    fuel = r.json()
    dup = admin_client.post(f"{F}/accounts", json={"code": "6201", "name": {"ar": "س", "en": "x"}, "type": "expense"})
    assert dup.status_code == 409 and dup.json()["code"] == "account_code_taken"
    fuel_type = types(admin_client)["fuel"]
    assert admin_client.patch(f"{F}/expense-types/{fuel_type}", json={"account_id": fuel["id"]}).status_code == 200
    [draft] = entries(admin_client, status="draft")
    assert admin_client.get(f"{F}/entries/{draft['id']}").json()["lines"][0]["account"]["code"] == "5120"
    assert admin_client.post(f"{F}/entries/discard", json=period()).json() == {"count": 1}
    post(admin_client)
    [draft] = entries(admin_client, status="draft")
    assert admin_client.get(f"{F}/entries/{draft['id']}").json()["lines"][0]["account"]["code"] == "6201"
    admin_client.post(f"{F}/entries/approve", json={"ids": [draft["id"]]})
    # an account with entries keeps its code, takes a new name; one a role points at stays open
    used = admin_client.get(f"{F}/accounts").json()
    six = next(a for a in used if a["code"] == "6201")
    r = admin_client.patch(f"{F}/accounts/{six['id']}", json={"version": six["version"], "code": "6202"})
    assert r.status_code == 409 and r.json()["code"] == "account_in_use"
    r = admin_client.patch(
        f"{F}/accounts/{six['id']}", json={"version": six["version"], "name": {"ar": "الوقود", "en": "Fuel"}}
    )
    assert r.status_code == 200 and r.json()["used"]
    treasury = next(a for a in used if a["code"] == "1110")
    r = admin_client.patch(f"{F}/accounts/{treasury['id']}", json={"version": treasury["version"], "active": False})
    assert r.status_code == 409 and r.json()["code"] == "account_in_role"
    # a role without its account: that document waits, the others go in
    owner_db.execute(text("DELETE FROM finance.account_roles WHERE role = 'treasury'"))
    owner_db.commit()
    approve(admin_client, expense(admin_client, company["id"], amount="3"))
    result = post(admin_client)
    assert [x["code"] for x in result["errors"]] == ["account_role_missing"] and result["created"] == 0
    r = admin_client.put(f"{F}/roles", json={"roles": {"treasury": treasury["id"]}})
    assert r.status_code == 200
    assert post(admin_client)["created"] == 1
    r = admin_client.put(f"{F}/roles", json={"roles": {"not_a_role": treasury["id"]}})
    assert r.status_code == 422 and r.json()["code"] == "unknown_account_role"


def test_expense_rules(admin_client, companies):
    a, b = companies["a"]["id"], companies["b"]["id"]
    other_vehicle = make_vehicle(admin_client, b)
    body = {"company_id": a, "type_id": types(admin_client)["fuel"], "expense_date": str(today()), "amount": "5",
            "payment_method": "treasury"}  # fmt: skip
    r = admin_client.post(f"{F}/expenses", json=body | {"vehicle_id": other_vehicle["id"]})
    assert r.status_code == 422 and r.json()["code"] == "expense_company_mismatch"
    # a receipt photo attached, and served back
    sha = upload(admin_client)
    vehicle = make_vehicle(admin_client, a)
    e = expense(admin_client, a, vehicle_id=vehicle["id"], files=[sha], supplier="Oula", reference_no="R-1")
    assert e["vehicle"]["plate_number"] == vehicle["plate_number"] and e["files"] == [sha]
    assert admin_client.get(f"{F}/expenses/{e['id']}/files/{sha}").status_code == 200
    assert admin_client.get(f"{F}/expenses", params={"vehicle_id": vehicle["id"]}).json()[0]["id"] == e["id"]
    # refused without a reason; decided once; paid only when entered as a supplier invoice
    assert admin_client.post(f"{F}/expenses/{e['id']}/reject", json={}).status_code == 422
    approve(admin_client, e)
    assert admin_client.post(f"{F}/expenses/{e['id']}/approve", json={}).json()["code"] == "expense_not_pending"
    r = admin_client.post(f"{F}/expenses/{e['id']}/pay", json={"paid_from": "bank"})
    assert r.status_code == 409 and r.json()["code"] == "expense_not_payable"
    later = approve(admin_client, expense(admin_client, a, payment_method="payable"))
    assert admin_client.get(f"{F}/expenses", params={"unpaid": True}).json()[0]["id"] == later["id"]
    paid = admin_client.post(f"{F}/expenses/{later['id']}/pay", json={"paid_from": "treasury"}).json()
    assert paid["paid_from"] == "treasury" and not paid["unpaid"]
    twice = admin_client.post(f"{F}/expenses/{later['id']}/pay", json={"paid_from": "bank"})
    assert twice.status_code == 409 and twice.json()["code"] == "expense_not_payable"
    r = admin_client.post(f"{F}/expenses/{later['id']}/cancel", json={"reason": "too late now"})
    assert r.status_code == 409 and r.json()["code"] == "expense_not_cancellable"


def test_permissions_and_company_scope(admin_client, new_client, companies):
    a, b = companies["a"]["id"], companies["b"]["id"]
    e = expense(admin_client, b)
    make_user(admin_client, "clerk_a", permissions=["finance.view", "finance.create"], company_ids=[a])
    c = new_client()
    login(c, "clerk_a")
    assert c.get(f"{F}/expenses/{e['id']}").status_code == 404
    assert c.get(f"{F}/expenses").json() == []
    mine = expense(c, a)
    assert c.post(f"{F}/expenses/{mine['id']}/approve", json={}).status_code == 403  # finance.approve
    r = c.post(f"{F}/entries/post", json=period())
    assert r.status_code == 403 and r.json()["code"] == "all_companies_required"  # the ledger is the whole business's
    assert c.get(f"{F}/entries/export", params=period()).status_code == 403  # finance.export
    make_user(admin_client, "reader", permissions=["finance.view"])
    c = new_client()
    login(c, "reader")
    assert c.get(f"{F}/accounts").status_code == 200
    assert c.post(f"{F}/expenses", json={}).status_code in (403, 422)
    r = c.post(f"{F}/accounts", json={"code": "9", "name": {"ar": "س", "en": "x"}, "type": "asset"})
    assert r.status_code == 403


def test_a_reopened_payroll_and_a_cancelled_deduction_are_listed_to_reverse(admin_client, companies):
    b = companies["b"]["id"]
    first, _ = month()
    office = make_employee(admin_client, b, basic_salary="200.000", payment_method="cash")
    r = admin_client.post(
        "/api/v1/deductions",
        json={"employee_id": office["id"], "source_type": "sim", "reason": "sim card", "total": "12", "installments": 2,
              "start_month": str(first + timedelta(days=40))},
    )  # fmt: skip
    assert r.status_code == 201, r.text
    deduction = r.json()
    payroll_settings(admin_client)
    run = admin_client.post("/api/v1/payroll/runs", json={"company_id": b, "month": str(first)}).json()
    assert admin_client.post(f"/api/v1/payroll/runs/{run['id']}/approve").status_code == 200
    post(admin_client)
    admin_client.post(f"{F}/entries/approve", json=period())
    # a draft is never exported unless asked for
    approve(admin_client, expense(admin_client, b, amount="2"))
    post(admin_client)
    csv = admin_client.get(f"{F}/entries/export", params=period() | {"format": "csv"}).content.decode("utf-8-sig")
    assert "EXP-" not in csv and "PAY-" in csv
    assert (
        "EXP-" in admin_client.get(f"{F}/entries/export", params=period() | {"format": "csv", "status": "draft"}).text
    )
    # the run reopened, the deduction cancelled before anything was taken: both listed, nothing changed by itself
    assert (
        admin_client.post(f"/api/v1/payroll/runs/{run['id']}/reopen", json={"reason": "a wrong figure"}).status_code
        == 200
    )
    r = admin_client.post(f"/api/v1/deductions/{deduction['id']}/cancel", json={"reason": "card returned"})
    assert r.status_code == 200, r.text
    stale = post(admin_client)["stale"]
    assert sorted((s["source_ref"][:4], s["entered"], s["document"]) for s in stale) == [
        ("DED-", "12.000", "0.000"),
        ("PAY-", "200.000", "0.000"),
    ]


def test_the_nightly_posting_enters_the_last_days(admin_client, company, db):
    from app.modules.finance import tasks

    approve(
        admin_client, expense(admin_client, company["id"], amount="4", expense_date=str(today() - timedelta(days=40)))
    )
    approve(
        admin_client, expense(admin_client, company["id"], amount="6", expense_date=str(today() - timedelta(days=50)))
    )
    assert tasks.post_entries() == {"created": 1, "stale": 0, "errors": 0}  # 45 days back, not 50
    assert tasks.post_entries()["created"] == 0


def test_a_code_with_a_leading_zero_stays_text_in_excel():
    from app.core import sheets

    data = sheets.to_xlsx("t", ["code", "amount"], [["0101", "12.500"]], rtl=False, text_columns=(0,))
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert (ws.cell(2, 1).value, ws.cell(2, 1).data_type, ws.cell(2, 2).value) == ("0101", "s", D("12.5"))


def test_the_default_chart_is_structured_and_salaries_go_to_two_accounts(admin_client, company):
    """Four-digit codes whose first digit is the class; every role on an account; drivers' salaries an operating
    cost, the office's an administrative one."""
    accounts = admin_client.get(f"{F}/accounts").json()
    kind = {"1": "asset", "2": "liability", "3": "equity", "4": "income", "5": "expense", "6": "expense"}
    assert all(len(a["code"]) == 4 and a["type"] == kind[a["code"][0]] for a in accounts)
    by_code = {a["code"]: a for a in accounts}
    assert by_code["5110"]["roles"] == ["driver_salaries_expense"] and by_code["6110"]["roles"] == ["salaries_expense"]
    assert {"2210", "3120", "6150"} <= set(by_code)  # end-of-service provision, owner's account, government fees
    mapped = {r["role"] for r in admin_client.get(f"{F}/roles").json()}
    assert len(mapped) == 20
    assert {
        t["code"]: by_code_of(accounts, t["account_id"]) for t in admin_client.get(f"{F}/expense-types").json()
    } == {
        "fuel": "5120",
        "repairs": "5130",
        "tyres": "5140",
        "vehicle_insurance": "5150",
        "registration": "5160",
        "rent": "6130",
        "telecom": "6140",
        "gov_fees": "6150",
        "office": "6160",
        "bank_fees": "6170",
        "other": "6190",
    }
    # the chart in Excel for the accountant: class, type and what each account is used for
    r = admin_client.get(f"{F}/accounts/export", headers={"Accept-Language": "ar"})
    assert r.status_code == 200, r.text
    rows = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("الرمز", "الحساب", "الفئة", "النوع", "يُستخدم في", "الحالة")
    drivers = next(row for row in rows if row[0] == "5110")
    assert drivers[1:5] == ("رواتب السائقين", "تكاليف التشغيل", "مصروفات", "رواتب السائقين")
    fuel = next(row for row in rows if row[0] == "5120")
    assert fuel[4] == "وقود" and len(rows) == 1 + len(accounts)
    # a month with a driver (200, of which 20 an advance taken back) and an office employee (300): the drivers'
    # salaries are what they cost, net and installments, apart from the office's
    driver = make_driver(admin_client, company["id"], basic_salary="200.000", payment_method="cash")
    make_employee(admin_client, company["id"], basic_salary="300.000", payment_method="cash")
    payroll_settings(admin_client)
    first, _ = month()
    r = admin_client.post(
        "/api/v1/deductions",
        json={"employee_id": driver["id"], "source_type": "advance", "reason": "advance", "total": "20",
              "start_month": str(first)},
    )  # fmt: skip
    assert r.status_code == 201, r.text
    run = admin_client.post("/api/v1/payroll/runs", json={"company_id": company["id"], "month": str(first)}).json()
    r = admin_client.post(f"/api/v1/payroll/runs/{run['id']}/approve")
    assert r.status_code == 200, r.text
    post(admin_client)
    entry = next(e for e in entries(admin_client, status="draft") if e["source_kind"] == "payroll_run")
    lines = admin_client.get(f"{F}/entries/{entry['id']}").json()["lines"]
    assert [(ln["account"]["code"], ln["debit"], ln["credit"]) for ln in lines] == [
        ("5110", "200.000", "0.000"),
        ("6110", "300.000", "0.000"),
        ("2120", "0.000", "480.000"),
        ("1150", "0.000", "20.000"),
    ]


def by_code_of(accounts, account_id) -> str:
    return next(a["code"] for a in accounts if a["id"] == account_id)
